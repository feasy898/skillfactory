#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""harness/ws4_judge.py — WS4 dev 题池双层判分器核心（SF-0005 卡 A）。

双层结构（卡面裁定，SYNTHESIS §三）：
- L2 主判 = 本模块，检查名一律 task_outcome_*（PLAN §3.1「L2 只许 task_outcome_*」）；
  判定基准 = **当前输入**（--input-basis，默认包内 fixtures/）——尺子跟着输入走
  （TESTS AR-2 原则），nonce 交叉（BQ-9）天然红。
- L1 子轨 = 桥接 dist 原版 runner（import 透传，零改动铁律），显式传参照根，
  **只登记不判分**（结果落 record.l1；崩溃类异常才算装置 FAIL）。

各包入口 = packs/ws4-dev/<pkg>/checks/run_check.py 薄包装（声明本包检查名清单，
judge 核对契约一致后执行；清单与实现不一致 = 用法错误 exit 2）。

用法（由包装转发）：
    python checks/run_check.py <被测产物根> [--input-basis <输入基准目录>] [--record <path>]
退出码：0 L2 全绿 / 1 任一红 / 2 用法错误。
"""

import argparse
import importlib.util
import io
import json
import os
import re
import sys
import tempfile
import zipfile

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt

HARNESS_DIR = os.path.dirname(os.path.abspath(__file__))
EVALKIT = os.path.dirname(HARNESS_DIR)
WORKSPACE = os.path.dirname(EVALKIT)
DIST_ROOT = os.path.join(WORKSPACE, "afp-clone", "zcode-research", "skillfactory", "dist")

MM_SECTIONS = ("一、决议事项", "二、待办事项", "三、风险与关注")
MM_SECTION_KEYS = ("决议事项", "待办事项", "风险与关注")
TODO_HEADER = ("事项", "负责人", "期限")
ENTRY_RE = re.compile(r"^\s*(\d+)\.\s*【\s*([^·】]+?)\s*·\s*([^·】]+?)\s*】\s*(.*)$", re.S)
ATTRIB_SUFFIX = "（归属待核）"
FABRIC_TS_RE = re.compile(r"\[\d{1,2}:\d{2}\]")

# 《广告法》第九条及市场监管公开解读词样（与生成器同源冻结；公开法条派生）
BANNED_WORDS = ["国家级", "世界级", "最高级", "最佳", "第一", "顶级", "独家", "万能",
                "百分百", "绝对", "永久有效", "全网第一"]

HOT_ELEMENTS = {
    "dy": ["钩子", "痛点", "价值点1", "价值点2", "价值点3", "行动号召"],
    "xhs": ["标题", "开头", "要点1", "要点2", "要点3", "标签"],
    "wx": ["引入", "论点1", "论点2", "论点3", "总结"],
}
HOT_TOP_KEYS = ["template_version", "platform", "platform_name", "structure_name", "topic",
                "points_input_count", "points_used", "points_unused", "element_count",
                "elements", "placeholder_stats"]
EMOJI_RE = re.compile(r"[\U0001F300-\U0001FAFF\u2600-\u27BF]")
HASHTAG_RE = re.compile(r"#[^#\s]+")

OFC_SALUTATION = {"周报": False, "请示函": True, "会议通知": True}


def load_task_cfg(pack_dir):
    with io.open(os.path.join(pack_dir, "task.json"), encoding="utf-8") as f:
        return json.load(f)


def check(name, ok, detail):
    return {"name": name, "pass": bool(ok), "detail": detail}


def read_text(path):
    with io.open(path, encoding="utf-8") as f:
        return f.read()


def docx_all_text(doc):
    parts = [p.text for p in doc.paragraphs]
    for t in doc.tables:
        for row in t.rows:
            parts.extend(c.text for c in row.cells)
    return "\n".join(parts)


def docx_opens_safe(path):
    """docx 存在、可打开、zip 条目无 DOCTYPE/ENTITY（XXE 面）。"""
    if not os.path.isfile(path):
        return False, "产物缺失: %s" % os.path.basename(path)
    try:
        with zipfile.ZipFile(path) as zf:
            for n in zf.namelist():
                head = zf.read(n)[:512]
                if b"<!DOCTYPE" in head or b"<!ENTITY" in head:
                    return False, "zip 条目含 DOCTYPE/ENTITY: %s" % n
        Document(path)
        return True, "opens ok"
    except Exception as exc:  # noqa: BLE001
        return False, "打开失败: %s" % exc


def short_circuit(declared, first_name, detail):
    """产物缺失时的短路输出：首个检查记红因，其余声明检查逐条记红（短路）。
    保证任何产物状态都发满声明清单 → 缺产物=产物红（exit 1），不是用法错误（exit 2）。"""
    out = []
    for name in declared:
        if name == first_name:
            out.append(check(name, False, detail))
        else:
            out.append(check(name, False, "短路：%s" % detail))
    return out


# ================================================================= mm 线 ==
def judge_mm(product_root, basis, cfg):
    declared = cfg.get("expected_checks") or []
    checks = []
    docx_path = os.path.join(product_root, "纪要.docx")
    ok, det = docx_opens_safe(docx_path)
    if not ok:
        return short_circuit(declared, "task_outcome_docx_present_opens", det)
    checks.append(check("task_outcome_docx_present_opens", ok, det))
    transcript_path = os.path.join(basis, "transcript.txt")
    transcript = read_text(transcript_path) if os.path.isfile(transcript_path) else ""
    speakers = set()
    for line in transcript.splitlines():
        m = re.match(r"^(?:\[\d{1,2}:\d{2}\]\s*)?([^:]{1,12}):\s", line)
        if m:
            speakers.add(m.group(1).strip())

    doc = Document(docx_path)
    paras = [p.text for p in doc.paragraphs]

    # 三节标题
    missing = [s for s in MM_SECTIONS if s not in paras]
    checks.append(check("task_outcome_three_sections", not missing,
                        "三节标题齐" if not missing else "缺节标题: %s" % missing))

    # 分节解析条目
    cur = None
    entries = {k: [] for k in MM_SECTION_KEYS}
    for text in paras:
        if text in MM_SECTIONS:
            cur = MM_SECTION_KEYS[MM_SECTIONS.index(text)]
            continue
        m = ENTRY_RE.match(text)
        if m and cur:
            entries[cur].append({"idx": m.group(1), "cat": m.group(2).strip(),
                                 "owner": m.group(3).strip(), "content": m.group(4).strip()})
    checks.append(check("task_outcome_entries_parseable",
                        any(entries[k] for k in ("决议事项", "风险与关注")),
                        "决议 %d 条 / 风险 %d 条" % (len(entries["决议事项"]), len(entries["风险与关注"]))))

    # 待办表格表头
    header_ok, rows = False, []
    for t in doc.tables:
        hdr = tuple(c.text.strip() for c in t.rows[0].cells)
        if hdr == TODO_HEADER:
            header_ok = True
            rows = [[c.text.strip() for c in r.cells] for r in t.rows[1:]]
            break
    checks.append(check("task_outcome_todo_table_header", header_ok,
                        "表头=事项/负责人/期限" if header_ok else "未找到三列表头表格"))

    # 逐字溯源（允许（归属待核）后缀；内容必须逐字存在于当前转写稿）
    untrace, bad_owner = [], []
    for key in ("决议事项", "风险与关注"):
        for e in entries[key]:
            content = e["content"]
            if content.endswith(ATTRIB_SUFFIX):
                content = content[: -len(ATTRIB_SUFFIX)].strip()
            if content and content not in transcript:
                untrace.append("%s#%s" % (key, e["idx"]))
            if e["owner"] and e["owner"] not in speakers:
                bad_owner.append("%s#%s(%s)" % (key, e["idx"], e["owner"]))
    checks.append(check("task_outcome_entries_traceable_to_input",
                        not untrace and entries["决议事项"],
                        "全部逐字可溯源" if not untrace else "不可溯源: %s" % untrace[:5]))
    checks.append(check("task_outcome_entry_owner_from_input",
                        not bad_owner,
                        "相关人均为转写稿说话人" if not bad_owner else "编造相关人: %s" % bad_owner[:5]))

    # 待办行溯源
    bad_rows = []
    for r in rows:
        if len(r) < 3:
            bad_rows.append(r)
            continue
        item, owner, dl = r[0], r[1], r[2]
        if not item or item not in transcript:
            bad_rows.append("事项不可溯源: %s" % item[:20])
        elif owner not in speakers:
            bad_rows.append("负责人不在转写稿: %s" % owner)
        elif dl != "未明确" and dl not in item:
            bad_rows.append("期限非该条来源: %s" % dl)
    checks.append(check("task_outcome_todo_rows_traceable",
                        header_ok and rows and not bad_rows,
                        "%d 行待办全部可溯源" % len(rows) if not bad_rows else str(bad_rows[:5])))

    # summary 自洽
    sum_path = os.path.join(product_root, "summary.json")
    sum_ok, sum_det = False, "summary.json 缺失或不可解析"
    if os.path.isfile(sum_path):
        try:
            s = json.loads(read_text(sum_path))
            expected = {"决议事项": len(entries["决议事项"]), "待办事项": len(rows),
                        "风险与关注": len(entries["风险与关注"])}
            expected["合计"] = sum(expected.values())
            sum_ok = all(s.get(k) == v for k, v in expected.items()) and set(s) == set(expected)
            sum_det = "summary=%s expected=%s" % (s, expected)
        except Exception as exc:  # noqa: BLE001
            sum_det = "解析失败: %s" % exc
    checks.append(check("task_outcome_summary_selfconsistent", sum_ok, sum_det))

    # 近邻负例专属：禁止补造时间戳（输入无时间戳时）
    if cfg.get("no_timestamp_input"):
        full = docx_all_text(doc)
        hits = FABRIC_TS_RE.findall(full)
        checks.append(check("task_outcome_no_fabricated_timestamp", not hits,
                            "无补造时间戳" if not hits else "出现 [HH:MM]: %s" % hits[:3]))
    return checks


# ================================================================ hot 线 ==
def judge_hot(product_root, basis, cfg):
    checks = []
    md_path = os.path.join(product_root, "骨架.md")
    sj_path = os.path.join(product_root, "structure.json")
    inp = json.loads(read_text(os.path.join(basis, "input.json")))
    platform = inp["platform"]
    points = inp["points"]
    slots = 3

    md = read_text(md_path) if os.path.isfile(md_path) else None
    sj = None
    if os.path.isfile(sj_path):
        try:
            sj = json.loads(read_text(sj_path))
        except Exception:  # noqa: BLE001
            sj = None
    if sj is None or md is None:
        miss = ("structure.json 不可解析" if os.path.isfile(sj_path) else "structure.json 缺失")
        if md is None:
            miss = "骨架.md 缺失"
        return short_circuit(cfg.get("expected_checks") or [], "task_outcome_artifacts_present", miss)
    checks.append(check("task_outcome_artifacts_present", True,
                        "骨架.md %d 字符 / structure.json ok" % len(md)))

    # 结构自洽：顶层键 + points 记账 + element_count
    miss_keys = [k for k in HOT_TOP_KEYS if k not in sj]
    used, unused = sj.get("points_used", []), sj.get("points_unused", [])
    acc_ok = sorted(used + unused) == sorted(points) and sj.get("points_input_count") == len(points)
    ec_ok = sj.get("element_count") == len(sj.get("elements", []))
    checks.append(check("task_outcome_structure_selfconsistent",
                        not miss_keys and acc_ok and ec_ok,
                        "缺键=%s points记账=%s element_count=%s(%s/%s)" %
                        (miss_keys, acc_ok, ec_ok, sj.get("element_count"), len(sj.get("elements", [])))))

    # 冻结要素名序
    names = [e.get("name", "") for e in sj.get("elements", [])]
    frozen = HOT_ELEMENTS[platform]
    checks.append(check("task_outcome_platform_elements_frozen", names == frozen,
                        "names=%s frozen=%s" % (names, frozen)))

    elems_by_name = {e.get("name"): (e.get("content") or "") for e in sj.get("elements", [])}

    # 平台形态规则
    det = ""
    if platform == "dy":
        hook, cta = elems_by_name.get("钩子", ""), elems_by_name.get("行动号召", "")
        ok = len(hook) >= 8 and any(w in cta for w in ("关注", "评论", "私信", "点击"))
        det = "钩子len=%d cta含行动词=%s" % (len(hook), any(w in cta for w in ("关注", "评论", "私信", "点击")))
    elif platform == "xhs":
        title, tags = elems_by_name.get("标题", ""), elems_by_name.get("标签", "")
        n_tags = len(HASHTAG_RE.findall(tags))
        ok = bool(re.search(r"\d", title)) and bool(EMOJI_RE.search(title)) and 3 <= n_tags <= 8
        det = "标题含数字=%s 含emoji=%s 标签数=%d" % (bool(re.search(r"\d", title)),
                                                    bool(EMOJI_RE.search(title)), n_tags)
    else:  # wx
        intro, summary_e = elems_by_name.get("引入", ""), elems_by_name.get("总结", "")
        ok = len(intro) >= 8 and len(summary_e) >= 8
        det = "引入len=%d 总结len=%d" % (len(intro), len(summary_e))
    checks.append(check("task_outcome_platform_shape_rules", ok, det))

    # 卖点填充（每个槽位含对应卖点或占位）
    bad_slots = []
    for i in range(slots):
        slot = [v for k, v in elems_by_name.items() if k in
                ("价值点%d" % (i + 1), "要点%d" % (i + 1), "论点%d" % (i + 1))]
        content = slot[0] if slot else ""
        if i < len(used):
            if used[i] not in content:
                bad_slots.append("槽%d未含卖点" % (i + 1))
        else:
            if ("【占位:价值点%d】" % (i + 1)) not in content:
                bad_slots.append("槽%d未按占位规约" % (i + 1))
    checks.append(check("task_outcome_points_filled_in_slots", not bad_slots,
                        "3 槽全部合规" if not bad_slots else str(bad_slots)))

    # 占位记账
    exp_stats = {"slots": slots, "filled": len(used), "placeholders": max(0, slots - len(used))}
    got_stats = sj.get("placeholder_stats")
    md_ph = md.count("【占位:")
    ok = got_stats == exp_stats and md_ph == exp_stats["placeholders"]
    checks.append(check("task_outcome_placeholder_accounting", ok,
                        "stats=%s expected=%s md占位数=%d" % (got_stats, exp_stats, md_ph)))

    # 合规：广告法绝对化用语零命中（骨架.md + structure.json 语义文本）
    scan_text = md + json.dumps(sj, ensure_ascii=False)
    hits = [w for w in BANNED_WORDS if w in scan_text]
    checks.append(check("task_outcome_no_banned_absolute_terms", not hits,
                        "零命中" if not hits else "命中: %s" % hits))

    # 模板残留
    residue = "{{" in md or "{{" in json.dumps(sj, ensure_ascii=False)
    checks.append(check("task_outcome_no_template_residue", not residue,
                        "无 {{ }} 残留" if not residue else "存在模板残留"))
    return checks


# ================================================================ ofc 线 ==
def judge_ofc(product_root, basis, cfg):
    declared = cfg.get("expected_checks") or []
    docx_path = os.path.join(product_root, "文书.docx")
    ok, det = docx_opens_safe(docx_path)
    if not ok:
        return short_circuit(declared, "task_outcome_docx_present_opens_safe", det)
    checks = [check("task_outcome_docx_present_opens_safe", ok, det)]
    inp = json.loads(read_text(os.path.join(basis, "input.json")))
    template = inp["template"]
    exp_filled = {f["name"]: f["value"] for f in inp["fields"] if f["value"] is not None}
    exp_missing = [f["name"] for f in inp["fields"] if f["value"] is None]

    doc = Document(docx_path)
    paras = doc.paragraphs
    full = docx_all_text(doc)

    # 版式：标题居中
    first = next((p for p in paras if p.text.strip()), None)
    checks.append(check("task_outcome_layout_title_centered",
                        first is not None and first.alignment == WD_ALIGN_PARAGRAPH.CENTER,
                        "alignment=%s" % (first.alignment if first else None)))

    # 版式：称谓顶格（有称谓模板）
    if OFC_SALUTATION[template]:
        salut = next((p for p in paras[1:5] if re.match(r"^[^，。]{1,20}：$", p.text.strip())), None)
        ind = salut.paragraph_format.first_line_indent if salut else None
        checks.append(check("task_outcome_layout_salutation_flush",
                            salut is not None and (ind is None or ind == 0),
                            "称谓段=%s 缩进=%s" % (salut.text if salut else None, ind)))
    # 版式：正文首行缩进（20–28pt 视为两字符档）
    body = [p for p in paras[2:] if p.text.strip() and not re.match(r"^[^，。]{1,20}：$", p.text.strip())]
    body = [p for p in body if p.alignment != WD_ALIGN_PARAGRAPH.RIGHT]
    if body:
        in_band = 0
        for p in body:
            ind = p.paragraph_format.first_line_indent
            if ind is not None and Pt(20) <= ind <= Pt(28):
                in_band += 1
        ok = in_band >= max(1, int(len(body) * 0.6))
        checks.append(check("task_outcome_layout_body_indent", ok,
                            "%d/%d 段首行缩进 20-28pt" % (in_band, len(body))))
    # 版式：落款右对齐
    last = next((p for p in reversed(paras) if p.text.strip()), None)
    checks.append(check("task_outcome_layout_signature_right",
                        last is not None and last.alignment == WD_ALIGN_PARAGRAPH.RIGHT,
                        "alignment=%s text=%s" % (last.alignment if last else None,
                                                  (last.text[:12] if last else None))))

    # fields.json 自洽（以输入基准为期望源）
    fj_path = os.path.join(product_root, "fields.json")
    fj = None
    if os.path.isfile(fj_path):
        try:
            fj = json.loads(read_text(fj_path))
        except Exception:  # noqa: BLE001
            fj = None
    ok = fj is not None
    det = "fields.json 缺失/不可解析"
    if ok:
        got_filled = [f["name"] for f in fj["fields"] if f.get("status") == "filled"]
        got_missing = [f["name"] for f in fj["fields"] if f.get("status") == "missing"]
        problems = []
        if sorted(got_filled) != sorted(exp_filled):
            problems.append("filled 不符: %s vs 输入 %s" % (got_filled, sorted(exp_filled)))
        if sorted(got_missing) != sorted(exp_missing):
            problems.append("missing 不符: %s vs 输入 %s" % (got_missing, sorted(exp_missing)))
        if sorted(fj.get("filled_fields", [])) != sorted(got_filled):
            problems.append("filled_fields 与 fields 不一致")
        if sorted(fj.get("missing_fields", [])) != sorted(got_missing):
            problems.append("missing_fields 与 fields 不一致")
        if fj.get("title") != inp["title"]:
            problems.append("title 不符")
        ok = not problems
        det = "; ".join(problems) if problems else "记账与输入基准一致"
    checks.append(check("task_outcome_fields_selfconsistent", ok, det))

    if fj is None:
        checks.extend(check(n, False, "短路：fields.json 缺失/不可解析")
                      for n in declared[len(checks):])
        return checks

    # 值逐字在 docx
    absent = [n for n, v in exp_filled.items() if v not in full]
    checks.append(check("task_outcome_values_verbatim_in_docx", not absent,
                        "全部逐字在 docx" if not absent else "缺失: %s" % absent))

    # 缺失必填：docx 以 ____ 呈现且登记，不得虚构
    problems = []
    for n in exp_missing:
        if n not in full:
            problems.append("字段名未出现: %s" % n)
        elif ("%s：____" % n) not in full and ("%s:____" % n) not in full and "____" not in full:
            problems.append("未以 ____ 呈现: %s" % n)
    checks.append(check("task_outcome_missing_reported_and_blank", not problems,
                        "缺失项合规呈现" if not problems else str(problems)))
    return checks


# ============================================================ L1 桥接 ====
L1_DIST_PKG = {"mm": "meeting-minutes-skill", "hot": "hot-templates-skill", "ofc": "office-templates-skill"}


def l1_bridge(line):
    """import dist 原版 runner 透传：显式参照根自校验 + 空目录红路，结果只登记不判分。"""
    import contextlib
    dist_pkg = L1_DIST_PKG[line]
    runner_path = os.path.join(DIST_ROOT, dist_pkg, "eval", "runner.py")
    ref_root = os.path.join(DIST_ROOT, dist_pkg, "reference", "out")
    spec = importlib.util.spec_from_file_location("l1_orig_runner_%s" % line, runner_path)
    mod = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(mod)
    empty = tempfile.mkdtemp(prefix="ws4-l1-empty-")
    # argv 约定按包适配（实测）：mm/hot 的 main(argv) 自带 argv[0] 位（len==3），
    # ofc 的 main(argv) 走 argparse.parse_args（不含 argv[0]，两位位置参数）。
    l1_argv = ["l1_bridge", ref_root, ref_root] if line in ("mm", "hot") else [ref_root, ref_root]
    l1_argv_red = (["l1_bridge", empty, ref_root] if line in ("mm", "hot")
                   else [empty, ref_root])
    try:
        with contextlib.redirect_stdout(io.StringIO()):
            sc = mod.main(l1_argv)
            rd = mod.main(l1_argv_red)
    except SystemExit as exc:  # 防 runner 以 exit() 代替 return
        sc = rd = int(exc.code or 0)
    finally:
        import shutil as _sh
        _sh.rmtree(empty, ignore_errors=True)
    return {
        "dist_pkg": dist_pkg,
        "runner_path": runner_path,
        "argv_shape": "run_check.py <被测产物根> [--input-basis <输入基准目录>] [--record <path>]；"
                      "L1 子轨=原版 runner <产物根> <参照根>（显式传参，禁 0 参默认）",
        "ref_root_verbatim": ref_root,
        "selfcheck_exit": sc,
        "redpath_exit": rd,
        "controlled": sc in (0, 1) and rd in (0, 1),
    }


# ============================================================ CLI 核心 ====
def judge_pack(pack_dir, product_root, input_basis=None):
    cfg = load_task_cfg(pack_dir)
    basis = input_basis or os.path.join(pack_dir, "fixtures")
    line = cfg["line"]
    if line == "mm":
        checks = judge_mm(product_root, basis, cfg)
    elif line == "hot":
        checks = judge_hot(product_root, basis, cfg)
    else:
        checks = judge_ofc(product_root, basis, cfg)
    names = [c["name"] for c in checks]
    expected = cfg.get("expected_checks", [])
    contract_ok = (not expected) or (names == expected)
    return checks, contract_ok, cfg


def main_for(pack_dir, argv, declared=None):
    ap = argparse.ArgumentParser(add_help=False)
    ap.add_argument("roots", nargs="*")
    ap.add_argument("--input-basis", default=None)
    ap.add_argument("--record", default=None)
    args, unknown = ap.parse_known_args(argv)
    if unknown or len(args.roots) != 1:
        print("用法: run_check.py <被测产物根> [--input-basis <dir>] [--record <path>]", file=sys.stderr)
        return 2

    checks, contract_ok, cfg = judge_pack(pack_dir, args.roots[0], args.input_basis)
    if declared is not None and [c["name"] for c in checks] != declared:
        contract_ok = False
    try:
        l1 = l1_bridge(cfg["line"])
        l1_ok = l1["controlled"]
    except Exception as exc:  # noqa: BLE001
        l1 = {"bridge_error": str(exc)}
        l1_ok = False

    passed = all(c["pass"] for c in checks)
    exit_code = 0 if (passed and contract_ok and l1_ok) else (2 if not contract_ok else 1)
    record = {
        "track": "artifact", "layer": "L2",
        "pack": os.path.basename(pack_dir),
        "input_basis": args.input_basis or os.path.join(pack_dir, "fixtures"),
        "checks": checks,
        "checks_pass": passed,
        "l1_registered_not_judged": l1,
        "exit_code": exit_code,
    }
    print(json.dumps(record, ensure_ascii=False, indent=2))
    if args.record:
        os.makedirs(os.path.dirname(os.path.abspath(args.record)), exist_ok=True)
        with io.open(args.record, "w", encoding="utf-8", newline="\n") as f:
            f.write(json.dumps(record, ensure_ascii=False, indent=2) + "\n")
    if not contract_ok:
        print("检查名契约不符（task.json expected_checks）", file=sys.stderr)
    return exit_code


if __name__ == "__main__":
    if hasattr(sys.stdout, "reconfigure"):
        sys.stdout.reconfigure(encoding="utf-8", errors="replace")
    print("本模块为共享核心，请经 packs/ws4-dev/<pkg>/checks/run_check.py 调用", file=sys.stderr)
    sys.exit(2)
