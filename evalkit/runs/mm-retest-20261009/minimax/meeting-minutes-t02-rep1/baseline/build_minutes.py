# -*- coding: utf-8 -*-
"""从 brief.md 中的会议转写稿生成三段式会议纪要（docx + summary.json）。"""

import json
import os

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Pt, RGBColor

BASE = os.path.dirname(os.path.abspath(__file__))
TRANSCRIPT = os.path.join(BASE, "bailuzhou_lib_1014.txt")
OUT_DIR = os.path.join(BASE, "out")
DOCX_PATH = os.path.join(OUT_DIR, "纪要.docx")
JSON_PATH = os.path.join(OUT_DIR, "summary.json")

# ---------------------------------------------------------------- 转写稿落盘
with open(os.path.join(BASE, "brief.md"), encoding="utf-8") as f:
    lines = f.read().splitlines()

# brief.md 第 1 行为任务说明，第 2 行为空行，其后至首个空行为转写稿全文
start = next(i for i, l in enumerate(lines) if l.startswith("[14:00]"))
end = start
while end < len(lines) and lines[end].strip():
    end += 1
transcript_lines = lines[start:end]
with open(TRANSCRIPT, "w", encoding="utf-8") as f:
    f.write("\n".join(transcript_lines) + "\n")

# ---------------------------------------------------------------- 纪要内容
MEETING = {
    "title": "白鹭洲图书馆十周年活动筹备会议纪要",
    "meta": [
        ("会议时间", "14:00—14:41（转写稿时段）"),
        ("会议主题", "十周年活动筹备事项通气与决议"),
        ("主持人", "馆长 远山"),
        ("参会部门", "馆长室、活动部、古籍部、借阅部、少儿部、宣传、财务部、后勤组、志愿队"),
        ("转写来源", "bailuzhou_lib_1014.txt（共 %d 条发言记录）" % len(transcript_lines)),
    ],
}

DECISIONS = [
    ("展品清单定稿", "主题展「纸上山河」展品清单经策展组终审后，于今日下班前定稿。", "14:01／14:03"),
    ("开幕式时间", "开幕式定于 11 月 1 日上午 9:00 举行。", "14:08"),
    ("仪式合并", "捐赠仪式与开幕式合并举办，不单独排期。", "14:14"),
    ("活动三项结论", "展期十天；活动期间免预约；闭馆时间顺延一小时。", "14:21"),
    ("加班餐列支", "布展期间的加班餐从工会经费列支。", "14:27"),
    ("伴手礼核算", "伴手礼数量按报名人数核算，由活动部芦芽负责统计，核算口径采用 10 月 25 日的报名数。", "14:25／14:26"),
]

# 序号、事项、负责人/责任方、完成时限、状态、转写时间戳
TODO_COLUMNS = ["序号", "待办事项", "责任方", "完成时限", "状态", "时间戳"]
TODOS = [
    ("主题展「纸上山河」展品清单终审并定稿", "活动部芦芽、策展组", "今日下班前", "进行中", "14:01／14:03"),
    ("接收文创雪糕试制样", "活动部芦芽", "下周三", "进行中", "14:05"),
    ("绘本剧场场地对接与小剧场布置", "后勤组、后勤贺黎坚", "周六上午布置完成", "进行中", "14:06／14:07"),
    ("查看并反馈海报文案初稿", "宣传柏舟、全体部门", "未设定", "进行中", "14:09"),
    ("就印刷厂报价是否含装裱再问一轮", "宣传柏舟", "未设定", "进行中", "14:10"),
    ("发送预算科目拆分表至各位", "财务部水清越", "今日下班前", "进行中", "14:11"),
    ("将逾期费减免口径报修订组（口径待馆里拍板）", "借阅部何渡", "本月 15 日前", "待口径确定", "14:12"),
    ("下发人偶巡游服装尺码表至各队", "少儿部乌兰", "明日", "进行中", "14:13"),
    ("志愿者排班表由初版完善至定稿", "志愿队佟安", "活动前", "进行中", "14:15"),
    ("借调 8 台对讲机用于活动当天人流疏导（器材科库房周日开放）", "志愿队佟安、器材科", "周日下午", "进行中", "14:16／14:17"),
    ("雪糕试制发票抬头使用单位全称", "财务部水清越", "开票时", "待办", "14:19"),
    ("纪念徽章打样进度每周五与厂家碰会", "活动部芦芽", "每周五", "持续跟进", "14:20"),
    ("统计伴手礼数量（以 10 月 25 日报名数为口径）", "活动部芦芽", "报名数据确定后", "进行中", "14:25／14:26"),
    ("古籍修复室暂停预约一周，展品回库体检", "古籍部米鹿鸣", "一周内", "待办", "14:28"),
    ("绘本剧音效文件交付剧务", "少儿部乌兰", "已交付", "已完成", "14:24"),
    ("绘本剧彩排时间与剧务再次确认", "少儿部乌兰", "未设定", "待办", "14:34"),
    ("预算执行率月报照常报送", "财务部水清越", "每月", "持续跟进", "14:35"),
    ("邀约开幕式司仪人选（台里老朋友）", "宣传柏舟", "未设定", "待办", "14:36／14:37"),
    ("完成志愿者保险审批后方可上岗", "志愿队佟安", "活动前", "办理中", "14:38"),
    ("展签繁体字校对（已托中文系老师）", "古籍部米鹿鸣", "布展前", "进行中", "14:39"),
    ("各口进展材料报送馆长邮箱", "各部门负责人", "周五前", "待办", "14:32"),
    ("媒体邀请函模板定稿后交宣传柏舟一份", "责任方未明确", "未设定", "待办", "14:33"),
]

# 编号、风险与关注事项、说明、涉及方、关注度、转写时间戳
RISKS = [
    ("展柜备用锁具到货时间不确定", "备用锁具未到货将直接冲击布展排期，存在布展延期的可能。", "古籍部米鹿鸣", "高", "14:18"),
    ("逾期费减免口径尚未拍板", "口径需馆里拍板，且须于本月 15 日前报修订组，时间窗紧。", "借阅部何渡", "高", "14:12"),
    ("志愿者保险审批未完成", "保险审批走完才能上岗，可能压缩后续培训与排班时间。", "志愿队佟安", "高", "14:38"),
    ("印刷厂报价是否含装裱未确认", "需再问一轮报价，报价口径不清将影响制作成本与印制排期。", "宣传柏舟", "中", "14:10"),
    ("周年当天借阅上限是否放宽未定", "是否临时放宽借阅上限尚未拍板，高峰日接待能力存在不确定性。", "借阅部何渡", "中", "14:04"),
    ("活动期间自习区占座", "占座现象在活动期间风险最高，需提前安排现场管理措施。", "借阅部何渡", "中", "14:23"),
    ("志愿者排班表仍为初版", "排班表初版仅在共享盘，尚未定稿，与上岗审批存在衔接风险。", "志愿队佟安", "中", "14:15"),
    ("周末停车场管制", "周末停车场有马拉松管制，需提前引导观众改乘地铁。", "后勤贺黎坚", "中", "14:30"),
    ("媒体邀请函模板未定稿", "模板定稿后需交付宣传柏舟一份，交接责任方在会上未明确。", "责任方未明确", "中", "14:33"),
]


# ---------------------------------------------------------------- 文档生成
def set_cjk(run, font="仿宋", size=None, bold=False):
    run.font.name = font
    run.font.bold = bold
    if size is not None:
        run.font.size = Pt(size)
    rpr = run._element.get_or_add_rPr()
    rfonts = rpr.find(qn("w:rFonts"))
    if rfonts is None:
        rfonts = rpr.makeelement(qn("w:rFonts"), {})
        rpr.append(rfonts)
    rfonts.set(qn("w:eastAsia"), font)
    rfonts.set(qn("w:ascii"), font)
    rfonts.set(qn("w:hAnsi"), font)


def add_paragraph(doc, text, font="仿宋", size=12, bold=False, align=None,
                  space_after=6, indent_first=True):
    p = doc.add_paragraph()
    pf = p.paragraph_format
    pf.space_after = Pt(space_after)
    pf.line_spacing = 1.5
    if align is not None:
        p.alignment = align
    if indent_first:
        pf.first_line_indent = Pt(size * 2)
    set_cjk(p.add_run(text), font=font, size=size, bold=bold)
    return p


def fix_layout(table, widths_pt):
    """固定表格布局与列宽，避免 Word/LibreOffice 按内容重新分配列宽。"""
    table.autofit = False
    tbl_pr = table._tbl.tblPr
    layout = tbl_pr.find(qn("w:tblLayout"))
    if layout is None:
        layout = tbl_pr.makeelement(qn("w:tblLayout"), {})
        tbl_pr.append(layout)
    layout.set(qn("w:type"), "fixed")
    total = int(sum(widths_pt) * 20)
    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = tbl_pr.makeelement(qn("w:tblW"), {})
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), str(total))
    tbl_w.set(qn("w:type"), "dxa")
    grid = table._tbl.find(qn("w:tblGrid"))
    if grid is not None:
        for col, w in zip(grid.findall(qn("w:gridCol")), widths_pt):
            col.set(qn("w:w"), str(int(w * 20)))
    for row in table.rows:
        for cell, w in zip(row.cells, widths_pt):
            cell.width = Pt(w)


def add_heading(doc, text):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(12)
    p.paragraph_format.space_after = Pt(6)
    set_cjk(p.add_run(text), font="黑体", size=14, bold=True)
    return p


def build_docx():
    doc = Document()
    style = doc.styles["Normal"]
    style.font.name = "仿宋"
    style.font.size = Pt(12)
    style.element.rPr.rFonts.set(qn("w:eastAsia"), "仿宋")

    add_paragraph(doc, MEETING["title"], font="黑体", size=18, bold=True,
                  align=WD_ALIGN_PARAGRAPH.CENTER, space_after=12, indent_first=False)
    add_paragraph(doc, "（本纪要依据会议转写稿整理，共三部分：决议事项、待办事项、风险与关注）",
                  font="楷体", size=10.5, align=WD_ALIGN_PARAGRAPH.CENTER,
                  space_after=10, indent_first=False)

    # 会议基本信息
    info = doc.add_table(rows=0, cols=2)
    info.style = "Table Grid"
    info.alignment = WD_TABLE_ALIGNMENT.CENTER
    for k, v in MEETING["meta"]:
        row = info.add_row().cells
        for cell, text, bold in ((row[0], k, True), (row[1], v, False)):
            cell.text = ""
            p = cell.paragraphs[0]
            set_cjk(p.add_run(text), font="仿宋", size=11, bold=bold)
    fix_layout(info, [72, 368])

    # 一、决议事项
    add_heading(doc, "一、决议事项（共 %d 条）" % len(DECISIONS))
    for i, (name, detail, ts) in enumerate(DECISIONS, 1):
        add_paragraph(doc, "%d. %s：%s（转写稿 %s）" % (i, name, detail, ts),
                      indent_first=False, space_after=4)

    # 二、待办事项表格
    add_heading(doc, "二、待办事项表格（共 %d 条）" % len(TODOS))
    table = doc.add_table(rows=1, cols=len(TODO_COLUMNS))
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    for cell, name in zip(table.rows[0].cells, TODO_COLUMNS):
        cell.text = ""
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        set_cjk(p.add_run(name), font="黑体", size=10.5, bold=True)
    for idx, item in enumerate(TODOS, 1):
        cells = table.add_row().cells
        values = [str(idx)] + list(item)
        for j, (cell, value) in enumerate(zip(cells, values)):
            cell.text = ""
            p = cell.paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            if j != 1:
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            set_cjk(p.add_run(value), font="仿宋", size=10)
    widths = [32, 176, 56, 58, 50, 76]
    fix_layout(table, widths)
    # 表头跨页重复
    tr_pr = table.rows[0]._tr.get_or_add_trPr()
    tr_pr.append(tr_pr.makeelement(qn("w:tblHeader"), {}))

    # 三、风险与关注
    add_heading(doc, "三、风险与关注（共 %d 条）" % len(RISKS))
    for i, (name, detail, owner, level, ts) in enumerate(RISKS, 1):
        add_paragraph(doc, "%d. 【%s】%s —— %s（涉及：%s；转写稿 %s）"
                      % (i, level, name, detail, owner, ts),
                      indent_first=False, space_after=4)

    doc.save(DOCX_PATH)


def build_json():
    summary = {
        "决议事项": len(DECISIONS),
        "待办事项": len(TODOS),
        "风险与关注": len(RISKS),
        "合计": len(DECISIONS) + len(TODOS) + len(RISKS),
    }
    with open(JSON_PATH, "w", encoding="utf-8") as f:
        json.dump(summary, f, ensure_ascii=False, indent=2)
        f.write("\n")
    return summary


if __name__ == "__main__":
    os.makedirs(OUT_DIR, exist_ok=True)
    build_docx()
    summary = build_json()
    print("转写稿：", TRANSCRIPT, os.path.getsize(TRANSCRIPT), "bytes")
    print("纪要：", DOCX_PATH, os.path.getsize(DOCX_PATH), "bytes")
    print("统计：", json.dumps(summary, ensure_ascii=False))
