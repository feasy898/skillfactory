# -*- coding: utf-8 -*-
"""Build 纪要.docx + summary.json from xiangqi_fashion_1009.txt."""
import json
import os

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt

BASE = os.path.dirname(os.path.abspath(__file__))
OUT = os.path.join(BASE, "out")

# ---------------------------------------------------------------- content ---
MEETING = {
    "title": "秋装首秀复盘会 会议纪要",
    "meta": [
        ("会议时间", "10:00–10:10"),
        ("会议主题", "秋装首秀复盘"),
        ("参会角色", "主播统筹、运营费晔、客服总监宦悦、物控祁妙、仓运蔚少霆"),
        ("下次会议", "下周三同一时间，复盘双11预案"),
    ],
}

DECISIONS = [
    ("报销同意书审批流", "报销同意书的审批流程统一改走新版OA系统，不再沿用原审批路径。", "主播统筹", "10:01"),
    ("下周主播排班", "主播冯议定的下周排班方案确认通过，按此执行。", "主播统筹", "10:04"),
    ("下次复盘安排", "本次复盘结束，下次于下周三同一时间召开，议题为双11预案。", "主播统筹", "10:10"),
]

# 序号 / 待办事项 / 责任人 / 截止时间 / 状态
TODO_ROWS = [
    ("1", "将本批退货的延期赔偿口径准确传达至各销区", "运营费晔", "未指定", "待办"),
    ("2", "跟进面料回修工单，责任落实到车间", "物控祁妙", "未指定", "待办"),
    ("3", "重新配置直播间优惠券的风控规则", "运营费晔", "本月底前", "待办"),
    ("4", "在群内接龙确认应急小组名单", "客服总监宦悦", "未指定", "待办"),
]

RISKS = [
    ("尺码相关差评集中", "客服总监宦悦", "差评中关于尺码的结论性吐槽高度集中，需关注尺码问题对售后的持续影响。", "10:03"),
    ("尾货清仓定价未定", "仓运蔚少霆", "尾货清仓定价尚未确定，会上未形成结论，需另行听取各方意见后定档。", "10:06"),
    ("快递履约时效延期", "仓运蔚少霆", "快递合作伙伴的履约时效存在延期风险，可能影响发货与用户收货体验。", "10:09"),
]


# ------------------------------------------------------------------ styles ---
def set_cn_font(run, size, bold=False, name="微软雅黑"):
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.name = name
    run._element.rPr.rFonts.set(qn("w:eastAsia"), name)


def add_para(doc, text, size=10.5, bold=False, space_after=4, indent=False):
    p = doc.add_paragraph()
    if indent:
        p.paragraph_format.left_indent = Pt(18)
    p.paragraph_format.space_after = Pt(space_after)
    set_cn_font(p.add_run(text), size, bold)
    return p


def add_heading(doc, text):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(9)
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.keep_with_next = True
    set_cn_font(p.add_run(text), 13, bold=True)
    return p


# ------------------------------------------------------------------- build ---
doc = Document()
style = doc.styles["Normal"]
style.font.name = "微软雅黑"
style.font.size = Pt(10.5)
style.paragraph_format.space_after = Pt(2)
style.element.rPr.rFonts.set(qn("w:eastAsia"), "微软雅黑")

for section in doc.sections:
    section.top_margin = Cm(2.0)
    section.bottom_margin = Cm(2.0)
    section.left_margin = Cm(2.2)
    section.right_margin = Cm(2.2)

title = doc.add_paragraph()
title.alignment = WD_ALIGN_PARAGRAPH.CENTER
title.paragraph_format.space_after = Pt(8)
set_cn_font(title.add_run(MEETING["title"]), 16, bold=True)

for k, v in MEETING["meta"]:
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(1)
    set_cn_font(p.add_run("%s：" % k), 10.5, bold=True)
    set_cn_font(p.add_run(v), 10.5)

# 一、决议事项
add_heading(doc, "一、决议事项")
for idx, (item, detail, owner, ts) in enumerate(DECISIONS, 1):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(1)
    set_cn_font(p.add_run("%d. %s" % (idx, item)), 10.5, bold=True)
    add_para(doc, detail, indent=True, space_after=1)
    add_para(doc, "—— %s（%s）" % (owner, ts), indent=True, space_after=4)

# 二、待办事项
add_heading(doc, "二、待办事项")
table = doc.add_table(rows=1, cols=5)
table.style = "Table Grid"
table.alignment = WD_TABLE_ALIGNMENT.CENTER
table.autofit = False
COL_WIDTHS = (Cm(1.3), Cm(7.6), Cm(2.6), Cm(2.3), Cm(1.6))
headers = ["序号", "待办事项", "责任人", "截止时间", "状态"]
for cell, text in zip(table.rows[0].cells, headers):
    cell.text = ""
    p = cell.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_cn_font(p.add_run(text), 10, bold=True)
for row in TODO_ROWS:
    cells = table.add_row().cells
    for i, (cell, text) in enumerate(zip(cells, row)):
        cell.text = ""
        p = cell.paragraphs[0]
        if i in (0, 2, 3, 4):
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        set_cn_font(p.add_run(text), 10)
for row in table.rows:
    for cell, width in zip(row.cells, COL_WIDTHS):
        cell.width = width
for col, width in zip(table.columns, COL_WIDTHS):
    col.width = width

# LibreOffice/Word only honour explicit column widths under a fixed layout.
_tblPr = table._tbl.tblPr
_layout = OxmlElement("w:tblLayout")
_layout.set(qn("w:type"), "fixed")
_tblPr.append(_layout)
_tblW = OxmlElement("w:tblW")
_tblW.set(qn("w:w"), str(int(sum(w.cm for w in COL_WIDTHS) * 567)))
_tblW.set(qn("w:type"), "dxa")
_tblPr.append(_tblW)

# 三、风险与关注
add_heading(doc, "三、风险与关注")
for idx, (item, owner, detail, ts) in enumerate(RISKS, 1):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(1)
    set_cn_font(p.add_run("%d. %s" % (idx, item)), 10.5, bold=True)
    add_para(doc, detail, indent=True, space_after=1)
    add_para(doc, "—— 提示人：%s（%s）" % (owner, ts), indent=True, space_after=4)

os.makedirs(OUT, exist_ok=True)
docx_path = os.path.join(OUT, "纪要.docx")
doc.save(docx_path)

summary = {
    "决议事项": len(DECISIONS),
    "待办事项": len(TODO_ROWS),
    "风险与关注": len(RISKS),
    "合计": len(DECISIONS) + len(TODO_ROWS) + len(RISKS),
}
with open(os.path.join(OUT, "summary.json"), "w", encoding="utf-8") as f:
    json.dump(summary, f, ensure_ascii=False, indent=2)
    f.write("\n")

print(json.dumps(summary, ensure_ascii=False))
print("saved:", docx_path)