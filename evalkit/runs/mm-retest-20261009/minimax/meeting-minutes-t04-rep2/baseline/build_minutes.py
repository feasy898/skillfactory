#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""由 xiangqi_fashion_1009.txt 生成三段式会议纪要：out/纪要.docx + out/summary.json"""
import json
import os

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Pt, RGBColor, Cm

BASE = os.path.dirname(os.path.abspath(__file__))
OUT = os.path.join(BASE, "out")
SRC = os.path.join(BASE, "xiangqi_fashion_1009.txt")

CN_FONT = "微软雅黑"

# ---------- 内容结构（逐条对应转写稿时间戳） ----------

DECISIONS = [
    ("报销同意书审批流改走新版OA", "自即日起，报销同意书审批统一在新版OA流程中流转。", "主播统筹", "10:01"),
    ("确认下周主播排班", "下周主播排班以冯议确定的方案为准，不再调整。", "主播统筹／冯议", "10:04"),
    ("下次复盘安排", "本次会议结束，下周三同一时间复盘双11预案。", "主播统筹", "10:10"),
]

TODO_HEADER = ("序号", "待办事项", "责任人", "时限", "来源")
TODOS = [
    ("这批退货的延期赔偿口径需传达至各销区，确保口径一致、传达到位。", "运营费晔（各销区配合）", "尽快", "10:02"),
    ("跟进面料回修工单，责任落实到车间，闭环处理。", "物控祁妙（车间）", "本周内", "10:05"),
    ("尾货清仓定价方案尚不确定，需征求意见后定案。", "仓运蔚少霆", "待定（会上未收敛）", "10:06"),
    ("重新配置直播间优惠券的风控规则。", "运营费晔", "月底前完成", "10:07"),
    ("应急小组名单在群内接龙汇总。", "客服总监宦悦", "尽快", "10:08"),
]

RISKS = [
    ("尺码相关差评高度集中", "差评中关于尺码的结论性吐槽很集中，反映尺码版型与用户预期存在系统性偏差。", "客服总监宦悦", "10:03"),
    ("快递合作伙伴履约时效延期", "快递合作伙伴的履约时效存在延期风险，可能影响发货与售后体验。", "仓运蔚少霆", "10:09"),
]

TOTAL = len(DECISIONS) + len(TODOS) + len(RISKS)


# ---------- 样式助手 ----------

def set_cn(run, name=CN_FONT, size=None, bold=None, color=None):
    run.font.name = name
    run._element.rPr.rFonts.set(qn("w:eastAsia"), name)
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.font.bold = bold
    if color is not None:
        run.font.color.rgb = color
    return run


def style_doc(doc):
    normal = doc.styles["Normal"]
    normal.font.name = CN_FONT
    normal.font.size = Pt(10.5)
    normal.element.rPr.rFonts.set(qn("w:eastAsia"), CN_FONT)
    pf = normal.paragraph_format
    pf.space_after = Pt(4)
    pf.line_spacing = 1.4


def heading(doc, text):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(14)
    p.paragraph_format.space_after = Pt(6)
    p.paragraph_format.keep_with_next = True   # 标题不单独留在页尾
    set_cn(p.add_run(text), size=14, bold=True, color=RGBColor(0x1F, 0x38, 0x64))
    return p


def body(doc, text, indent=True):
    p = doc.add_paragraph()
    if indent:
        p.paragraph_format.left_indent = Cm(0.6)
    set_cn(p.add_run(text), size=10.5)
    return p


def shade(cell, hex_color):
    from docx.oxml import OxmlElement
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:fill"), hex_color)
    tcPr.append(shd)


def fix_widths(table, widths):
    """关闭自动调整并写入固定列宽，否则 Word/LibreOffice 会忽略单元格宽度。"""
    from docx.oxml import OxmlElement
    from docx.shared import Cm as _Cm
    table.autofit = False
    tblPr = table._tbl.tblPr
    layout = OxmlElement("w:tblLayout")
    layout.set(qn("w:type"), "fixed")
    tblPr.append(layout)
    for row in table.rows:
        for c, w in enumerate(widths):
            row.cells[c].width = w
    for c, w in enumerate(widths):
        table.columns[c].width = w
    return _Cm


def main():
    os.makedirs(OUT, exist_ok=True)
    doc = Document()
    style_doc(doc)
    for s in doc.sections:
        s.top_margin = s.bottom_margin = Cm(2.2)
        s.left_margin = s.right_margin = Cm(2.4)

    # 标题区
    t = doc.add_paragraph()
    t.alignment = WD_ALIGN_PARAGRAPH.CENTER
    t.paragraph_format.space_after = Pt(2)
    set_cn(t.add_run("秋装首秀复盘会 会议纪要"), size=20, bold=True,
           color=RGBColor(0x1F, 0x38, 0x64))
    st = doc.add_paragraph()
    st.alignment = WD_ALIGN_PARAGRAPH.CENTER
    st.paragraph_format.space_after = Pt(10)
    set_cn(st.add_run("2026-10-09 10:00–10:10　|　来源：xiangqi_fashion_1009.txt"),
           size=9, color=RGBColor(0x60, 0x60, 0x60))

    # 会议信息
    info = doc.add_table(rows=2, cols=4)
    info.alignment = WD_TABLE_ALIGNMENT.CENTER
    info.style = "Table Grid"
    info_vals = [
        ("会议主题", "秋装首秀复盘", "参会角色", "主播统筹、运营、客服、物控、仓运"),
        ("条数统计", "决议 %d ／ 待办 %d ／ 风险 %d ／ 合计 %d"
         % (len(DECISIONS), len(TODOS), len(RISKS), TOTAL), "会议时长", "约 10 分钟"),
    ]
    for r, vals in enumerate(info_vals):
        for c, v in enumerate(vals):
            cell = info.cell(r, c)
            cell.text = ""
            p = cell.paragraphs[0]
            is_label = c % 2 == 0
            set_cn(p.add_run(v), size=9.5, bold=is_label)
            shade(cell, "EEF2F8" if is_label else "FFFFFF")
    for row in info.rows:
        for c, w in enumerate((Cm(2.4), Cm(6.6), Cm(2.4), Cm(5.6))):
            row.cells[c].width = w
    fix_widths(info, [Cm(2.4), Cm(6.6), Cm(2.4), Cm(5.6)])

    # 一、决议事项
    heading(doc, "一、决议事项")
    for i, (title, detail, owner, ts) in enumerate(DECISIONS, 1):
        p = doc.add_paragraph()
        p.paragraph_format.left_indent = Cm(0.6)
        p.paragraph_format.space_after = Pt(3)
        set_cn(p.add_run("%d. %s" % (i, title)), size=10.5, bold=True)
        b = doc.add_paragraph()
        b.paragraph_format.left_indent = Cm(1.4)
        b.paragraph_format.space_after = Pt(6)
        set_cn(b.add_run("内容：%s　|　落实：%s　|　依据：转写稿 [%s]" % (detail, owner, ts)),
               size=9.5, color=RGBColor(0x44, 0x44, 0x44))

    # 二、待办事项
    heading(doc, "二、待办事项")
    tb = doc.add_table(rows=1, cols=5)
    tb.style = "Table Grid"
    tb.alignment = WD_TABLE_ALIGNMENT.CENTER
    for c, h in enumerate(TODO_HEADER):
        cell = tb.rows[0].cells[c]
        cell.text = ""
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        set_cn(p.add_run(h), size=10, bold=True, color=RGBColor(0xFF, 0xFF, 0xFF))
        shade(cell, "1F3864")
    for i, (item, owner, due, ts) in enumerate(TODOS, 1):
        cells = tb.add_row().cells
        vals = [str(i), item, owner, due, ts]
        for c, v in enumerate(vals):
            cells[c].text = ""
            p = cells[c].paragraphs[0]
            if c in (0, 3, 4):
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            set_cn(p.add_run(v), size=9.5)
        if i % 2 == 0:
            for cell in cells:
                shade(cell, "F5F7FA")
    from docx.oxml import OxmlElement
    hdr = OxmlElement("w:tblHeader")           # 跨页时重复表头
    tb.rows[0]._tr.get_or_add_trPr().append(hdr)
    for row in tb.rows:                          # 禁止单行跨页断开
        cant = OxmlElement("w:cantSplit")
        row._tr.get_or_add_trPr().append(cant)

    widths = [Cm(1.2), Cm(7.4), Cm(3.2), Cm(2.8), Cm(1.6)]
    for row in tb.rows:
        for c, w in enumerate(widths):
            row.cells[c].width = w
    fix_widths(tb, widths)

    # 三、风险与关注
    heading(doc, "三、风险与关注")
    for i, (title, detail, owner, ts) in enumerate(RISKS, 1):
        p = doc.add_paragraph()
        p.paragraph_format.left_indent = Cm(0.6)
        p.paragraph_format.space_after = Pt(3)
        set_cn(p.add_run("%d. %s" % (i, title)), size=10.5, bold=True,
               color=RGBColor(0xA6, 0x1B, 0x1B))
        b = doc.add_paragraph()
        b.paragraph_format.left_indent = Cm(1.4)
        b.paragraph_format.space_after = Pt(6)
        set_cn(b.add_run("%s　|　提出方：%s　|　依据：转写稿 [%s]" % (detail, owner, ts)),
               size=9.5, color=RGBColor(0x44, 0x44, 0x44))

    # 说明
    note = doc.add_paragraph()
    note.paragraph_format.space_before = Pt(12)
    set_cn(note.add_run(
        "说明：转写稿共 11 条发言，[10:00] 为开场寒暄（未形成纪要条目），"
        "其余 10 条全部归入上述三段，无遗漏、无重复计数。合计 %d 条。" % TOTAL),
        size=9, color=RGBColor(0x70, 0x70, 0x70))

    docx_path = os.path.join(OUT, "纪要.docx")
    doc.save(docx_path)

    summary = {
        "决议事项": len(DECISIONS),
        "待办事项": len(TODOS),
        "风险与关注": len(RISKS),
        "合计": TOTAL,
    }
    json_path = os.path.join(OUT, "summary.json")
    with open(json_path, "w", encoding="utf-8") as f:
        json.dump(summary, f, ensure_ascii=False, indent=2)
        f.write("\n")

    print("written:", docx_path)
    print("written:", json_path)
    print("counts:", summary)


if __name__ == "__main__":
    main()
