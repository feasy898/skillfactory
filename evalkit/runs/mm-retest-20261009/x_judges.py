#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""runs/mm-retest-20261009/x_judges.py — mm 重测裁判面板驱动（盲态、异族 ≥2、位置轮换、§3.4 抽检、§9-3 arith 首轮即生效）。

移植自 office-retest-20261009/x_judges.py + x_judge_t08_rejudge.py（arith/重发机制全题内建），按本轮登记（round-plan.json）适配：
  面板：被测 glm（t01/t03/t05/t07）→ [step-router-v1, Minimax-M3.1-Flash-Preview]；
        被测 minimax（t02/t04/t06/t08）→ [step-router-v1, GLM-5.3]（≥2 族互裁、与被测族交集为空）。
  step=模型 API 直调（无 assistant-prefill）；glm/minimax=headless M4/M2。
  盲态：产物匿名化（P1..P6 洗牌）+ docx→事实预抽取 + 路径/挂载痕迹漂白；映射仅存 records/judge-staging/。
  位置轮换：槽序按任务序奇偶互换。
判分协议（judge_prompt_version=S8-JUDGE-v2-MMRETEST-1，随轮落盘）：全 8 题=加权维度判分，
  judge-notes 首批全部嵌入 brief（装置教训）；判据面=rubric.md+judge-notes.md（不进臂）。
§9-3 judge_arith_check 首轮即生效：harness 逐产物按 rubric 权重机算 total（Decimal HALF_UP 0.1；
  本轮 8 题 rubric 无封顶条款）；判官 total 与机算不一致 → 判词无效 → 同裁判新会话重发 1 次 →
  仍不一致 fail-closed unjudgeable（禁止人工改数）。
D2 纪律：判词按 slot 先写后处理（raw 先落 verdict-slot<k>.json 再解析）；判词归属记实际会话。
§3.4 + §9-4：每题抽 1 份产物同裁判匿名复评（Q1）；|Δ|>2 → 换面板内另一异族裁判重评全任务
  （比较对象=换上裁判对被抽产物分数 vs 失格裁判首评）；仍>2 → unjudgeable（blind=false，§9-5 登记口径）。
用法：python x_judges.py [--task <task_id>]
退出码：0 全任务判分落盘 / 1 有任务 unjudgeable 或装置失败（如实，不放宽）。
"""
import hashlib
import json
import os
import random
import re
import shutil
import sys
import zlib
from decimal import Decimal, ROUND_HALF_UP

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
import x_judge_lib as J  # noqa: E402

ROOT = J.ROOT
RUN = J.RUN
PANEL = {"glm": ["step-router-v1", "Minimax-M3.1-Flash-Preview"],
         "minimax": ["step-router-v1", "GLM-5.3"]}
JUDGE_OF_MODEL = {"step-router-v1": "step", "Minimax-M3.1-Flash-Preview": "minimax",
                  "GLM-5.3": "glm"}

TASK_ORDER = ["t0%d" % i for i in range(1, 9)]
MERGED = os.path.join(ROOT, "planning", "taskset-redesign", "merged", "meeting-minutes-skill")
FAM_OF_TASK = {t: ("glm" if i % 2 == 1 else "minimax") for i, t in enumerate(TASK_ORDER, 1)}

SCORING_RULE = """评分规则（S8-JUDGE-v2-MMRETEST-1 / 加权维度协议）：
- rubric 表每行为一个维度（维度名/权重/满分档/中档/零分档）。对每维度独立定档并赋分：
  满分档=1.0；零分档=0.0；中档=该维度分（取 rubric 中明示的分值，如 0.4/0.5；rubric 未明示数值时取 0.5）。
- 定档必须附实测证据（引用暂存文件中的具体事实/原文/数字，逐条 ≤150 字）；无证据的档位无效。
- 产物加权分 total = round(10 × Σ(维度分 × 权重) ÷ Σ权重, 1)，区间 [0,10]。这是唯一总分，不再另设项分平均。
- 你不知道也不得推测产物的来源臂、模型或顺序；只按 rubric 与 judge-notes 与证据判分；实现路径与参照不同不扣分（端态原则），除非 rubric 显式考格式/结构。
- verdict.json 结构（紧凑，必须遵守）：{"scores": {"P1": {"total": x.x, "dims": {"<维度名>": 分数, ...}, "evidence": "≤150 字逐维证据摘要"}}, ...}（dims 键名须与 rubric 维度名逐字一致；除 scores 外不要其他顶层键）。"""

DIM_KEYS = ("正确性", "完整性", "可验证性")
_STAGING_ROOTS = {}
_SLOT_SEQ = {"n": 0}


def sha256_file(p):
    return hashlib.sha256(open(p, "rb").read()).hexdigest()


def redact_text(s, arm_root):
    s = s.replace(arm_root, "<ARMROOT>")
    s = s.replace(ROOT, "<REPOROOT>")
    s = re.sub(r"\.zcode/skills/[A-Za-z0-9_\-]+", "<SKILL_MOUNT>", s)
    s = re.sub(r"\.zcode\\\\skills\\\\[A-Za-z0-9_\-]+", "<SKILL_MOUNT>", s)
    for key in ("arm", "engine", "run_id", "runId", "side", "arm_label", "session_id",
                "sessionId", "skill_mount", "asset"):
        s = re.sub(r'"%s"\s*:\s*"[^"]*"' % key, '"%s": "<redacted>"' % key, s)
        s = re.sub(r"'%s'\s*:\s*'[^']*'" % key, "'%s': '<redacted>'" % key, s)
    return s


def docx_facts(path):
    """python-docx 确定性抽取：段落文本+对齐+缩进+字体/字号/行距 + 表格文本。"""
    from docx import Document
    from docx.oxml.ns import qn
    doc = Document(path)
    lines = []
    for i, p in enumerate(doc.paragraphs):
        if not p.text.strip() and p._p.xpath(".//w:drawing") == []:
            continue
        pf = p.paragraph_format
        fli_chars = None
        pPr = p._p.pPr
        if pPr is not None:
            ind = pPr.find(qn("w:ind"))
            if ind is not None:
                fli_chars = ind.get(qn("w:firstLineChars"))
        fonts = []
        for r in p.runs[:6]:
            sz = r.font.size.pt if r.font.size else None
            east = None
            rPr = r._r.rPr
            if rPr is not None:
                rf = rPr.find(qn("w:rFonts"))
                if rf is not None:
                    east = rf.get(qn("w:eastAsia"))
            fonts.append("%s/%s/%s" % (r.text[:14], east or r.font.name, sz))
        lines.append("P%03d align=%s firstLineChars=%s firstLineIndent=%s spacing=%s runs=[%s] text=%s"
                     % (i, p.alignment, fli_chars, pf.first_line_indent, pf.line_spacing,
                        " | ".join(fonts), p.text))
    for ti, tb in enumerate(doc.tables, 1):
        lines.append("TABLE%d %d行×%d列:" % (ti, len(tb.rows), len(tb.columns)))
        for row in tb.rows:
            lines.append("  " + " ┃ ".join(c.text.replace("\n", "⏎") for c in row.cells))
    return "\n".join(lines)


def stage_products(task, unit_dirs, staging, seed):
    """收集产物→漂白→洗牌匿名化。返回 (staged, mapping, meta)。"""
    rng = random.Random(seed)
    items = []
    for rep, ad in unit_dirs:
        files = {}
        out = os.path.join(ad, "out")
        if os.path.isdir(out):
            for dp, _dn, fns in os.walk(out):
                for fn in fns:
                    src = os.path.join(dp, fn)
                    files[os.path.relpath(src, out)] = src
        rl = os.path.join(ad, "run_log.json")
        if os.path.isfile(rl):
            files["run_log.json"] = rl
        items.append({"rep": rep, "arm": os.path.basename(os.path.normpath(ad)),
                      "dir": ad, "files": files})
    rng.shuffle(items)
    os.makedirs(staging, exist_ok=True)
    mapping, staged, meta = {}, [], []
    for i, it in enumerate(items, 1):
        pid = "P%d" % i
        mapping[pid] = {"rep": it["rep"], "arm": it["arm"]}
        pdir = os.path.join(staging, pid)
        os.makedirs(pdir, exist_ok=True)
        names = []
        for rel, src in sorted(it["files"].items()):
            ext = os.path.splitext(rel)[1].lower()
            base = rel.replace("/", "__")
            dst = os.path.join(pdir, base)
            if ext == ".docx":
                facts = docx_facts(src)
                with open(dst + ".extract.txt", "w", encoding="utf-8", newline="\n") as f:
                    f.write("EXTRACT OF %s (python-docx 确定性抽取，段落含对齐/缩进/字体/字号/行距事实)\n%s\n"
                            % (base, facts))
                names.append(base + ".extract.txt")
            elif ext in (".json", ".md", ".txt", ".log"):
                raw = open(src, encoding="utf-8", errors="replace").read()
                with open(dst, "w", encoding="utf-8", newline="\n") as f:
                    f.write(redact_text(raw, it["dir"]))
                names.append(base)
            else:
                shutil.copyfile(src, dst)
                names.append(base)
        staged.append({"pid": pid, "files": names})
        meta.append({"pid": pid, "n_files": len(names)})
    return staged, mapping, meta


def rubric_weights(task):
    """从 rubric.md 冻结原文机读权重表：| 维度名 | 权重 | ...。返回 ({dim_key: (name, Decimal w)}, Σw)。"""
    text = open(os.path.join(MERGED, task, "rubric.md"), encoding="utf-8").read()
    weights = {}
    for m in re.finditer(r"^\|\s*([^|]+?)\s*\|\s*(0\.\d+)\s*\|", text, re.M):
        name = m.group(1).strip()
        w = Decimal(m.group(2))
        key = next((k for k in DIM_KEYS if k in name), None)
        if key is None:
            raise RuntimeError("%s rubric 权重表出现未知维度: %s" % (task, name))
        weights[key] = (name, w)
    if len(weights) != 3:
        raise RuntimeError("%s rubric 权重表维度数≠3: %s" % (task, sorted(weights)))
    sum_w = sum(w for _n, w in weights.values())
    if sum_w != Decimal("1.0"):
        raise RuntimeError("%s rubric Σ权重=%s ≠ 1.0" % (task, sum_w))
    return weights, sum_w


def formula_total(dims, weights, sum_w):
    """§9-3：total = round_half_up(10 × Σ(维度分×权重)/Σ权重, 1)（Decimal 精确）。
    本轮 8 题 rubric 无封顶条款（开工扫描实证）→ 无 min(·,cap) 项（§9-1 上限语义：无封顶即无截断）。"""
    acc = Decimal(0)
    for k, v in dims.items():
        hit = next((kk for kk in weights if kk in str(k)), None)
        if hit is None:
            return None
        acc += Decimal(str(v)) * weights[hit][1]
    return (Decimal(10) * acc / sum_w).quantize(Decimal("0.1"), rounding=ROUND_HALF_UP)


def arith_check(verdict, weights, sum_w, pids):
    """逐产物机算对比。返回 (rows, invalid)。任一产物无法机算相符 → 判词整体无效（§9-3）。"""
    rows, invalid = [], False
    scores = (verdict or {}).get("scores") or {}
    for pid in pids:
        sv = scores.get(pid) or {}
        jt, dims = sv.get("total"), sv.get("dims") or {}
        f = formula_total(dims, weights, sum_w) if len(dims) == 3 else None
        match = (f is not None and jt is not None and Decimal(str(jt)) == f)
        if not match:
            invalid = True
        rows.append({"pid": pid, "dims": dims, "judge_total": jt,
                     "formula_total": (float(f) if f is not None else None), "match": match})
    return rows, invalid


def build_brief(task, staged):
    tj = json.load(open(os.path.join(MERGED, task, "task.json"), encoding="utf-8"))
    notes = open(os.path.join(MERGED, task, "judge-notes.md"), encoding="utf-8").read().rstrip()
    rubric = open(os.path.join(MERGED, task, "rubric.md"), encoding="utf-8").read().rstrip()
    lines = ["# mm 重测盲评判分任务（%s，通过线 %s）" % (task, tj.get("pass_line")),
             "", "你是盲评裁判。以下为本任务的匿名化产物（P1..Pn，顺序随机，来源对你保密）。",
             "", "## rubric（冻结原文，加权维度判分）", "", rubric,
             "", "## judge-notes（判官专用：判挂点与期望值，判分必照此执行）", "", notes,
             "", SCORING_RULE,
             "", "（题目指令原文备查，判分只按 rubric 与 judge-notes）：", "```", tj["instruction"], "```",
             "", "## 产物清单"]
    for s in staged:
        lines.append("- %s: %s" % (s["pid"], ", ".join(s["files"])))
    lines += ["", "请阅读同目录各 P* 子目录内的文件后，把 verdict.json 写入当前工作目录（仅输出该文件）。"]
    return "\n".join(lines)


def parse_verdict_json(text):
    """提取含 "scores" 的最后一个平衡 JSON 对象（容忍围栏/推理散文）。"""
    best = None
    for m in re.finditer(r'"scores"', text):
        i = text.rfind("{", 0, m.start())
        if i < 0:
            continue
        depth = 0
        for j in range(i, len(text)):
            if text[j] == "{":
                depth += 1
            elif text[j] == "}":
                depth -= 1
                if depth == 0:
                    cand = text[i:j + 1]
                    try:
                        obj = json.loads(cand)
                        if isinstance(obj, dict) and isinstance(obj.get("scores"), dict):
                            best = obj
                    except Exception:
                        pass
                    break
    return best


def judge_call(judge, workdir, brief, staged, task_label, phase="main"):
    """单裁判单次调用。D2 纪律：判词按 slot 先写后处理、判词归属记实际会话。
    step=HTTP（产物内容内联）；glm/minimax=headless CLI。返回 (rec, verdict, slot_rec)。"""
    import time
    chan = JUDGE_OF_MODEL[judge]
    _SLOT_SEQ["n"] += 1
    slot_no = _SLOT_SEQ["n"]
    slot = J.budget_judge({"judge": judge, "channel": chan, "task": task_label, "phase": phase,
                           "judge_prompt_version": J.JUDGE_PROMPT_VERSION,
                           "ts_intent": time.strftime("%H:%M:%S"), "status": "pending"})
    v = None
    try:
        if chan == "step":
            parts = ["【输出纪律（最高优先级）】你是评分器。禁止展示任何思考、草稿、advisor 对话或分步分析；你的回复必须直接以 { 开头、以 } 结尾，内容仅为最终 verdict JSON。", "", brief, "", "# 产物内容（暂存文件全文）"]
            for s in staged:
                src_base = _STAGING_ROOTS[task_label]
                for fn in s["files"]:
                    cand = [os.path.join(src_base, s["pid"], fn), os.path.join(src_base, fn)]
                    p = next((c for c in cand if os.path.isfile(c)), None)
                    if p is None:
                        continue
                    parts.append("=== %s/%s ===\n%s" % (s["pid"], fn,
                                 open(p, encoding="utf-8", errors="replace").read()[:24000]))
            r = J.step_chat([{"role": "user", "content": "\n".join(
                parts + ["", '现在直接输出最终评分 JSON：直接以 { 开头、以 } 结尾，最外层恰有一个 "scores" 键；每个产品一个键，键含 "total" 与 "dims" 与 "evidence"；除该 JSON 外不得输出任何文字，不得重复输出 JSON。'])}],
                max_tokens=32000)
            # 先写后处理：raw 判词落 slot 文件（归属=本次 HTTP 会话 id）
            raw_path = os.path.join(workdir, "verdict-slot%d.raw.txt" % slot_no)
            with open(raw_path, "w", encoding="utf-8", newline="\n") as f:
                f.write(r["text"])
            v = parse_verdict_json(r["text"])
            rec = {"exit_code": 0 if v else 1, "wall_ms": r["wall_ms"], "session_id": r["id"],
                   "model_id_observed": r["model"], "channel": "http-api",
                   "input_tokens": r["prompt_tokens"], "output_tokens": r["completion_tokens"]}
            with open(os.path.join(workdir, "response.txt"), "w", encoding="utf-8",
                      newline="\n") as f:
                f.write(r["text"])
        else:
            with open(os.path.join(workdir, "brief.md"), "w", encoding="utf-8",
                      newline="\n") as f:
                f.write(brief + "\n")
            for s in staged:
                src_pdir = os.path.join(_STAGING_ROOTS[task_label], s["pid"])
                dst_pdir = os.path.join(workdir, s["pid"])
                if os.path.isdir(src_pdir):
                    shutil.copytree(src_pdir, dst_pdir, dirs_exist_ok=True)
            rec = J.cli_judge(workdir, chan)
            vp = os.path.join(workdir, "verdict.json")
            if os.path.isfile(vp):
                # 先写后处理：模型落盘的 verdict.json 先拷 slot 文件（归属=实际会话）再解析
                slot_path = os.path.join(workdir, "verdict-slot%d.json" % slot_no)
                shutil.copyfile(vp, slot_path)
                try:
                    v = json.load(open(slot_path, encoding="utf-8"))
                except Exception:
                    v = None
                if not (isinstance(v, dict) and isinstance(v.get("scores"), dict)):
                    v = parse_verdict_json(open(slot_path, encoding="utf-8",
                                                errors="replace").read())
            with open(os.path.join(workdir, "transcript_stdout.txt"), "w", encoding="utf-8",
                      newline="\n") as f:
                f.write(rec["_raw_stdout"])
        J.budget_judge_update(slot, status="done" if v else "no-verdict",
                              exit_code=rec["exit_code"], wall_ms=rec["wall_ms"],
                              session_id=rec.get("session_id"),
                              tokens={"input": rec.get("input_tokens"),
                                      "output": rec.get("output_tokens")},
                              model_observed=rec.get("model_id_observed"),
                              verdict_slot="verdict-slot%d" % slot_no)
        slot_rec = {"slot": slot, "slot_no": slot_no, "session_id": rec.get("session_id"),
                    "model_id_observed": rec.get("model_id_observed"),
                    "exit_code": rec["exit_code"], "wall_ms": rec["wall_ms"],
                    "input_tokens": rec.get("input_tokens"),
                    "output_tokens": rec.get("output_tokens")}
        return rec, v, slot_rec
    except Exception as exc:
        J.budget_judge_update(slot, status="failed", exit_code=1,
                              error=type(exc).__name__ + ":" + str(exc)[:300])
        raise


def judge_validated(judge, workroot, brief, staged, task, weights, sum_w, pids, phase_prefix,
                    log):
    """主判 + arith 不一致同裁判新会话重发 1 次（§9-3）+ 通道故障重试 ≤1（登记 retry_policy）；
    两次均未产出有效判词 → fail-closed（arith 与通道故障均如实留痕，不放宽）。"""
    attempts = []
    v_final = None
    for attempt in (1, 2):
        label = "%s%s" % (phase_prefix, "" if attempt == 1 else "-resend")
        wdir = os.path.join(workroot, "%s-%s-a%d" % (judge.replace("/", "_"), label, attempt))
        os.makedirs(wdir, exist_ok=True)
        try:
            rec, v, slot_rec = judge_call(judge, wdir, brief, staged, label,
                                          phase="resend" if attempt == 2 else "main")
        except Exception as exc:  # 通道故障：如实留痕，重试 ≤1
            attempts.append({"attempt": attempt, "judge": judge, "phase": label,
                             "channel": JUDGE_OF_MODEL[judge],
                             "channel_failure": type(exc).__name__ + ":" + str(exc)[:300],
                             "arith_invalid": True, "had_verdict": False})
            log("  judge=%s attempt=%d channel_failure=%s" % (judge, attempt, type(exc).__name__))
            continue
        rows, invalid = arith_check(v, weights, sum_w, pids)
        attempts.append(dict({"attempt": attempt, "judge": judge, "phase": label,
                              "workdir": os.path.relpath(wdir, ROOT),
                              "channel": JUDGE_OF_MODEL[judge],
                              "arith_rows": rows, "arith_invalid": invalid,
                              "had_verdict": bool(v and isinstance(v.get("scores"), dict))},
                             **slot_rec))
        log("  judge=%s attempt=%d invalid=%s" % (judge, attempt, invalid))
        if not invalid:
            v_final = v
            return v_final, attempts, False
        if attempt == 1:
            log("  判词无效（arith 不一致或 no-verdict）→ 同裁判新会话重发 1 次（%s）" % judge)
    return None, attempts, True  # 仍不一致/通道故障 → fail-closed


def task_unit_dirs(task, fam, reps):
    out = []
    for rep in reps:
        ad = os.path.join(RUN, fam, "meeting-minutes-%s-rep%d" % (task, rep))
        out.append((rep, os.path.join(ad, "baseline")))
        out.append((rep, os.path.join(ad, "treatment")))
    return out


def run_task(task, results, jp):
    ti = TASK_ORDER.index(task) + 1
    fam = FAM_OF_TASK[task]
    reps = [1, 2, 3]
    units = task_unit_dirs(task, fam, reps)
    missing = [u for u in units
               if not os.path.isfile(os.path.join(u[1], "records", "arm_run.json"))]
    if missing:
        results[task] = {"status": "infra_error", "missing": [os.path.relpath(m, ROOT)
                                                              for m in missing]}
        return
    staging = os.path.join(RUN, "judge-work", "_staging", task)
    _STAGING_ROOTS[task] = staging
    seed = zlib.crc32(("mm-retest-20261009|%s" % task).encode()) & 0xffffff
    staged, mapping, meta = stage_products(task, units, staging, seed)
    J.write_json(os.path.join(RUN, "records", "judge-staging", "%s.mapping.json" % task),
                 {"mapping": mapping, "staged": staged, "seed": seed})
    brief = build_brief(task, staged)
    brief_sha = hashlib.sha256(brief.encode("utf-8")).hexdigest()
    rubric_sha = sha256_file(os.path.join(MERGED, task, "rubric.md"))
    notes_sha = sha256_file(os.path.join(MERGED, task, "judge-notes.md"))
    weights, sum_w = rubric_weights(task)
    panel = PANEL[fam][:]
    if ti % 2 == 0:            # 位置轮换：偶数任务槽序互换
        panel = panel[::-1]
    pids = [s["pid"] for s in staged]
    workroot = os.path.join(RUN, "judge-work", task)
    os.makedirs(workroot, exist_ok=True)
    record = {"task": task, "tested_family": fam, "panel": panel,
              "panel_rotation": "task_idx=%d parity=%s" % (ti, ti % 2),
              "staging_seed": seed, "brief_sha256": brief_sha, "rubric_sha256": rubric_sha,
              "judge_notes_sha256": notes_sha,
              "rubric_weights": {k: [n, float(w)] for k, (n, w) in weights.items()},
              "judge_prompt_version": J.JUDGE_PROMPT_VERSION, "protocol": "weighted-dims",
              "judges": {}, "fail_closed": None}
    print("JUDGE %s fam=%s panel=%s seed=%d" % (task, fam, panel, seed), flush=True)

    verdicts, invalid_any = {}, False
    for judge in panel:
        v, attempts, invalid = judge_validated(judge, workroot, brief, staged, task,
                                               weights, sum_w, pids, task, print)
        record["judges"][judge] = {"attempts": attempts, "final_invalid": invalid,
                                   "verdict": v}
        if invalid:
            invalid_any = True
            bad = judge
            ch_fail = any("channel_failure" in a for a in attempts)
            record["fail_closed"] = (
                "判官 %s 通道故障（重试 ≤1 仍失败）如实降级 → fail-closed 记 unjudgeable" % bad
                if ch_fail else
                "judge_arith_check 复发不一致（重发 1 次后仍无效）→ fail-closed 记 unjudgeable（§9-3；blind=false）")
    if invalid_any:
        jc_action = "unjudgeable"
        record["resample"] = None
        final_scores, per_judge = {}, {}
    else:
        for judge in panel:
            verdicts[judge] = record["judges"][judge]["verdict"]
        per_judge = {}
        for judge, v in verdicts.items():
            for pid, sv in v["scores"].items():
                per_judge.setdefault(pid, {})[judge] = sv.get("total")
        final_scores = {pid: (round(sum(x for x in d.values() if x is not None)
                                    / len([x for x in d.values() if x is not None]), 2)
                              if d and any(x is not None for x in d.values()) else None)
                        for pid, d in per_judge.items()}
        # §3.4 抽检：按任务序轮换抽 1 份（覆盖两臂），同裁判匿名复评（Q1）
        sampled_pid = staged[(ti - 1) % len(staged)]["pid"]
        judge0 = panel[0]
        first = per_judge.get(sampled_pid, {}).get(judge0)
        staged2 = [{"pid": "Q1", "files": next(s["files"] for s in staged
                                               if s["pid"] == sampled_pid)}]
        rstage = os.path.join(staging, sampled_pid + "_resample")
        os.makedirs(rstage, exist_ok=True)
        for fn in os.listdir(os.path.join(staging, sampled_pid)):
            shutil.copyfile(os.path.join(staging, sampled_pid, fn), os.path.join(rstage, fn))
        _STAGING_ROOTS["%s-resample" % task] = rstage
        brief2 = build_brief(task, staged2)
        r_attempts = []
        second = None
        for attempt in (1, 2):
            label = "%s-resample" % task if attempt == 1 else "%s-resample-resend" % task
            rdir = os.path.join(workroot, "%s-a%d" % (judge0.replace("/", "_"), attempt))
            os.makedirs(rdir, exist_ok=True)
            try:
                rec2, v2, slot_rec2 = judge_call(judge0, rdir, brief2, staged2, label,
                                                 phase="resample" if attempt == 1
                                                 else "resample-resend")
            except Exception as exc:
                r_attempts.append({"attempt": attempt, "judge": judge0, "phase": label,
                                   "channel_failure": type(exc).__name__ + ":" + str(exc)[:300],
                                   "arith_invalid": True, "had_verdict": False})
                continue
            rrows, rinvalid = arith_check(v2, weights, sum_w, ["Q1"])
            r_attempts.append(dict({"attempt": attempt, "judge": judge0, "phase": label,
                                    "workdir": os.path.relpath(rdir, ROOT),
                                    "arith_rows": rrows, "arith_invalid": rinvalid,
                                    "had_verdict": bool(v2 and isinstance(v2.get("scores"),
                                                                         dict))}, **slot_rec2))
            if not rinvalid:
                second = v2["scores"]["Q1"]["total"]
                break
            if attempt == 1:
                print("  复评 arith 不一致 → 重发 1 次", flush=True)
        if second is None:
            record["fail_closed"] = ("§3.4 复评判词 judge_arith_check 重发后仍不一致/通道故障 → "
                                     "fail-closed 记 unjudgeable")
        jc_delta = (abs(first - second) if (first is not None and second is not None) else None)
        action = "pass"
        if record["fail_closed"]:
            action = "unjudgeable"
        elif jc_delta is not None and jc_delta > 2:
            # §3.4-2 + §9-4：换面板内另一异族裁判重评全任务；比较对象=换上裁判对被抽产物
            # 分数 vs 失格裁判（judge0）对该产物的首评
            action = "re-judged"
            judge1 = panel[1]
            if judge1 in verdicts:
                news = {pid: sv.get("total") for pid, sv in verdicts[judge1]["scores"].items()}
                final_scores = {pid: news.get(pid) for pid in final_scores}
                if news.get(sampled_pid) is not None and first is not None \
                        and abs(news[sampled_pid] - first) > 2:
                    action = "unjudgeable"
            else:
                action = "unjudgeable"
        record["resample"] = {"judge": judge0, "sampled_pid": sampled_pid,
                              "anonymized_as": "Q1", "attempts": r_attempts,
                              "brief_sha256": hashlib.sha256(brief2.encode("utf-8")).hexdigest()}
        record["judge_consistency"] = {"task": task, "resample_of": sampled_pid,
                                       "first_score": first, "second_score": second,
                                       "delta": jc_delta, "action": action,
                                       "compare_sides": "同裁判首评（panel[0] 对 %s）vs 匿名复评（Q1）"
                                                        % sampled_pid}
        jc_action = action
    record["final_scores"] = final_scores
    record["per_judge_scores"] = per_judge
    J.write_json(os.path.join(RUN, "records", "judging", "%s.json" % task), record)
    status = "judged" if jc_action != "unjudgeable" else "unjudgeable"
    results[task] = {"status": status, "tested_family": fam, "panel": panel,
                     "final_scores": final_scores, "mapping": mapping,
                     "judge_consistency": record.get("judge_consistency"),
                     "fail_closed": record.get("fail_closed")}
    J.write_json(jp, results)
    print("JUDGED %s fam=%s status=%s action=%s" % (task, fam, status, jc_action), flush=True)


def main():
    only = None
    if "--task" in sys.argv:
        only = sys.argv[sys.argv.index("--task") + 1]
    jp = os.path.join(RUN, "records", "judging", "meeting-minutes.json")
    results = json.load(open(jp, encoding="utf-8")) if os.path.isfile(jp) else {}
    try:
        import docx  # noqa: F401
    except ImportError:
        vpy = os.path.join(ROOT, ".venv-dev", "bin", "python")
        if os.path.exists(vpy) and os.path.abspath(sys.executable) != os.path.abspath(vpy):
            os.execv(vpy, [vpy, os.path.abspath(__file__)] + list(sys.argv[1:]))
        print("python-docx 缺失且无 .venv-dev", file=sys.stderr)
        return 2
    for task in TASK_ORDER:
        if only and task != only:
            continue
        if isinstance(results.get(task), dict) and results[task].get("status") in (
                "judged", "unjudgeable"):
            print("SKIP judged %s" % task, flush=True)
            continue
        try:
            run_task(task, results, jp)
        except Exception as exc:
            import traceback
            results[task] = {"status": "infra_error",
                             "error": type(exc).__name__ + ":" + str(exc)[:300],
                             "trace": traceback.format_exc()[-1500:]}
            J.write_json(jp, results)
            print("INFRA_ERROR %s: %s" % (task, exc), flush=True)
            continue
    J.write_json(os.path.join(RUN, "records", "judge-results.json"),
                 {"meeting-minutes": results})
    bad = [1 for t in results.values() if t.get("status") in ("unjudgeable", "infra_error")]
    return 1 if bad else 0


if __name__ == "__main__":
    sys.exit(main())
