#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""会议通知 文书生成器（确定性实现）。

命令行契约：
    python3 generate.py --template 会议通知 --data <数据json路径> --outdir <输出目录>

每次调用在输出目录内生成恰好两个产物：
    文书.docx   —— 会议通知正文（公文版式）
    fields.json —— 字段回执（filled/missing + 输出路径）

确定性保证：
  - 不读取系统时间、不使用随机源；文档全部内容均由输入数据 + 固定连接文案拼装；
  - docx 保存后对 zip 包做归一化复写（全部条目 mtime 固定为 1980-01-01 00:00:00、
    统一 external_attr/create_system/压缩方式、保持条目顺序），
    因此同输入两次调用产物逐字节一致。
"""

import argparse
import json
import os
import posixpath
import sys
import tempfile
import zipfile

try:
    import docx  # noqa: F401
except ModuleNotFoundError:  # 本机解释器为 venv，python-docx 装在用户 site-packages
    import site

    for _cand in (site.getusersitepackages(),):
        if os.path.isdir(_cand) and _cand not in sys.path:
            sys.path.insert(0, _cand)
    import docx

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Cm, Pt

TEMPLATE_NAME = "会议通知"
DOCX_NAME = "文书.docx"
FIELDS_NAME = "fields.json"

# 会议通知模板的已知字段（顺序即正文字段回执顺序）
REQUIRED_FIELDS = [
    "会议名称", "召开单位", "主送对象", "会议时间", "会议地点",
    "参会人员", "会议议题", "会议要求", "联系人", "联系电话", "发文日期",
]

BODY_FONT_EAST = "仿宋_GB2312"
BODY_FONT_ASCII = "Times New Roman"
TITLE_FONT_EAST = "黑体"
BODY_SIZE = Pt(16)          # 三号
TITLE_SIZE = Pt(22)         # 二号
LINE_SPACING = Pt(28)       # 公文正文固定行距 28 磅
FIRST_LINE_INDENT = Pt(32)  # 首行缩进两字（16pt × 2）

_CN_NUM = "零一二三四五六七八九十"


def _cn_index(i: int) -> str:
    """1 -> 一, 2 -> 二, ... 11 -> 十一（列表序号用）。"""
    if 0 <= i <= 10:
        return _CN_NUM[i]
    if 11 <= i <= 19:
        return "十" + _CN_NUM[i - 10]
    return str(i)


def _set_east_asia(run, east: str) -> None:
    rpr = run._element.get_or_add_rPr()
    rfonts = rpr.get_or_add_rFonts()
    rfonts.set(qn("w:eastAsia"), east)


def _char_pt(text: str) -> float:
    """按 CJK 全宽 16pt、ASCII 半宽 8pt 估算字符串排版宽度（落款对位用）。"""
    return sum(16.0 if ord(ch) > 0x2E7F else 8.0 for ch in text)


def build_document(data: dict) -> Document:
    doc = Document()

    # 页面与版心（A4，公文页边距：上 3.7 下 3.5 左 2.8 右 2.6，单位 cm）
    sec = doc.sections[0]
    sec.page_width = Cm(21.0)
    sec.page_height = Cm(29.7)
    sec.top_margin = Cm(3.7)
    sec.bottom_margin = Cm(3.5)
    sec.left_margin = Cm(2.8)
    sec.right_margin = Cm(2.6)

    # Normal 样式兜底：三号仿宋，中西文分设
    normal = doc.styles["Normal"]
    normal.font.name = BODY_FONT_ASCII
    normal.font.size = BODY_SIZE
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), BODY_FONT_EAST)

    def para(align=WD_ALIGN_PARAGRAPH.JUSTIFY, indent=None, right_indent=None,
             line=LINE_SPACING, before=Pt(0), after=Pt(0)):
        p = doc.add_paragraph()
        pf = p.paragraph_format
        pf.alignment = align
        if indent is not None:
            pf.first_line_indent = indent
        if right_indent is not None:
            pf.right_indent = right_indent
        pf.line_spacing = line
        pf.space_before = before
        pf.space_after = after
        return p

    def body_run(p, text, *, east=BODY_FONT_EAST, ascii_=BODY_FONT_ASCII,
                 size=BODY_SIZE, bold=False):
        r = p.add_run(text)
        r.font.name = ascii_
        r.font.size = size
        r.font.bold = bold
        _set_east_asia(r, east)
        return r

    name = data.get("会议名称", "")

    # 标题：二号黑体居中
    if name:
        p_title = para(align=WD_ALIGN_PARAGRAPH.CENTER, line=Pt(33), after=Pt(24))
        body_run(p_title, f"关于召开{name}的通知",
                 east=TITLE_FONT_EAST, ascii_=TITLE_FONT_EAST, size=TITLE_SIZE)

    # 主送对象：顶格
    if data.get("主送对象"):
        p_recv = para(align=WD_ALIGN_PARAGRAPH.LEFT)
        body_run(p_recv, f"{data['主送对象']}：")

    # 开头段
    if data.get("会议时间") or data.get("会议地点") or name:
        p_open = para(indent=FIRST_LINE_INDENT)
        body_run(
            p_open,
            f"定于{data.get('会议时间', '')}在{data.get('会议地点', '')}召开{name}，"
            f"现将有关事项通知如下：",
        )

    # 一至四、固定字段条目（加粗条目名 + 原值）
    for label, key in ((
        ("一、会议时间：", "会议时间"),
        ("二、会议地点：", "会议地点"),
        ("三、参会人员：", "参会人员"),
        ("四、会议议题：", "会议议题"),
    )):
        if data.get(key):
            p_item = para(indent=FIRST_LINE_INDENT)
            body_run(p_item, label, bold=True)
            body_run(p_item, str(data[key]))

    # 五、会议要求（逐项（一）（二）……）
    reqs = [str(x) for x in (data.get("会议要求") or []) if str(x)]
    if reqs:
        p_req = para(indent=FIRST_LINE_INDENT)
        body_run(p_req, "五、会议要求：", bold=True)
        for i, item in enumerate(reqs, start=1):
            p_sub = para(indent=FIRST_LINE_INDENT)
            body_run(p_sub, f"（{_cn_index(i)}）{item}")

    # 联系方式与结语
    if data.get("联系人") or data.get("联系电话"):
        p_contact = para(indent=FIRST_LINE_INDENT)
        body_run(
            p_contact,
            f"如有疑问，请联系{data.get('联系人', '')}，"
            f"联系电话：{data.get('联系电话', '')}。",
        )
    p_close = para(indent=FIRST_LINE_INDENT)
    body_run(p_close, "特此通知。")

    # 落款：右对齐，成文日期右空四字、署名居日期上方（宽度估算法居中）
    unit, date_str = data.get("召开单位", ""), data.get("发文日期", "")
    date_right = Pt(64)  # 右空四字（16pt × 4）
    unit_right = Pt(64 + max(0.0, (_char_pt(date_str) - _char_pt(unit)) / 2.0))
    if unit:
        p_unit = para(align=WD_ALIGN_PARAGRAPH.RIGHT, right_indent=unit_right)
        body_run(p_unit, unit)
    if date_str:
        p_date = para(align=WD_ALIGN_PARAGRAPH.RIGHT, right_indent=date_right)
        body_run(p_date, date_str)

    # 清除模板 core.xml 中自带的固定历史时间戳（2013 常量，非本机时间；
    # 删除后产物不含任何时间类字符串）
    core = doc.core_properties._element
    _dcterms = "{http://purl.org/dc/terms/}"
    for tag in ("created", "modified"):
        el = core.find(_dcterms + tag)
        if el is not None:
            core.remove(el)

    return doc


def normalize_docx_zip(path: str) -> None:
    """复写 zip：固定 mtime/属性/压缩，保持条目顺序——保证逐字节可复现。"""
    with zipfile.ZipFile(path, "r") as zin:
        entries = [(zi.filename, zin.read(zi.filename)) for zi in zin.infolist()]
    with zipfile.ZipFile(path, "w") as zout:
        for name, blob in entries:
            zi = zipfile.ZipInfo(name, date_time=(1980, 1, 1, 0, 0, 0))
            zi.compress_type = zipfile.ZIP_DEFLATED
            zi.external_attr = 0o644 << 16
            zi.create_system = 3
            zout.writestr(zi, blob)


def main() -> int:
    ap = argparse.ArgumentParser(description="生成 会议通知 文书.docx 与 fields.json")
    ap.add_argument("--template", required=True, help="模板名（当前支持：会议通知）")
    ap.add_argument("--data", required=True, help="输入数据 json 路径")
    ap.add_argument("--outdir", required=True, help="输出目录")
    args = ap.parse_args()

    if args.template != TEMPLATE_NAME:
        print(f"错误：不支持的模板 {args.template!r}（当前支持：{TEMPLATE_NAME}）",
              file=sys.stderr)
        return 2

    with open(args.data, "r", encoding="utf-8") as f:
        data = json.load(f)

    filled = {k: data[k] for k in REQUIRED_FIELDS
              if k in data and data[k] not in (None, "", [])}
    missing = [k for k in REQUIRED_FIELDS if k not in filled]

    outdir = args.outdir.rstrip("/")
    os.makedirs(outdir, exist_ok=True)
    docx_path = posixpath.join(outdir, DOCX_NAME)
    fields_path = posixpath.join(outdir, FIELDS_NAME)

    # docx：先写临时文件，归一化 zip 后再落位（输出目录不出现中间文件）
    doc = build_document(data)
    fd, tmp = tempfile.mkstemp(suffix=".docx")
    os.close(fd)
    try:
        doc.save(tmp)
        normalize_docx_zip(tmp)
        os.replace(tmp, docx_path)
        os.chmod(docx_path, 0o644)  # mkstemp 出身是 0600，统一产物权限
    finally:
        if os.path.exists(tmp):
            os.remove(tmp)

    # fields.json：固定键序 + 固定序列化参数，逐字节可复现
    fields = {
        "template": TEMPLATE_NAME,
        "data": args.data,
        "outputs": {
            DOCX_NAME: docx_path,
            FIELDS_NAME: fields_path,
        },
        "filled_fields": filled,
        "missing_fields": missing,
    }
    with open(fields_path, "w", encoding="utf-8", newline="\n") as f:
        f.write(json.dumps(fields, ensure_ascii=False, indent=2) + "\n")

    return 0


if __name__ == "__main__":
    raise SystemExit(main())
