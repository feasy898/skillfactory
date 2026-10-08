#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""SF-0008 office 臂生成器（ab3-empty-and-string-lists）。

命令行契约：
  python3 gen_office.py --template <周报|工作总结> --data <数据json路径> --outdir <输出目录>
每次调用在 outdir 恰好生成两件产物：文书.docx 与 fields.json。

行为规则（见 brief.md）：
  * 空列表字段按必填缺失处理：正文渲染占位条目「1．____」，fields.json 记占位语义（value=null）。
  * 字符串型列表按单项化处理：正文渲染为单项条目，fields.json 中该字段值为单元素数组。
  * 选填小节缺失整节省略（正文不出现该节头），fields.json 仍记 missing（required=false）。
  * 必填缺失以 ____ 占位（标题如「____工作周报」「____2026年9月工作总结」）。
  * 小节编号「一、二、三、」，条目编号「1．2．3．」（全角实点）。
  * 落款右对齐，不因空日期/缺主体破坏（缺失的落款行以 ____ 占位保持结构）。
"""
from __future__ import annotations

import argparse
import json
import sys
from pathlib import Path

try:
    from docx import Document
except ImportError:
    # 本机 venv 未装 python-docx，回落到用户站点安装（存在性由调用环境保证）
    sys.path.append(str(Path.home() / ".local" / "lib" / "python3.12" / "site-packages"))
    from docx import Document

from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Cm, Pt

PLACEHOLDER = "____"
FULLWIDTH_COLON = "："  # U+FF1A
ITEM_DOT = "．"  # U+FF0E 全角实点
SECTION_NUMERALS = ["一", "二", "三", "四", "五", "六", "七", "八", "九", "十"]

TEMPLATES = {
    "周报": {
        "title_parts": ["周期"],
        "title_suffix": "工作周报",
        "fields": [
            {"name": "填报人", "kind": "scalar", "required": True, "placement": "signature", "signature_style": "labeled"},
            {"name": "周期", "kind": "scalar", "required": True, "placement": "header"},
            {"name": "本周工作内容", "kind": "list", "required": True, "placement": "section"},
            {"name": "下周工作计划", "kind": "list", "required": True, "placement": "section"},
            {"name": "问题与需协调事项", "kind": "list", "required": False, "placement": "section"},
            {"name": "报送日期", "kind": "scalar", "required": True, "placement": "signature", "signature_style": "labeled"},
        ],
    },
    "工作总结": {
        "title_parts": ["总结主体", "总结时段"],
        "title_suffix": "工作总结",
        "fields": [
            {"name": "总结主体", "kind": "scalar", "required": True, "placement": "signature", "signature_style": "bare"},
            {"name": "总结时段", "kind": "scalar", "required": True, "placement": "header"},
            {"name": "工作回顾", "kind": "list", "required": True, "placement": "section"},
            {"name": "主要成绩", "kind": "list", "required": True, "placement": "section"},
            {"name": "下一步工作打算", "kind": "list", "required": True, "placement": "section"},
            {"name": "存在问题", "kind": "list", "required": False, "placement": "section"},
            {"name": "成文日期", "kind": "scalar", "required": False, "placement": "signature", "signature_style": "bare"},
        ],
    },
}


def normalize_field(spec: dict, raw) -> dict:
    """把输入数据归一为 fields.json 记录（缺值时 value 必须为 null）。"""
    entry = {
        "name": spec["name"],
        "kind": spec["kind"],
        "required": spec["required"],
    }
    if spec["kind"] == "scalar":
        if raw is None:
            entry.update(filled=False, missing=True, value=None, missing_reason="absent")
        elif isinstance(raw, str) and raw.strip() == "":
            entry.update(filled=False, missing=True, value=None, missing_reason="empty_string")
        else:
            entry.update(filled=True, missing=False, value=raw if isinstance(raw, str) else str(raw))
    else:  # list
        if raw is None:
            entry.update(filled=False, missing=True, value=None, missing_reason="absent")
        elif isinstance(raw, str):
            if raw.strip() == "":
                entry.update(filled=False, missing=True, value=None, missing_reason="empty_string", input_type="string")
            else:
                # 字符串型列表 → 单项化：fields.json 值为单元素数组
                entry.update(filled=True, missing=False, value=[raw], input_type="string",
                             normalization="single_item_list")
        elif isinstance(raw, list):
            if len(raw) == 0:
                entry.update(filled=False, missing=True, value=None, missing_reason="empty_list")
            else:
                items = [x if isinstance(x, str) else json.dumps(x, ensure_ascii=False) for x in raw]
                entry.update(filled=True, missing=False, value=items, input_type="list")
        else:
            entry.update(filled=True, missing=False, value=[str(raw)], input_type="other",
                         normalization="single_item_list")
    return entry


def scalar_text(entry: dict) -> str:
    return entry["value"] if entry["filled"] else PLACEHOLDER


def add_paragraph(doc, text, *, size=12, bold=False, east="宋体", align=None,
                  indent_chars=0.0, space_before=0, space_after=6, line_spacing=1.5):
    p = doc.add_paragraph()
    if align is not None:
        p.alignment = align
    pf = p.paragraph_format
    pf.line_spacing = line_spacing
    pf.space_before = Pt(space_before)
    pf.space_after = Pt(space_after)
    if indent_chars:
        pf.first_line_indent = Pt(size * indent_chars)
    r = p.add_run(text)
    r.font.name = "Times New Roman"
    r.font.size = Pt(size)
    r.font.bold = bold
    rPr = r._element.get_or_add_rPr()
    rPr.get_or_add_rFonts().set(qn("w:eastAsia"), east)
    return p


def section_entries(entries):
    rendered = [e for e in entries
                if e_placement(e) == "section" and (e["filled"] or e["required"])]
    return rendered


def e_placement(entry):
    return entry["_placement"]


def build_docx(template_name: str, spec: dict, entries: list):
    doc = Document()
    section = doc.sections[0]
    section.page_width, section.page_height = Cm(21.0), Cm(29.7)
    for attr in ("top_margin", "bottom_margin", "left_margin", "right_margin"):
        setattr(section, attr, Cm(2.54))
    normal = doc.styles["Normal"]
    normal.font.name = "Times New Roman"
    normal.font.size = Pt(12)
    normal.element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:eastAsia"), "宋体")

    by_name = {e["name"]: e for e in entries}

    # 标题：缺失部分以 ____ 占位（如「____工作周报」「____2026年9月工作总结」）
    title_text = "".join(scalar_text(by_name[n]) for n in spec["title_parts"]) + spec["title_suffix"]
    add_paragraph(doc, title_text, size=22, bold=True, east="黑体",
                  align=WD_ALIGN_PARAGRAPH.CENTER, space_after=18)

    # 文首信息行（如 周期：____ / 总结时段：2026年9月），全角冒号
    for e in entries:
        if e_placement(e) == "header" and (e["filled"] or e["required"]):
            add_paragraph(doc, f"{e['name']}{FULLWIDTH_COLON}{scalar_text(e)}", indent_chars=2)

    # 小节：仅渲染已填或必填（缺失去占位）的小节；选填缺失整节省略
    for idx, e in enumerate(section_entries(entries), start=1):
        add_paragraph(doc, f"{SECTION_NUMERALS[idx - 1]}、{e['name']}",
                      size=14, bold=True, east="黑体", indent_chars=2,
                      space_before=6, space_after=3)
        if e["filled"]:
            items = e["value"]
        else:
            items = [PLACEHOLDER]  # 空列表 → 占位条目「1．____」
        for i, item in enumerate(items, start=1):
            add_paragraph(doc, f"{i}{ITEM_DOT}{item}", indent_chars=2)

    # 落款：右对齐；缺失行以 ____ 占位，不因空日期/缺主体破坏结构
    sig_entries = [e for e in entries if e_placement(e) == "signature"]
    for i, e in enumerate(sig_entries):
        style = e.get("_signature_style", "labeled")
        if e["filled"]:
            text = f"{e['name']}{FULLWIDTH_COLON}{e['value']}" if style == "labeled" else e["value"]
        else:
            text = f"{e['name']}{FULLWIDTH_COLON}{PLACEHOLDER}" if style == "labeled" else PLACEHOLDER
        add_paragraph(doc, text, align=WD_ALIGN_PARAGRAPH.RIGHT,
                      space_before=12 if i == 0 else 0, space_after=3)

    return doc, title_text


def build_fields_json(template_name: str, data_source: str, spec: dict, entries: list) -> dict:
    rendered_sections = {e["name"] for e in section_entries(entries)}
    for e in entries:
        where = []
        if e["name"] in spec["title_parts"]:
            where.append("title")
        if e_placement(e) == "header" and (e["filled"] or e["required"]):
            where.append("header")
        if e["name"] in rendered_sections:
            where.append("section")
        if e_placement(e) == "signature":
            where.append("signature")
        e["rendered_in"] = where
        if not e["filled"]:
            if e["kind"] == "list":
                if e["required"]:
                    # 空列表按必填缺失处理 → 占位语义
                    e["placeholder"] = True
                    e["placeholder_entries"] = [f"1{ITEM_DOT}{PLACEHOLDER}"]
                else:
                    e["section_omitted"] = True
            else:
                e["placeholder_text"] = PLACEHOLDER
    return {
        "template": template_name,
        "data_source": data_source,
        "filled_fields": [e["name"] for e in entries if e["filled"]],
        "missing_fields": [e["name"] for e in entries if not e["filled"]],
        "fields": [{k: v for k, v in e.items() if not k.startswith("_")} for e in entries],
    }


def main(argv=None) -> int:
    parser = argparse.ArgumentParser(description="生成 周报/工作总结 文书.docx 与 fields.json")
    parser.add_argument("--template", required=True, choices=sorted(TEMPLATES), help="模板名：周报 或 工作总结")
    parser.add_argument("--data", required=True, help="数据 json 路径")
    parser.add_argument("--outdir", required=True, help="输出目录")
    args = parser.parse_args(argv)

    data_path = Path(args.data)
    if not data_path.is_file():
        parser.error(f"--data 文件不存在: {args.data}")
    data = json.loads(data_path.read_text(encoding="utf-8"))
    if not isinstance(data, dict):
        parser.error("--data 顶层必须是 JSON 对象")

    spec = TEMPLATES[args.template]
    entries = []
    for field_spec in spec["fields"]:
        entry = normalize_field(field_spec, data.get(field_spec["name"]))
        entry["_placement"] = field_spec["placement"]
        entry["_signature_style"] = field_spec.get("signature_style", "labeled")
        entries.append(entry)

    outdir = Path(args.outdir)
    outdir.mkdir(parents=True, exist_ok=True)

    doc, title_text = build_docx(args.template, spec, entries)
    docx_path = outdir / "文书.docx"
    doc.save(docx_path)

    fields = build_fields_json(args.template, args.data, spec, entries)
    fields_path = outdir / "fields.json"
    fields_path.write_text(json.dumps(fields, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")

    print(f"[gen_office] template={args.template} title=「{title_text}」")
    print(f"[gen_office] 产物: {docx_path} / {fields_path}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
