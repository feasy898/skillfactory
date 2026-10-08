#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""verify_out.py — 对 out/ 产物按技能验收清单逐条自查（等价 eval 五项检查）。"""
import json
import re
import sys

from docx import Document

OUT_DIR = "out"
TRANSCRIPT = "对公系统切换部署会.txt"

LINE_RE = re.compile(r"^\s*\[(\d{1,2}:\d{2})\]\s*([^：:]+?)\s*[：:]\s*(.+?)\s*$")
ENTRY_RE = re.compile(r"^(\d+)\. 【([^】]+?) · ([^】]+?)】(.*)$")

SECTIONS = ["一、决议事项", "二、待办事项", "三、风险与关注"]

failures = []


def check(cond, msg):
    print(("PASS " if cond else "FAIL ") + msg)
    if not cond:
        failures.append(msg)


# 转写稿逐字溯源底表：(when, speaker, content) -> 原始行
lines = {}
with open(TRANSCRIPT, encoding="utf-8-sig") as f:
    for line in f.read().splitlines():
        m = LINE_RE.match(line)
        if m:
            lines[(m.group(1), m.group(2).strip(), m.group(3))] = line

doc = Document(OUT_DIR + "/纪要.docx")

# ---- 文档结构：段落流 + 表格 ----
paras = [p.text for p in doc.paragraphs]
print("== 段落 ==")
for p in paras:
    print(repr(p))
print("== 表格 ==")
for t in doc.tables:
    for row in t.rows:
        print(" | ".join(c.text for c in row.cells))

check(paras[0] == "会议纪要", "标题段落为「会议纪要」")
for s in SECTIONS:
    check(s in paras, "节标题逐字存在：%s" % s)
check(paras.index(SECTIONS[0]) < paras.index(SECTIONS[1]) < paras.index(SECTIONS[2]),
      "三节标题顺序正确")

# 各节条目切分
def entries_between(a, b):
    ia, ib = paras.index(a), paras.index(b)
    return [p for p in paras[ia + 1:ib] if p.strip()]

decision_entries = entries_between(SECTIONS[0], SECTIONS[1])
risk_entries = entries_between(SECTIONS[2], paras[-1] + "\x00") if False else \
    [p for p in paras[paras.index(SECTIONS[2]) + 1:] if p.strip()]

tables = doc.tables
check(len(tables) == 1, "恰有 1 个待办表格")
todo_table = tables[0]
check(len(todo_table.columns) == 3, "待办表格为 3 列")
hdr = [c.text for c in todo_table.rows[0].cells]
check(hdr == ["事项", "负责人", "期限"], "表头为 事项/负责人/期限：%r" % hdr)
todo_rows = [[c.text for c in r.cells] for r in todo_table.rows[1:]]

# ---- 逐字溯源（决议/风险条目 + 待办事项列） ----
def trace(when, speaker, content):
    return (when, speaker, content) in lines

def check_numbered(entries, label):
    for i, e in enumerate(entries, start=1):
        m = ENTRY_RE.match(e)
        if not m:
            check(False, "%s 第 %d 条格式不符：%r" % (label, i, e))
            continue
        n, when, speaker, content = m.groups()
        check(n == str(i), "%s 第 %d 条编号连续（实际 %s）" % (label, i, n))
        check(trace(when, speaker, content),
              "%s 第 %d 条可逐字溯源：[%s] %s" % (label, i, when, content[:30]))

check_numbered(decision_entries, "决议")
check_numbered(risk_entries, "风险")

for i, (item, owner, ddl) in enumerate(todo_rows, start=1):
    hit = None
    for (when, speaker, content), raw in lines.items():
        if content == item:
            hit = (when, speaker, content)
    check(hit is not None, "待办第 %d 条事项列可逐字溯源：%r" % (i, item[:30]))
    if hit:
        _, speaker, _ = hit
        owner_ok = owner == "待定" or owner == speaker and "负责人是我" in item \
            or owner in item
        check(owner_ok, "待办第 %d 条负责人合规（%r）" % (i, owner))
        ddl_ok = ddl == "待定" or ddl in item
        check(ddl_ok, "待办第 %d 条期限合规（%r）" % (i, ddl))

# ---- summary.json 四键一致性 ----
with open(OUT_DIR + "/summary.json", encoding="utf-8") as f:
    summary = json.load(f)
print("== summary.json ==")
print(json.dumps(summary, ensure_ascii=False, indent=2))

check(list(summary.keys()) == ["决议事项", "待办事项", "风险与关注", "合计"],
      "summary 四键名精确且有序")
check(summary["决议事项"] == len(decision_entries), "决议条数一致（docx %d / json %d）"
      % (len(decision_entries), summary["决议事项"]))
check(summary["待办事项"] == len(todo_rows), "待办条数一致（docx %d / json %d）"
      % (len(todo_rows), summary["待办事项"]))
check(summary["风险与关注"] == len(risk_entries), "风险条数一致（docx %d / json %d）"
      % (len(risk_entries), summary["风险与关注"]))
check(summary["合计"] == summary["决议事项"] + summary["待办事项"] + summary["风险与关注"],
      "合计=三者和")

# 元信息行一致性
meta = paras[1]
m = re.match(r"来源：(.+?) ｜ 解析发言 (\d+) 条（决议 (\d+) · 待办 (\d+) · 风险 (\d+)）", meta)
check(m is not None, "元信息行格式解析成功：%r" % meta)
if m:
    src, total, a, b, c = m.group(1), *map(int, m.groups()[1:])
    check(src == TRANSCRIPT, "元信息来源=输入文件名")
    check((a, b, c) == (len(decision_entries), len(todo_rows), len(risk_entries)),
          "元信息三数与实际一致")
    check(total == a + b + c == summary["合计"], "元信息合计=实际合计=summary 合计")

print()
if failures:
    print("验收失败 %d 项" % len(failures))
    sys.exit(1)
print("验收全部通过")
