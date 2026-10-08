#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""会议通知 文书生成器（确定性实现）。

命令行契约：
    python3 generate.py --template 会议通知 --data <数据json路径> --outdir <输出目录>
每次调用在 --outdir 下生成恰两个产物：
    文书.docx   会议通知正文（对齐/缩进/字体/行距均显式设定，不含时间戳/随机内容等不可复现源）
    fields.json 字段填写台账（固定缩进与分隔符的规范序列化，除 outputs 类路径键外逐字节可复现）

确定性保障：
  1. docx 内容完全由输入数据决定；core 属性（作者/创建时间/修改时间/修订号等）固定为常量，
     不取当前时间；
  2. 保存后对 zip 容器做规范化重写：固定条目时间戳 1980-01-01T00:00:00、固定外部属性、
     保持原条目顺序与压缩方式，消除 python-docx/zipfile 默认把当前时间写进 zip 条目
     头/中央目录导致的字节漂移；
  3. fields.json 用 ensure_ascii=False + indent=2 的 UTF-8 规范序列化并固定结尾换行。
"""

from __future__ import annotations

import argparse
import io
import json
import sys
import zipfile
from datetime import datetime
from pathlib import Path

try:
    from docx import Document
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.oxml.ns import qn
    from docx.shared import Cm, Pt
except ModuleNotFoundError:  # 兼容：venv 无包但用户目录已装 python-docx 的环境
    import os

    _user_site = os.path.expanduser("~/.local/lib/python3.12/site-packages")
    if os.path.isdir(_user_site) and _user_site not in sys.path:
        sys.path.insert(0, _user_site)
    from docx import Document
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.oxml.ns import qn
    from docx.shared import Cm, Pt

# ---------------------------------------------------------------- 常量（模板契约）

TEMPLATE_NAME = "会议通知"

# 模板字段（固定顺序，fields.json 与正文渲染均按此顺序，保证可复现）
FIELDS = [
    "会议名称",
    "召开单位",
    "主送对象",
    "会议时间",
    "会议地点",
    "参会人员",
    "会议议题",
    "会议要求",  # 列表型字段
    "联系人",
    "联系电话",
    "发文日期",
]
LIST_FIELDS = {"会议要求"}
PLACEHOLDER = "____"

# 版式常量（GB/T 9704 公文版式近似：三号仿宋正文、二号小标宋标题、28 磅固定行距）
PAGE_A4 = (Cm(21.0), Cm(29.7))
MARGINS = {"top": Cm(3.7), "bottom": Cm(3.5), "left": Cm(2.8), "right": Cm(2.6)}
FONT_TITLE_EA = "方正小标宋简体"
FONT_BODY_EA = "仿宋_GB2312"
FONT_HEAD_EA = "黑体"
FONT_ASCII = "Times New Roman"
SIZE_TITLE = Pt(22)  # 二号
SIZE_BODY = Pt(16)  # 三号
LINE_BODY = Pt(28)  # 固定行距 28 磅
LINE_TITLE = Pt(34)
INDENT_FIRST_LINE = Pt(32)  # 首行缩进 2 字符（16pt × 2）

# 固定 core 属性（不取当前时间，保证可复现）
FIXED_CREATED = datetime(2026, 1, 1, 0, 0, 0)

# zip 规范化常量
FIXED_ZIP_DATE = (1980, 1, 1, 0, 0, 0)  # zip 纪元，早于一切真实生成时间

CN_DIGITS = "零一二三四五六七八九"


def cn_num(n: int) -> str:
    """1..99 的中文数字（确定性纯函数）。"""
    if n < 1 or n > 99:
        raise ValueError(f"序号超出范围: {n}")
    if n < 10:
        return CN_DIGITS[n]
    if n == 10:
        return "十"
    if n < 20:
        return "十" + CN_DIGITS[n - 10]
    tens, ones = divmod(n, 10)
    return CN_DIGITS[tens] + "十" + (CN_DIGITS[ones] if ones else "")


# ---------------------------------------------------------------- 数据解析


def resolve_fields(data: dict) -> dict:
    """按模板字段顺序解析填写状态。缺失定义：键缺席 / None / 空串 / 空列表。"""
    resolved = {}
    for name in FIELDS:
        raw = data.get(name)
        if name in LIST_FIELDS:
            if isinstance(raw, list) and raw and all(str(x).strip() for x in raw):
                resolved[name] = {"status": "filled", "value": [str(x) for x in raw]}
            else:
                resolved[name] = {"status": "missing", "value": None}
        else:
            if raw is None or (isinstance(raw, str) and not raw.strip()):
                resolved[name] = {"status": "missing", "value": None}
            else:
                resolved[name] = {"status": "filled", "value": raw if isinstance(raw, str) else str(raw)}
    return resolved


# ---------------------------------------------------------------- docx 构建


def _add_run(paragraph, text: str, *, ea: str, size=SIZE_BODY, bold=False):
    run = paragraph.add_run(text)
    run.font.name = FONT_ASCII  # 西文/数字
    run.font.size = size
    run.font.bold = bold
    rpr = run._element.get_or_add_rPr()
    rfonts = rpr.get_or_add_rFonts()
    rfonts.set(qn("w:eastAsia"), ea)  # 中文字体
    return run


def _add_para(doc: Document, *, align=WD_ALIGN_PARAGRAPH.JUSTIFY, first_line=None,
              line=LINE_BODY, before=Pt(0), after=Pt(0)):
    p = doc.add_paragraph()
    p.alignment = align
    pf = p.paragraph_format
    if first_line is not None:
        pf.first_line_indent = first_line
    pf.line_spacing = line  # Length 赋值 → 固定行距(exact)
    pf.space_before = before
    pf.space_after = after
    return p


def build_document(resolved: dict) -> Document:
    def val(name: str) -> str:
        item = resolved[name]
        return item["value"] if item["status"] == "filled" else PLACEHOLDER

    doc = Document()

    # 页面设置（显式固定，A4 纵向 + 公文页边距）
    sec = doc.sections[0]
    sec.page_width, sec.page_height = PAGE_A4
    sec.top_margin = MARGINS["top"]
    sec.bottom_margin = MARGINS["bottom"]
    sec.left_margin = MARGINS["left"]
    sec.right_margin = MARGINS["right"]

    # 标题（居中、二号小标宋）
    p = _add_para(doc, align=WD_ALIGN_PARAGRAPH.CENTER, line=LINE_TITLE, after=Pt(12))
    _add_run(p, f"关于召开{val('会议名称')}的通知", ea=FONT_TITLE_EA, size=SIZE_TITLE)

    # 主送对象（顶格）
    p = _add_para(doc, align=WD_ALIGN_PARAGRAPH.LEFT)
    _add_run(p, f"{val('主送对象')}：", ea=FONT_BODY_EA)

    # 正文开头（首行缩进 2 字符）
    p = _add_para(doc, first_line=INDENT_FIRST_LINE)
    _add_run(
        p,
        f"经研究，定于{val('会议时间')}在{val('会议地点')}召开{val('会议名称')}，"
        f"现将有关事项通知如下：",
        ea=FONT_BODY_EA,
    )

    # 事项条目（小标题黑体 + 内容仿宋，同段两 run）
    items = [
        ("会议时间", "会议时间"),
        ("会议地点", "会议地点"),
        ("参会人员", "参会人员"),
        ("会议议题", "会议议题"),
    ]
    for idx, (field, label) in enumerate(items, start=1):
        p = _add_para(doc, first_line=INDENT_FIRST_LINE)
        _add_run(p, f"{cn_num(idx)}、{label}", ea=FONT_HEAD_EA)
        _add_run(p, f"：{val(field)}", ea=FONT_BODY_EA)

    p = _add_para(doc, first_line=INDENT_FIRST_LINE)
    _add_run(p, f"{cn_num(len(items) + 1)}、会议要求", ea=FONT_HEAD_EA)
    _add_run(p, "：", ea=FONT_HEAD_EA)

    req = resolved["会议要求"]
    if req["status"] == "filled":
        for i, item_text in enumerate(req["value"], start=1):
            p = _add_para(doc, first_line=INDENT_FIRST_LINE)
            _add_run(p, f"（{cn_num(i)}）{item_text}" + ("；" if i < len(req["value"]) else "。"),
                     ea=FONT_BODY_EA)
    else:
        p = _add_para(doc, first_line=INDENT_FIRST_LINE)
        _add_run(p, f"（一）{PLACEHOLDER}。", ea=FONT_BODY_EA)

    # 结语与联系方式
    p = _add_para(doc, first_line=INDENT_FIRST_LINE)
    _add_run(p, "特此通知。", ea=FONT_BODY_EA)
    p = _add_para(doc, first_line=INDENT_FIRST_LINE)
    _add_run(p, f"联系人：{val('联系人')}；联系电话：{val('联系电话')}", ea=FONT_BODY_EA)

    # 落款（右对齐：召开单位 / 发文日期）
    p = _add_para(doc, align=WD_ALIGN_PARAGRAPH.RIGHT, before=Pt(28))
    _add_run(p, val("召开单位"), ea=FONT_BODY_EA)
    p = _add_para(doc, align=WD_ALIGN_PARAGRAPH.RIGHT)
    _add_run(p, val("发文日期"), ea=FONT_BODY_EA)

    # core 属性全部固定为常量（杜绝"当前时间"类不可复现源）
    cp = doc.core_properties
    cp.title = f"关于召开{val('会议名称')}的通知"
    cp.subject = TEMPLATE_NAME
    cp.author = val("召开单位")
    cp.last_modified_by = val("召开单位")
    cp.category = "通知"
    cp.keywords = TEMPLATE_NAME
    cp.revision = 1
    cp.created = FIXED_CREATED
    cp.modified = FIXED_CREATED

    return doc


def normalize_docx_zip(path: Path) -> None:
    """zip 容器规范化：固定条目时间戳/外部属性，保持条目顺序与压缩方式。

    python-docx 保存时 zipfile 会把当前系统时间写进每个条目的 local file header
    与 central directory，导致同输入两次保存字节不一致；本函数用固定时间戳重写
    全部条目以消除该漂移。
    """
    with zipfile.ZipFile(path) as zin:
        entries = [(info.filename, zin.read(info.filename)) for info in zin.infolist()]
    with zipfile.ZipFile(path, "w", zipfile.ZIP_DEFLATED) as zout:
        for name, blob in entries:
            zi = zipfile.ZipInfo(name, date_time=FIXED_ZIP_DATE)
            zi.compress_type = zipfile.ZIP_DEFLATED
            zi.create_system = 3  # unix
            zi.external_attr = 0o644 << 16  # 固定权限位
            zout.writestr(zi, blob)


# ---------------------------------------------------------------- fields.json


def build_fields_payload(resolved: dict, template: str, outdir_arg: str) -> dict:
    filled = [n for n in FIELDS if resolved[n]["status"] == "filled"]
    missing = [n for n in FIELDS if resolved[n]["status"] == "missing"]
    posix_out = Path(outdir_arg).as_posix()
    return {
        "template": template,
        "fields": [
            {"name": n, "status": resolved[n]["status"], "value": resolved[n]["value"]}
            for n in FIELDS
        ],
        "filled_fields": filled,
        "missing_fields": missing,
        "outputs": {
            "outdir": posix_out,
            "document": f"{posix_out}/文书.docx",
            "fields": f"{posix_out}/fields.json",
        },
    }


def dump_json(payload: dict) -> bytes:
    """规范序列化：固定缩进、非转义中文、结尾一个换行。"""
    return (json.dumps(payload, ensure_ascii=False, indent=2) + "\n").encode("utf-8")


# ---------------------------------------------------------------- 主流程


def main(argv=None) -> int:
    parser = argparse.ArgumentParser(description="会议通知 文书生成器（确定性）")
    parser.add_argument("--template", required=True, help="模板名，本实现仅支持：会议通知")
    parser.add_argument("--data", required=True, help="输入数据 json 路径")
    parser.add_argument("--outdir", required=True, help="输出目录（每次调用生成 文书.docx 与 fields.json）")
    args = parser.parse_args(argv)

    if args.template != TEMPLATE_NAME:
        print(f"error: 不支持的模板 {args.template!r}，本生成器仅实现 {TEMPLATE_NAME!r}", file=sys.stderr)
        return 2

    data_path = Path(args.data)
    if not data_path.is_file():
        print(f"error: 数据文件不存在: {data_path}", file=sys.stderr)
        return 2
    data = json.loads(data_path.read_text(encoding="utf-8"))
    if not isinstance(data, dict):
        print("error: 数据 json 顶层必须是对象", file=sys.stderr)
        return 2

    resolved = resolve_fields(data)

    outdir = Path(args.outdir)
    outdir.mkdir(parents=True, exist_ok=True)

    doc = build_document(resolved)
    buf = io.BytesIO()
    doc.save(buf)
    docx_path = outdir / "文书.docx"
    docx_path.write_bytes(buf.getvalue())
    normalize_docx_zip(docx_path)

    fields_path = outdir / "fields.json"
    fields_path.write_bytes(dump_json(build_fields_payload(resolved, args.template, args.outdir)))

    print(f"generated: {docx_path.as_posix()} ({docx_path.stat().st_size} bytes), "
          f"{fields_path.as_posix()} ({fields_path.stat().st_size} bytes)")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
