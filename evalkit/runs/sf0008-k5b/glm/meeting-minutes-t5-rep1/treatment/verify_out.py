#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""verify_out.py — 按 docx-and-summary.md 验收清单逐条核对 out/ 产物。"""
import json
import re
import sys
from docx import Document

TR = "transcript.txt"
OUTDIR = "out"

LINE_RE = re.compile(r"^\s*\[(\d{1,2}:\d{2})\]\s*([^：:]+?)\s*[：:]\s*(.+?)\s*$")
CATS = (
    ("决议事项", ("决定", "议定", "同意", "结论")),
    ("待办事项", ("需要", "负责", "跟进", "截止")),
    ("风险与关注", ("风险", "延期", "不确定")),
)

# 期望分类：独立重算
lines = [l for l in open(TR, encoding="utf-8-sig").read().splitlines() if l.strip()]
expect = {"决议事项": [], "待办事项": [], "风险与关注": []}
for line in lines:
    m = LINE_RE.match(line)
    if not m:
        continue
    when, who, content = m.group(1), m.group(2).strip(), m.group(3)
    for cat, kws in CATS:
        if any(k in content for k in kws):
            expect[cat].append((when, who, content))
            break

fails = []


def check(cond, msg):
    print(("PASS " if cond else "FAIL ") + msg)
    if not cond:
        fails.append(msg)


doc = Document(OUTDIR + "/纪要.docx")
paras = [p.text for p in doc.paragraphs]
tables = doc.tables

# 1) 三节标题逐字齐全、顺序正确
for h in ("一、决议事项", "二、待办事项", "三、风险与关注"):
    check(h in paras, "节标题存在且逐字：%s" % h)
idx = [paras.index(h) for h in ("一、决议事项", "二、待办事项", "三、风险与关注") if h in paras]
check(len(idx) == 3 and idx == sorted(idx), "三节标题顺序正确")
check(paras[0] == "会议纪要", "首段标题=会议纪要")

# 2) 逐节条目：格式 + 与独立重算的期望集合完全一致
ENTRY_RE = re.compile(r"^(\d+)\. 【(.+?) · (.+?)】(.+)$")


def section_entries(start, end):
    out = []
    for p in paras[start + 1:end]:
        if p.strip():
            out.append(p)
    return out


dec = section_entries(idx[0], idx[1])
risk = section_entries(idx[2], len(paras))
for name, got in (("决议", dec), ("风险", risk)):
    ok_fmt = all(ENTRY_RE.match(p) for p in got)
    check(ok_fmt, "%s条目格式均为 `<i>. 【hh:mm · 说话人】内容`" % name)
    nums = [int(ENTRY_RE.match(p).group(1)) for p in got if ENTRY_RE.match(p)]
    check(nums == list(range(1, len(got) + 1)), "%s编号从 1 连续" % name)
    exp = expect[name + "事项" if name != "风险" else "风险与关注"]
    got_tuples = [(ENTRY_RE.match(p).group(2), ENTRY_RE.match(p).group(3), ENTRY_RE.match(p).group(4)) for p in got]
    check(got_tuples == exp, "%s条目与独立重算逐条一致（%d 条，含顺序）" % (name, len(exp)))

# 3) 待办表格
check(len(tables) == 1, "文档恰含 1 个表格")
t = tables[0]
check([c.text for c in t.rows[0].cells] == ["事项", "负责人", "期限"], "表头=事项/负责人/期限（3 列）")
rows = [[c.text for c in r.cells] for r in t.rows[1:]]
check(len(rows) == len(expect["待办事项"]), "待办数据行数=%d 与重算一致" % len(expect["待办事项"]))
exp_todo = [(w, s, c) for (w, s, c) in expect["待办事项"]]
check([(None, None, r[0]) for r in rows] == [(None, None, c) for (_, _, c) in exp_todo],
      "待办事项列逐字与原句一致（含行序）")

# 4) 负责人/期限：原文逐字可寻 或 待定（负责人允许 speaker 替换）
for r, (when, who, content) in zip(rows, exp_todo):
    owner, ddl = r[1], r[2]
    ok_o = owner == "待定" or owner in content or owner == who  # 我→说话人替换
    ok_d = ddl == "待定" or ddl in content
    check(ok_o and ok_d, "待办[%s…]: 负责人=%r 期限=%r 合规" % (content[:12], owner, ddl))

# 5) summary 四键一致
s = json.load(open(OUTDIR + "/summary.json", encoding="utf-8"))
check(list(s.keys()) == ["决议事项", "待办事项", "风险与关注", "合计"], "summary 四键名精确")
check(s["决议事项"] == len(dec) and s["待办事项"] == len(rows) and s["风险与关注"] == len(risk),
      "summary 数值与 docx 实际条数一致")
check(s["合计"] == s["决议事项"] + s["待办事项"] + s["风险与关注"], "合计=三者之和")

# 6) 溯源：每个条目都能在转写稿找到含该时间+说话人的出处行
src = open(TR, encoding="utf-8-sig").read()
for p in dec + risk:
    m = ENTRY_RE.match(p)
    w, who, content = m.group(2), m.group(3), m.group(4)
    src_lines = [l for l in lines if LINE_RE.match(l) and LINE_RE.match(l).group(1) == w
                 and LINE_RE.match(l).group(2).strip() == who]
    check(any(content == LINE_RE.match(l).group(3) for l in src_lines),
          "溯源 OK：【%s · %s】" % (w, who))
for r in rows:
    content = r[0]
    check(any(content == LINE_RE.match(l).group(3) for l in lines if LINE_RE.match(l)),
          "表格溯源 OK：%s…" % content[:12])

# 7) 中文字体 eastAsia=宋体 抽查
from docx.oxml.ns import qn
ok_font = True
for p in doc.paragraphs:
    for run in p.runs:
        rPr = run._element.rPr
        if rPr is None or rPr.rFonts is None or rPr.rFonts.get(qn("w:eastAsia")) != "宋体":
            ok_font = False
check(ok_font, "全文段落 run 的 eastAsia 字体=宋体")

# 8) 元信息行与合计一致
meta = paras[1]
check("解析发言 %d 条" % s["合计"] in meta and "来源：transcript.txt" in meta,
      "元信息行含来源文件名与合计条数：%r" % meta)

print("\n%s（%d 项失败）" % ("ALL CHECKS PASSED" if not fails else "FAILED", len(fails)))
sys.exit(1 if fails else 0)
