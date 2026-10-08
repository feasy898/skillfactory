#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""SF-0008 office 臂：四类公文模板生成器（周报 / 请示函 / 会议通知 / 工作总结）。

命令行契约：
    python3 gen_doc.py --template <周报|请示函|会议通知|工作总结> \
                       --data <数据json路径> --outdir <输出目录>

每次调用在 outdir 下生成两个产物：
    文书.docx   GB/T 9704 公文风格（A4 版面）
    fields.json 字段记账：title / fields 明细（名称、是否必填、状态、值、渲染方式）/ summary

缺数据不编造：必填缺失以 ____ 占位，选填缺失省略，均如实记入 fields.json（missing 的 value 为 null）。
"""

import argparse
import json
import os
import re
import sys

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Pt, Cm

# ---- GB/T 9704 版式常量 ----
FONT_TITLE_EA = "方正小标宋简体"    # 标题：二号小标宋
FONT_BODY_EA = "仿宋_GB2312"       # 正文：三号仿宋
FONT_LATIN = "Times New Roman"     # 西文与数字
SZ_TITLE = 22                      # 二号 ≈ 22pt
SZ_BODY = 16                       # 三号 = 16pt
LINE_TITLE = Pt(36)                # 标题固定行距
LINE_BODY = Pt(28)                 # 正文固定行距（每面约 22 行）
PLACEHOLDER = "____"               # 必填缺失占位符

# 字段角色：salutation=称谓顶格 / body=正文段落 / signature=落款右对齐 / note=成文日期下附注
# 标题为格式化模式串：{字段} 取已填值，缺失以 ____ 代入，不编造。
TEMPLATES = {
    "周报": {
        "fields": [
            ("部门", True, "body"),
            ("填报人", True, "body"),
            ("周期", True, "body"),
            ("本周工作内容", True, "body"),
            ("下周工作计划", True, "body"),
            ("问题与需协调事项", False, "body"),
            ("报送日期", True, "signature"),
        ],
        "title_pattern": "{部门}工作周报",
        "closing": None,
    },
    "请示函": {
        "fields": [
            ("请示事由", True, "body"),
            ("主送机关", True, "salutation"),
            ("请示缘由", True, "body"),
            ("请示事项", True, "body"),
            ("请示单位", True, "signature"),
            ("联系人", False, "note"),
            ("联系电话", False, "note"),
            ("成文日期", True, "signature"),
        ],
        "title_pattern": "关于{请示事由}的请示",
        "closing": "妥否，请批示。",
    },
    "会议通知": {
        "fields": [
            ("会议名称", True, "body"),
            ("召开单位", True, "signature"),
            ("主送对象", True, "salutation"),
            ("会议时间", True, "body"),
            ("会议地点", True, "body"),
            ("参会人员", False, "body"),
            ("会议议题", True, "body"),
            ("会议要求", False, "body"),
            ("联系人", False, "body"),
            ("联系电话", False, "body"),
            ("发文日期", True, "signature"),
        ],
        "title_pattern": "关于召开{会议名称}的通知",
        "closing": "特此通知。",
    },
    "工作总结": {
        "fields": [
            ("总结主体", True, "body"),
            ("总结时段", True, "body"),
            ("工作回顾", True, "body"),
            ("主要成绩", False, "body"),
            ("存在问题", False, "body"),
            ("下一步工作打算", False, "body"),
            ("成文日期", True, "signature"),
        ],
        "title_pattern": "{总结主体}{总结时段}工作总结",
        "closing": None,
    },
}


def is_missing(v):
    """空值 / 空串 / 空数组均按缺失记账，不编造。"""
    if v is None:
        return True
    if isinstance(v, str) and v.strip() == "":
        return True
    if isinstance(v, (list, tuple, dict)) and len(v) == 0:
        return True
    return False


def render_value(v):
    """标量原样；列表编号内联（多项「1. …；2. …」，单项不编号）。"""
    if isinstance(v, (list, tuple)):
        items = [str(x) for x in v]
        if len(items) > 1:
            return "；".join("{}. {}".format(i, t) for i, t in enumerate(items, 1))
        return items[0] if items else ""
    return str(v)


def style_run(run, size, ea=FONT_BODY_EA, bold=False):
    run.font.size = Pt(size)
    run.font.bold = bold
    rPr = run._element.get_or_add_rPr()
    rFonts = rPr.get_or_add_rFonts()
    rFonts.set(qn("w:ascii"), FONT_LATIN)
    rFonts.set(qn("w:hAnsi"), FONT_LATIN)
    rFonts.set(qn("w:eastAsia"), ea)


def set_first_line_chars(p, chars=2, size_pt=SZ_BODY):
    """首行缩进两字：同时写 firstLineChars（中文版式）与 firstLine（兜底）。"""
    p.paragraph_format.first_line_indent = Pt(size_pt * chars)
    ind = p._p.get_or_add_pPr().find(qn("w:ind"))
    if ind is not None:
        ind.set(qn("w:firstLineChars"), str(chars * 100))


def add_para(doc, text, *, align=None, size=SZ_BODY, ea=FONT_BODY_EA, bold=False,
             indent_chars=0, line=LINE_BODY, space_before=None, space_after=None,
             right_indent_pt=None):
    p = doc.add_paragraph()
    pf = p.paragraph_format
    if align is not None:
        p.alignment = align
    pf.line_spacing = line
    pf.space_before = Pt(space_before) if space_before else Pt(0)
    pf.space_after = Pt(space_after) if space_after else Pt(0)
    if indent_chars:
        set_first_line_chars(p, indent_chars, size)
    if right_indent_pt:
        pf.right_indent = Pt(right_indent_pt)
    run = p.add_run(text)
    style_run(run, size, ea=ea, bold=bold)
    return p


def build_docx(template, title, entries, outdir):
    """entries: [(name, required, role, value, missing, rendered_line), ...]"""
    doc = Document()
    sec = doc.sections[0]
    # A4 版面， GB/T 9704 页边距：上 3.7 / 下 3.5 / 左 2.8 / 右 2.6 cm
    sec.page_width, sec.page_height = Cm(21.0), Cm(29.7)
    sec.top_margin, sec.bottom_margin = Cm(3.7), Cm(3.5)
    sec.left_margin, sec.right_margin = Cm(2.8), Cm(2.6)

    # Normal 样式兜底为三号仿宋
    normal = doc.styles["Normal"]
    normal.font.name = FONT_LATIN
    normal.font.size = Pt(SZ_BODY)
    rFonts = normal.element.get_or_add_rPr().get_or_add_rFonts()
    rFonts.set(qn("w:ascii"), FONT_LATIN)
    rFonts.set(qn("w:hAnsi"), FONT_LATIN)
    rFonts.set(qn("w:eastAsia"), FONT_BODY_EA)

    # 标题（居中，二号小标宋）
    add_para(doc, title, align=WD_ALIGN_PARAGRAPH.CENTER, size=SZ_TITLE,
             ea=FONT_TITLE_EA, line=LINE_TITLE, space_after=28)

    # 称谓（顶格，值 + 全角冒号）
    for name, req, role, val, missing, _line in entries:
        if role != "salutation":
            continue
        if missing:
            text = "{}：{}".format(name, PLACEHOLDER)
        else:
            text = "{}：".format(render_value(val))
        add_para(doc, text)

    # 正文段落（首行缩进两字，「字段名：值」全角冒号；必填缺失 ____ 占位，选填缺失省略）
    for name, req, role, val, missing, _line in entries:
        if role != "body":
            continue
        if missing and not req:
            continue
        text = "{}：{}".format(name, PLACEHOLDER if missing else render_value(val))
        add_para(doc, text, indent_chars=2)

    closing = TEMPLATES[template]["closing"]
    if closing:
        add_para(doc, closing, indent_chars=2)

    # 落款（右对齐，右空四字；与正文之间空一行）
    sig_entries = [e for e in entries if e[2] == "signature"]
    for i, (name, req, role, val, missing, _line) in enumerate(sig_entries):
        text = "{}：{}".format(name, PLACEHOLDER if missing else render_value(val))
        add_para(doc, text, align=WD_ALIGN_PARAGRAPH.RIGHT,
                 right_indent_pt=SZ_BODY * 4, space_before=28 if i == 0 else 0)

    # 附注（成文日期下一行，居左空二字，圆括号）
    note_items = [(e[0], e[3]) for e in entries if e[2] == "note" and not e[4]]
    if note_items:
        text = "（{}）".format("，".join("{}：{}".format(n, render_value(v)) for n, v in note_items))
        add_para(doc, text, indent_chars=2)

    doc.core_properties.title = title
    doc.core_properties.author = ""

    docx_path = os.path.join(outdir, "文书.docx")
    doc.save(docx_path)
    return docx_path


def main():
    ap = argparse.ArgumentParser(
        description="四类公文模板生成器（周报/请示函/会议通知/工作总结），生成 文书.docx 与 fields.json")
    ap.add_argument("--template", required=True, choices=sorted(TEMPLATES),
                    help="模板类型：周报|请示函|会议通知|工作总结")
    ap.add_argument("--data", required=True, help="数据 JSON 文件路径")
    ap.add_argument("--outdir", required=True, help="输出目录（在其下生成 文书.docx 与 fields.json）")
    args = ap.parse_args()

    try:
        with open(args.data, "r", encoding="utf-8") as f:
            data = json.load(f)
    except (OSError, json.JSONDecodeError) as e:
        print("读取数据失败：{}".format(e), file=sys.stderr)
        return 2
    if not isinstance(data, dict):
        print("数据 JSON 顶层必须是对象（字段名 → 值）", file=sys.stderr)
        return 2

    spec = TEMPLATES[args.template]
    spec_names = [f[0] for f in spec["fields"]]

    entries = []
    for name, req, role in spec["fields"]:
        val = data.get(name)
        entries.append([name, req, role, val, is_missing(val), None])
    # 数据中多出的字段：追加为选填正文段，同样如实记账
    for name in data:
        if name not in spec_names:
            val = data[name]
            entries.append([name, False, "body", val, is_missing(val), None])

    # 标题（缺失部分以 ____ 代入，不编造）
    class _G(dict):
        def __missing__(self, k):
            return PLACEHOLDER

    g = _G({e[0]: render_value(e[3]) for e in entries if not e[4]})
    title = spec["title_pattern"].format_map(g)
    title_parts = set(re.findall(r"\{(.+?)\}", spec["title_pattern"]))

    # 逐字段渲染文本（用于 fields.json 的「渲染方式」记账）
    fields = []
    for name, req, role, val, missing, _ in entries:
        rv = None if missing else render_value(val)
        if missing:
            if role == "body":
                where = "正文段落「{}：{}」".format(name, PLACEHOLDER) if req else "选填缺失，正文省略"
            elif role == "salutation":
                where = "称谓顶格「{}：{}」".format(name, PLACEHOLDER)
            elif role == "signature":
                where = "落款右对齐「{}：{}」".format(name, PLACEHOLDER)
            else:
                where = "选填缺失，附注省略"
        else:
            if role == "salutation":
                where = "称谓顶格「{}：」".format(rv)
            elif role == "signature":
                where = "落款右对齐「{}：{}」".format(name, rv)
            elif role == "note":
                where = "成文日期下附注「{}：{}」".format(name, rv)
            else:
                where = "正文段落「{}：{}」".format(name, rv)
        if name in title_parts:
            where += "；并用于标题《{}》".format(title)
        fields.append({
            "name": name,
            "required": bool(req),
            "status": "missing" if missing else "filled",
            "value": None if missing else val,
            "rendered_as": where,
        })

    filled = [f["name"] for f in fields if f["status"] == "filled"]
    missing = [f["name"] for f in fields if f["status"] == "missing"]

    os.makedirs(args.outdir, exist_ok=True)
    build_docx(args.template, title, [tuple(e) for e in entries], args.outdir)

    record = {
        "title": title,
        "template": args.template,
        "source_data": args.data,
        "fields": fields,
        "filled_fields": filled,
        "missing_fields": missing,
        "summary": {
            "filled": len(filled),
            "missing": len(missing),
            "total": len(fields),
        },
    }
    with open(os.path.join(args.outdir, "fields.json"), "w", encoding="utf-8") as f:
        json.dump(record, f, ensure_ascii=False, indent=2)
        f.write("\n")

    print("[ok] template={} title={} filled={} missing={} -> {}".format(
        args.template, title, len(filled), len(missing), args.outdir))
    return 0


if __name__ == "__main__":
    sys.exit(main())
