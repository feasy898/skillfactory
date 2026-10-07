# -*- coding: utf-8 -*-
"""根据 fixtures/transcript.txt 生成 out/纪要.docx 与 out/summary.json。
所有条目内容均逐字取自转写稿，不做归纳或改写。
"""
import json
import os

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Pt

BASE = os.path.dirname(os.path.abspath(__file__))
OUT = os.path.join(BASE, "out")

TITLE = "星帆订单系统项目例会会议纪要"

DECISIONS = [
    ("赵静怡", "星帆订单系统整体进度正常，会议同意按原计划推进，里程碑保持不变"),
    ("黄雨桐", "会议同意星帆订单系统项目预算300万元，走专项列支，由市场部跟进执行"),
]

TODOS = [
    ("黄雨桐在本周五前完成复盘纪要并同步给财务部", "黄雨桐", "本周五前"),
    ("徐静怡负责把数据迁移脚本整理成7月4日归档", "徐静怡", "7月4日"),
]

RISKS = [
    ("吴志远", "财务部侧存在履约风险，去年高峰期出现过积压，今年需要提前准备预案"),
    ("赵静怡", "灰度放量方案的数据口径尚未对齐，复盘前需要补齐基线数据"),
]


def set_font(run, size=None, bold=None):
    run.font.name = "SimSun"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "SimSun")
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.font.bold = bold


def add_heading(doc, text):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(12)
    p.paragraph_format.space_after = Pt(6)
    set_font(p.add_run(text), size=14, bold=True)


def add_numbered(doc, category, items):
    for idx, (person, content) in enumerate(items, 1):
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(4)
        set_font(p.add_run("%d. 【%s · %s】%s" % (idx, category, person, content)), size=11)


def main():
    os.makedirs(OUT, exist_ok=True)
    doc = Document()

    title_p = doc.add_paragraph()
    title_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_font(title_p.add_run(TITLE), size=18, bold=True)

    add_heading(doc, "一、决议事项")
    add_numbered(doc, "决议事项", DECISIONS)

    add_heading(doc, "二、待办事项")
    table = doc.add_table(rows=1, cols=3)
    table.style = "Table Grid"
    for cell, name in zip(table.rows[0].cells, ["事项", "负责人", "期限"]):
        cell.text = ""
        run = cell.paragraphs[0].add_run(name)
        set_font(run, size=11, bold=True)
    for item, owner, due in TODOS:
        cells = table.add_row().cells
        for cell, value in zip(cells, [item, owner, due]):
            cell.text = ""
            set_font(cell.paragraphs[0].add_run(value), size=11)

    add_heading(doc, "三、风险与关注")
    add_numbered(doc, "风险与关注", RISKS)

    doc.save(os.path.join(OUT, "纪要.docx"))

    summary = {
        "决议事项": len(DECISIONS),
        "待办事项": len(TODOS),
        "风险与关注": len(RISKS),
        "合计": len(DECISIONS) + len(TODOS) + len(RISKS),
    }
    with open(os.path.join(OUT, "summary.json"), "w", encoding="utf-8") as f:
        json.dump(summary, f, ensure_ascii=False)

    print(summary)


if __name__ == "__main__":
    main()