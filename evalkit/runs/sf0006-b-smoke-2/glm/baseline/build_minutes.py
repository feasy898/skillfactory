#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""从 fixtures/transcript.txt 生成 out/纪要.docx 与 out/summary.json。

条目内容全部逐字取自转写稿，构建时强制校验（非逐字即报错）。
"""
import json
import os
import re

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Cm, Pt

BASE = os.path.dirname(os.path.abspath(__file__))
OUT_DIR = os.path.join(BASE, "out")
TRANSCRIPT = os.path.join(BASE, "fixtures", "transcript.txt")

with open(TRANSCRIPT, encoding="utf-8") as f:
    raw_lines = [ln.rstrip("\n") for ln in f if ln.strip()]

utterances = []  # (时间, 说话人, 内容)
for ln in raw_lines:
    m = re.match(r"^\[(\d{2}:\d{2})]\s*(.+?):\s*(.+)$", ln)
    if not m:
        raise ValueError(f"无法解析的行: {ln!r}")
    utterances.append((m.group(1), m.group(2), m.group(3).strip()))

speakers = {sp for _, sp, _ in utterances}
full_text = "".join(tx for _, _, tx in utterances)

def require_verbatim(s, what):
    if s not in full_text:
        raise AssertionError(f"{what} 未逐字见于转写稿: {s!r}")

# ---- 从转写稿逐字摘取的条目（说话人=该句发言人） ----
decisions = [
    ("赵静怡", "星帆订单系统整体进度正常，会议同意按原计划推进，里程碑保持不变"),
    ("黄雨桐", "会议同意星帆订单系统项目预算300万元，走专项列支，由市场部跟进执行"),
]
todos = [
    ("完成复盘纪要并同步给财务部", "黄雨桐", "本周五前"),
    ("把数据迁移脚本整理成7月4日归档", "徐静怡", "7月4日"),
]
risks = [
    ("吴志远", "财务部侧存在履约风险，去年高峰期出现过积压，今年需要提前准备预案"),
    ("赵静怡", "灰度放量方案的数据口径尚未对齐，复盘前需要补齐基线数据"),
]

for person, content in decisions + risks:
    if person not in speakers:
        raise AssertionError(f"相关人不是转写稿说话人: {person}")
    require_verbatim(content, "条目内容")
for item, owner, deadline in todos:
    if owner not in speakers:
        raise AssertionError(f"负责人不是转写稿说话人: {owner}")
    require_verbatim(item, "待办事项")
    if deadline != "未明确":
        require_verbatim(deadline, "期限")

# ---- 生成 docx ----
doc = Document()

normal = doc.styles["Normal"]
normal.font.name = "Times New Roman"
normal.font.size = Pt(12)
normal.element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")

def styled_run(par, text, east="宋体", size=None, bold=None):
    run = par.add_run(text)
    run.font.name = "Times New Roman"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), east)
    if size is not None:
        run.font.size = size
    if bold is not None:
        run.bold = bold
    return run

# 标题（首段）
title_par = doc.add_paragraph()
title_par.alignment = WD_ALIGN_PARAGRAPH.CENTER
styled_run(title_par, "星帆订单系统项目例会会议纪要", east="黑体", size=Pt(18), bold=True)

def add_heading(text):
    par = doc.add_paragraph()
    par.paragraph_format.space_before = Pt(10)
    par.paragraph_format.space_after = Pt(6)
    styled_run(par, text, east="黑体", size=Pt(15), bold=True)

def add_item(idx, category, person, content):
    par = doc.add_paragraph()
    par.paragraph_format.space_after = Pt(4)
    styled_run(par, f"{idx}. 【{category} · {person}】{content}")

# 一、决议事项
add_heading("一、决议事项")
for i, (person, content) in enumerate(decisions, 1):
    add_item(i, "决议事项", person, content)

# 二、待办事项（表格）
add_heading("二、待办事项")
table = doc.add_table(rows=1 + len(todos), cols=3)
table.style = "Table Grid"
for cell, text, width in zip(
    table.rows[0].cells, ("事项", "负责人", "期限"), (Cm(9), Cm(3.2), Cm(3.2))
):
    cell.width = width
    par = cell.paragraphs[0]
    styled_run(par, text, east="黑体", bold=True)
for r, (item, owner, deadline) in enumerate(todos, 1):
    for cell, text, width in zip(
        table.rows[r].cells, (item, owner, deadline), (Cm(9), Cm(3.2), Cm(3.2))
    ):
        cell.width = width
        par = cell.paragraphs[0]
        styled_run(par, text)

# 三、风险与关注
add_heading("三、风险与关注")
for i, (person, content) in enumerate(risks, 1):
    add_item(i, "风险与关注", person, content)

os.makedirs(OUT_DIR, exist_ok=True)
doc.save(os.path.join(OUT_DIR, "纪要.docx"))

# ---- 生成 summary.json ----
summary = {
    "决议事项": len(decisions),
    "待办事项": len(todos),
    "风险与关注": len(risks),
}
summary["合计"] = sum(summary.values())
with open(os.path.join(OUT_DIR, "summary.json"), "w", encoding="utf-8") as f:
    json.dump(summary, f, ensure_ascii=False)
    f.write("\n")

print("utterances:", len(utterances), "speakers:", sorted(speakers))
print("summary:", summary)
print("written:", os.path.join(OUT_DIR, "纪要.docx"), "and summary.json")
