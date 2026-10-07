#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""独立复核 out/纪要.docx 与 out/summary.json 是否满足 brief.md 的全部判据。"""
import json
import os
import re

from docx import Document

BASE = os.path.dirname(os.path.abspath(__file__))
OUT_DIR = os.path.join(BASE, "out")

with open(os.path.join(BASE, "fixtures", "transcript.txt"), encoding="utf-8") as f:
    raw = [ln.rstrip("\n") for ln in f if ln.strip()]
speakers, texts = set(), []
for ln in raw:
    m = re.match(r"^\[\d{2}:\d{2}]\s*(.+?):\s*(.+)$", ln)
    speakers.add(m.group(1))
    texts.append(m.group(2).strip())
full_text = "".join(texts)

doc = Document(os.path.join(OUT_DIR, "纪要.docx"))
paras = [p.text for p in doc.paragraphs]
print(f"段落数={len(paras)}")

# 1) 首段为标题且含会议主题
title = paras[0]
assert "星帆订单系统" in title, f"标题未含会议主题: {title!r}"
print(f"  PASS 首段标题: {title!r}")

# 2) 三个节标题逐字、唯一、按序
HEADINGS = ["一、决议事项", "二、待办事项", "三、风险与关注"]
for h in HEADINGS:
    assert paras.count(h) == 1, f"节标题缺失或不唯一: {h!r}（出现 {paras.count(h)} 次）"
found = [t for t in paras if t in HEADINGS]
assert found == HEADINGS, found
print("  PASS 三节标题逐字且按序")

def body_between(start, stop=None):
    """start 节标题之后、stop 节标题之前的非空段落。"""
    i = paras.index(start) + 1
    j = paras.index(stop) if stop else len(paras)
    return [t for t in paras[i:j] if t.strip()]

decisions = body_between("一、决议事项", "二、待办事项")
risks = body_between("三、风险与关注")

# 3) 决议事项 / 风险与关注 条目格式与溯源
item_pat = re.compile(r"^(\d+)\. 【([^·]+) · ([^】]+)】(.+)$")
for section, items in (("决议事项", decisions), ("风险与关注", risks)):
    assert items, f"{section} 无条目"
    for i, t in enumerate(items, 1):
        m = item_pat.match(t)
        assert m, f"{section} 条目格式不符: {t!r}"
        num, cat, person, content = m.groups()
        assert num == str(i), f"编号不连续: {t!r}"
        assert cat == section, f"类别应为「{section}」: {t!r}"
        assert person in speakers, f"相关人不是转写稿说话人: {person}"
        assert content in full_text or content + "。" in full_text, \
            f"内容非逐字取自转写稿: {content!r}"
    print(f"  PASS {section}: {len(items)} 条，编号/类别/相关人/逐字内容全部合规")

# 4) 待办表格
tables = doc.tables
assert len(tables) == 1, f"应恰有 1 张表，实为 {len(tables)}"
tb = tables[0]
header = [c.text for c in tb.rows[0].cells]
assert header == ["事项", "负责人", "期限"], header
print(f"  PASS 表头三列逐字: {header}")
todo_count = 0
for row in tb.rows[1:]:
    item, owner, deadline = (c.text for c in row.cells)
    todo_count += 1
    assert item in full_text, f"事项无逐字出处: {item!r}"
    assert owner in speakers, f"负责人不是转写稿说话人: {owner}"
    assert deadline == "未明确" or deadline in full_text, f"期限无出处: {deadline!r}"
    print(f"  PASS 待办 {todo_count}: {item!r} | {owner} | {deadline}")

# 5) summary.json：键名逐字、计数与 docx 实际条目数一致
with open(os.path.join(OUT_DIR, "summary.json"), encoding="utf-8") as f:
    summary = json.load(f)
assert list(summary) == ["决议事项", "待办事项", "风险与关注", "合计"], list(summary)
expected = {"决议事项": len(decisions), "待办事项": todo_count, "风险与关注": len(risks)}
expected["合计"] = sum(expected.values())
assert summary == expected, f"计数不一致: {summary} vs {expected}"
print(f"  PASS summary.json 键名与计数一致: {json.dumps(summary, ensure_ascii=False)}")

print("\nALL CHECKS PASSED")
