#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""runs/office-retest-20261009/x_judges.py — office 重测裁判面板驱动（盲态、异族 ≥2、位置轮换、§3.4 抽检）。

移植自 sf0008-k5b/s8_judges.py，按本轮登记（round-plan.json）适配：
  面板：被测 glm → [step-router-v1, Minimax-M3.1-Flash-Preview]；被测 minimax → [step-router-v1, GLM-5.3]
  （≥2 族互裁；与被测族交集为空）。step=模型 API 直调；glm/minimax=headless M4/M2。
  盲态：产物匿名化（P1..Pn 洗牌）+ docx→事实预抽取 + 绝对路径/挂载痕迹漂白；映射仅存
  records/judge-staging/。位置轮换：槽序按任务序奇偶互换。
双判分协议（judge_prompt_version=S8-JUDGE-v2-OFXRETEST-1，随轮落盘）：
  A. 原始 5 ab 任务：golden.json rubric（冻结原文）逐条 0/0.5/1/1.5/2 判分，总分=
     round(10×Σ项÷(2×项数),1)（照 K5B 口径）。
     ab5 特例（planner 裁决 9827318）：③④ 由裁判在臂外对交付产物重放计分——暂存目录已附
     _replay_facts.json（x_ab5_replay.py 机械产出：同命令双跑 sha256 对拍、runner 红绿复检、
     交付单元对拍），①②⑤ 仍按原口径（产物自含 + run_log 留痕辅助）。
  B. 增订 8 新题：rubric.md 加权维度判分——每维度按档位赋分（满分档=1.0、零分档=0.0、
     中档=rubric 明示分，未明示取 0.5），加权分=round(10×Σ(维度分×权重)÷Σ权重,1)；
     通过线=task.json.pass_line。verdict 须含 dims 明细。
§3.4：每任务抽 1 份产物同裁判匿名复评；|Δ|>2 → 换面板内另一异族裁判重评全任务；
  仍>2 → task unjudgeable（fail-closed）。
用法：python x_judges.py [--task <task_id>]
退出码：0 全任务判分落盘 / 1 有任务 unjudgeable 或装置失败（如实，不放宽）。
"""

import json
import os
import random
import re
import subprocess
import sys
import zlib

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
import x_judge_lib as J  # noqa: E402

ROOT = J.ROOT
RUN = J.RUN
PANEL = {"glm": ["step-router-v1", "Minimax-M3.1-Flash-Preview"],
         "minimax": ["step-router-v1", "GLM-5.3"]}
JUDGE_OF_MODEL = {"step-router-v1": "step", "Minimax-M3.1-Flash-Preview": "minimax",
                  "GLM-5.3": "glm"}

TASK_ORDER = ["ab1-standard-four-templates", "ab2-boundary-missing-required",
              "ab3-empty-and-string-lists", "ab4-invalid-input-rejection",
              "ab5-determinism-layout-selfeval",
              "t01", "t02", "t03", "t04", "t05", "t06", "t07", "t08"]
AB_SET = set(TASK_ORDER[:5])
MERGED = os.path.join(ROOT, "planning", "taskset-redesign", "merged", "office-templates-skill")

SCORING_RULE_AB = """评分规则（S8-JUDGE-v2-OFXRETEST-1 / ab 题协议，照 K5B）：
- 对每条 rubric 项独立打分：0=完全不符 / 0.5 / 1=部分符合 / 1.5 / 2=完全符合。
- 每条打分必须附实测证据（引用暂存文件中的具体事实/原文/数字）；无证据的分数无效。证据必须精简（每条 rubric 项 ≤150 字），优先保证 verdict JSON 完整输出。
- ab5 协议注记（planner 裁决 9827318）：第 3、4 条（runner 对拍、红绿自校验）以暂存目录所附 _replay_facts.json（裁判侧臂外机械重放）为计分依据；第 1、2、5 条按产物自含与 run_log 留痕辅助证据正常判。
- 产物总分 total = round(10 × Σ项分 ÷ (2×项数), 1)，区间 [0,10]。
- 你不知道也不得推测产物的来源臂、模型或顺序；只按 rubric 与证据判分。
- verdict.json 结构（紧凑，必须遵守）：{"scores": {"P1": {"total": x.x, "evidence": "≤150 字逐项证据摘要"}}, ...}（每个产品一个键；不要 items 数组、不要其他顶层键）。"""

SCORING_RULE_NEW = """评分规则（S8-JUDGE-v2-OFXRETEST-1 / 增订题加权维度协议）：
- rubric 表每行为一个维度（维度名/权重/满分档/中档/零分档）。对每维度独立定档并赋分：
  满分档=1.0；零分档=0.0；中档=该维度分（取 rubric 中明示的分值，如 0.4/0.5；rubric 未明示数值时取 0.5）。
- 定档必须附实测证据（引用暂存文件中的具体事实/原文/数字，逐条 ≤150 字）；无证据的档位无效。
- 产物加权分 total = round(10 × Σ(维度分 × 权重) ÷ Σ权重, 1)，区间 [0,10]。这是唯一总分，不再另设项分平均。
- 你不知道也不得推测产物的来源臂、模型或顺序；只按 rubric 与证据判分；实现路径与参照不同不扣分（端态原则），除非 rubric 显式考格式/结构。
- verdict.json 结构（紧凑，必须遵守）：{"scores": {"P1": {"total": x.x, "dims": {"<维度名>": 分数, ...}, "evidence": "≤150 字逐维证据摘要"}}, ...}（dims 键名须与 rubric 维度名逐字一致；除 scores 外不要其他顶层键）。"""


def redact_text(s, arm_root):
    s = s.replace(arm_root, "<ARMROOT>")
    s = s.replace(ROOT, "<REPOROOT>")
    s = re.sub(r"\.zcode/skills/[A-Za-z0-9_\-]+", "<SKILL_MOUNT>", s)
    s = re.sub(r"\.zcode\\\\skills\\\\[A-Za-z0-9_\-]+", "<SKILL_MOUNT>", s)
    for key in ("arm", "engine", "run_id", "runId", "side", "arm_label", "session_id",
                "sessionId", "skill_mount", "asset"):
        s = re.sub(r'"%s"\s*:\s*"[^"]*"' % key, '"%s": "<redacted>"' % key, s)
        s = re.sub(r"'%s'\s*:\s*'[^']*'" % key, "'%s': '<redacted>'" % key, s)
    return s


def docx_facts(path):
    """python-docx 确定性抽取：段落文本+对齐+缩进+字体/字号/行距 + 表格文本。"""
    from docx import Document
    from docx.oxml.ns import qn
    doc = Document(path)
    lines = []
    for i, p in enumerate(doc.paragraphs):
        if not p.text.strip() and p._p.xpath(".//w:drawing") == []:
            continue
        align = p.alignment
        pf = p.paragraph_format
        fli = pf.first_line_indent
        fli_chars = None
        pPr = p._p.pPr
        if pPr is not None:
            ind = pPr.find(qn("w:ind"))
            if ind is not None:
                fli_chars = ind.get(qn("w:firstLineChars"))
        fonts = []
        for r in p.runs[:6]:
            sz = r.font.size.pt if r.font.size else None
            east = None
            rPr = r._r.rPr
            if rPr is not None:
                rf = rPr.find(qn("w:rFonts"))
                if rf is not None:
                    east = rf.get(qn("w:eastAsia"))
            spacing = pf.line_spacing
            fonts.append("%s/%s/%s" % (r.text[:14], east or r.font.name, sz))
        lines.append("P%03d align=%s firstLineChars=%s firstLineIndent=%s spacing=%s runs=[%s] text=%s"
                     % (i, align, fli_chars, fli, pf.line_spacing, " | ".join(fonts), p.text))
    for ti, tb in enumerate(doc.tables, 1):
        lines.append("TABLE%d %d行×%d列:" % (ti, len(tb.rows), len(tb.columns)))
        for row in tb.rows:
            lines.append("  " + " ┃ ".join(c.text.replace("\n", "⏎") for c in row.cells))
    return "\n".join(lines)


def stage_products(task, unit_dirs, staging, seed):
    """收集产物→漂白→洗牌匿名化。返回 (staged, mapping, meta)。"""
    rng = random.Random(seed)
    items = []
    for rep, ad in unit_dirs:
        files = {}
        out = os.path.join(ad, "out")
        if os.path.isdir(out):
            for dp, _dn, fns in os.walk(out):
                for fn in fns:
                    src = os.path.join(dp, fn)
                    files[os.path.relpath(src, out)] = src
        rl = os.path.join(ad, "run_log.json")
        if os.path.isfile(rl):
            files["run_log.json"] = rl
        items.append({"rep": rep, "arm": os.path.basename(os.path.normpath(ad)),
                      "dir": ad, "files": files})
    rng.shuffle(items)
    os.makedirs(staging, exist_ok=True)
    mapping, staged, meta = {}, [], []
    for i, it in enumerate(items, 1):
        pid = "P%d" % i
        mapping[pid] = {"rep": it["rep"], "arm": it["arm"]}
        pdir = os.path.join(staging, pid)
        os.makedirs(pdir, exist_ok=True)
        names = []
        for rel, src in sorted(it["files"].items()):
            ext = os.path.splitext(rel)[1].lower()
            base = rel.replace("/", "__")
            dst = os.path.join(pdir, base)
            if ext == ".docx":
                facts = docx_facts(src)
                with open(dst + ".extract.txt", "w", encoding="utf-8", newline="\n") as f:
                    f.write("EXTRACT OF %s (python-docx 确定性抽取，段落含对齐/缩进/字体/字号/行距事实)\n%s\n"
                            % (base, facts))
                names.append(base + ".extract.txt")
            elif ext in (".json", ".md", ".txt", ".log"):
                raw = open(src, encoding="utf-8", errors="replace").read()
                with open(dst, "w", encoding="utf-8", newline="\n") as f:
                    f.write(redact_text(raw, it["dir"]))
                names.append(base)
            else:
                import shutil
                shutil.copyfile(src, dst)
                names.append(base)
        if task == "ab5-determinism-layout-selfeval":
            # 裁决：③④ 裁判侧臂外重放——重放事实随暂存进判据面（映射仍保密）
            fj = os.path.join(RUN, "records", "ab5-replay",
                              "%s-rep%d-%s.facts.json" % (task, it["rep"], it["arm"]))
            if os.path.isfile(fj):
                raw = open(fj, encoding="utf-8").read()
                with open(os.path.join(pdir, "_replay_facts.json"), "w",
                          encoding="utf-8", newline="\n") as f:
                    f.write(redact_text(raw, it["dir"]))
                names.append("_replay_facts.json")
        staged.append({"pid": pid, "files": names})
        meta.append({"pid": pid, "n_files": len(names)})
    return staged, mapping, meta


def rubric_ab(task):
    d = json.load(open(os.path.join(ROOT, "dist", "office-templates-skill", "eval/golden.json"),
                       encoding="utf-8"))
    return next(t for t in d["ab_tasks"] if t["id"] == task)["rubric"]


def rubric_new(task):
    return open(os.path.join(MERGED, task, "rubric.md"), encoding="utf-8").read()


def build_brief(task, staged, is_ab, pass_line=None):
    if is_ab:
        lines = ["# office 重测盲评判分任务（ab 题：%s）" % task, "",
                 "你是盲评裁判。以下为本任务的匿名化产物（P1..Pn，顺序随机，来源对你保密）。", "",
                 "## rubric（golden.json 冻结原文，逐条判分）"]
        for i, r in enumerate(rubric_ab(task), 1):
            lines.append("%d. %s" % (i, r))
        lines += ["", SCORING_RULE_AB]
    else:
        tj = json.load(open(os.path.join(MERGED, task, "task.json"), encoding="utf-8"))
        notes_p = os.path.join(MERGED, task, "judge-notes.md")
        notes = open(notes_p, encoding="utf-8").read().rstrip() if os.path.isfile(notes_p) else ""
        # 判据面修正（2026-10-09）：rubric.md 判官执行节明确「判据与期望值见 judge-notes.md（判官专用）」，
        # 首批 t01-t04 brief 漏嵌 notes 致双裁在档位边界发散（t04 P3/P5/P6 7.9 vs 10.0 在案）；
        # 本修正补齐判官专用 notes，rubric/评分规则文本零改动，受影响任务全量重判（留痕 incidents）。
        lines = ["# office 重测盲评判分任务（增订题：%s，通过线 %s）" % (task, tj.get("pass_line")),
                 "", "你是盲评裁判。以下为本任务的匿名化产物（P1..Pn，顺序随机，来源对你保密）。",
                 "", "## rubric（冻结原文，加权维度判分）", "", rubric_new(task).rstrip()]
        if notes:
            lines += ["", "## judge-notes（判官专用：判挂点与期望值，判分必照此执行）", "", notes]
        lines += ["", SCORING_RULE_NEW,
                  "", "（题目指令原文备查，判分只按 rubric 与 judge-notes）：", "```", tj["instruction"], "```"]
    lines += ["", "## 产物清单"]
    for s in staged:
        lines.append("- %s: %s" % (s["pid"], ", ".join(s["files"])))
    lines += ["", "请阅读同目录各 P* 子目录内的文件后，把 verdict.json 写入当前工作目录（仅输出该文件）。"]
    return "\n".join(lines)


def parse_verdict_json(text):
    """提取含 "scores" 的最后一个平衡 JSON 对象（容忍围栏/推理散文）。"""
    best = None
    for m in re.finditer(r'"scores"', text):
        i = text.rfind("{", 0, m.start())
        if i < 0:
            continue
        depth = 0
        for j in range(i, len(text)):
            if text[j] == "{":
                depth += 1
            elif text[j] == "}":
                depth -= 1
                if depth == 0:
                    cand = text[i:j + 1]
                    try:
                        obj = json.loads(cand)
                        if isinstance(obj, dict) and isinstance(obj.get("scores"), dict):
                            best = obj
                    except Exception:
                        pass
                    break
    return best


STAGING_ROOTS = {}


def judge_call(judge, workdir, brief, staged, task_label):
    """单裁判单次调用。step=HTTP（产物内容内联）；glm/minimax=headless CLI。返回 (rec, verdict)。"""
    import time
    chan = JUDGE_OF_MODEL[judge]
    slot = J.budget_judge({"judge": judge, "channel": chan,
                           "task": task_label, "judge_prompt_version": J.JUDGE_PROMPT_VERSION,
                           "ts_intent": time.strftime("%H:%M:%S"), "status": "pending"})
    try:
        if chan == "step":
            parts = ["【输出纪律（最高优先级）】你是评分器。禁止展示任何思考、草稿、advisor 对话或分步分析；你的回复必须直接以 { 开头、以 } 结尾，内容仅为最终 verdict JSON。", "", brief, "", "# 产物内容（暂存文件全文）"]
            for s in staged:
                src_base = STAGING_ROOTS[task_label]
                for fn in s["files"]:
                    cand = [os.path.join(src_base, s["pid"], fn), os.path.join(src_base, fn)]
                    p = next((c for c in cand if os.path.isfile(c)), None)
                    if p is None:
                        continue
                    parts.append("=== %s/%s ===\n%s" % (s["pid"], fn,
                                 open(p, encoding="utf-8", errors="replace").read()[:24000]))
            r = J.step_chat([{"role": "user", "content": "\n".join(
                parts + ["", '现在直接输出最终评分 JSON：直接以 { 开头、以 } 结尾，最外层恰有一个 "scores" 键；每个产品一个键，键含 "total" 与 "evidence"；除该 JSON 外不得输出任何文字，不得重复输出 JSON。'])}],
                max_tokens=32000)
            # 通道修复注记（2026-10-09）：原 assistant_prefill 续写式拼接在 step-router-v1 上
            # 实测产生「prefill+模型重启输出」双层嵌套 JSON 致解析 None（judge-work/*/response.txt
            # 在案）；改为无 prefill + 强化输出纪律，评分规则与 brief 文本零改动。
            v = parse_verdict_json(r["text"])
            rec = {"exit_code": 0 if v else 1, "wall_ms": r["wall_ms"], "session_id": r["id"],
                   "model_id_observed": r["model"], "channel": "http-api",
                   "input_tokens": r["prompt_tokens"], "output_tokens": r["completion_tokens"]}
            with open(os.path.join(workdir, "response.txt"), "w", encoding="utf-8",
                      newline="\n") as f:
                f.write(r["text"])
        else:
            with open(os.path.join(workdir, "brief.md"), "w", encoding="utf-8",
                      newline="\n") as f:
                f.write(brief + "\n")
            for s in staged:
                src_pdir = os.path.join(STAGING_ROOTS[task_label], s["pid"])
                dst_pdir = os.path.join(workdir, s["pid"])
                if os.path.isdir(src_pdir):
                    import shutil as sh
                    sh.copytree(src_pdir, dst_pdir, dirs_exist_ok=True)
            rec = J.cli_judge(workdir, chan)
            vp = os.path.join(workdir, "verdict.json")
            v = json.load(open(vp, encoding="utf-8")) if os.path.isfile(vp) else None
            with open(os.path.join(workdir, "transcript_stdout.txt"), "w", encoding="utf-8",
                      newline="\n") as f:
                f.write(rec["_raw_stdout"])
        J.budget_judge_update(slot, status="done" if v else "no-verdict",
                              exit_code=rec["exit_code"], wall_ms=rec["wall_ms"],
                              session_id=rec.get("session_id"),
                              tokens={"input": rec.get("input_tokens"),
                                      "output": rec.get("output_tokens")},
                              model_observed=rec.get("model_id_observed"))
        return rec, v
    except Exception as exc:
        J.budget_judge_update(slot, status="failed", exit_code=1,
                              error=type(exc).__name__ + ":" + str(exc)[:300])
        raise


def task_unit_dirs(task, fam, reps):
    out = []
    for rep in reps:
        ad = os.path.join(RUN, fam, "office-templates-%s-rep%d" % (task, rep))
        out.append((rep, os.path.join(ad, "baseline")))
        out.append((rep, os.path.join(ad, "treatment")))
    return out


def fam_of_task(task):
    return "glm" if TASK_ORDER.index(task) % 2 == 0 else "minimax"  # 1-based 奇=glm


def run_ab5_replays(reps):
    """ab5 ③④ 裁判侧重放（零模型）：每 rep 每臂跑 x_ab5_replay.py，落 records/ab5-replay/。"""
    task = "ab5-determinism-layout-selfeval"
    fam = fam_of_task(task)
    outd = os.path.join(RUN, "records", "ab5-replay")
    os.makedirs(outd, exist_ok=True)
    for rep in reps:
        for arm in ("baseline", "treatment"):
            ad = os.path.join(RUN, fam, "office-templates-%s-rep%d" % (task, rep), arm)
            outj = os.path.join(outd, "%s-rep%d-%s.facts.json" % (task, rep, arm))
            if os.path.isfile(outj):
                continue
            r = subprocess.run([sys.executable, os.path.join(RUN, "x_ab5_replay.py"), ad, outj],
                               capture_output=True, timeout=600)
            if r.returncode != 0:
                raise RuntimeError("ab5 replay 失败 rep%d/%s: %s"
                                   % (rep, arm, r.stderr.decode("utf-8", "replace")[-400:]))


def main():
    only = None
    if "--task" in sys.argv:
        only = sys.argv[sys.argv.index("--task") + 1]
    reps = [1, 2, 3]
    jp = os.path.join(RUN, "records", "judging", "office-templates.json")
    results = json.load(open(jp, encoding="utf-8")) if os.path.isfile(jp) else {}
    for ti, task in enumerate(TASK_ORDER, 1):
        if only and task != only:
            continue
        if isinstance(results.get(task), dict) and results[task].get("status") in ("judged", "unjudgeable"):
            print("SKIP judged %s" % task, flush=True)
            continue
        try:
            fam = fam_of_task(task)
            is_ab = task in AB_SET
            if task == "ab5-determinism-layout-selfeval":
                run_ab5_replays(reps)
            units = task_unit_dirs(task, fam, reps)
            missing = [u for u in units if not os.path.isfile(os.path.join(u[1], "records", "arm_run.json"))]
            if missing:
                results[task] = {"status": "infra_error", "missing": [m for m in missing]}
                continue
            staging = os.path.join(RUN, "judge-work", "_staging", task)
            STAGING_ROOTS[task] = staging
            seed = zlib.crc32(("office-retest-20261009|%s" % task).encode()) & 0xffffff
            staged, mapping, meta = stage_products(task, units, staging, seed)
            J.write_json(os.path.join(RUN, "records", "judge-staging", "%s.mapping.json" % task),
                         {"mapping": mapping, "staged": staged, "seed": seed})
            brief = build_brief(task, staged, is_ab)
            panel = PANEL[fam][:]
            if ti % 2 == 0:            # 位置轮换：偶数任务槽序互换
                panel = panel[::-1]
            scores, verdicts, ok = {}, {}, True

            def _one_judge(judge):
                wdir = os.path.join(RUN, "judge-work", task, judge.replace("/", "_"))
                os.makedirs(wdir, exist_ok=True)
                return judge_call(judge, wdir, brief, staged, task)

            for judge in panel:
                rec, v = _one_judge(judge)
                verdicts[judge] = v
                if not v or not isinstance(v.get("scores"), dict):
                    ok = False
                    continue
                for pid, sv in v["scores"].items():
                    scores.setdefault(pid, {})[judge] = sv.get("total")

            def _scored_judges():
                return {j for d in scores.values() for j in d if d[j] is not None}

            if len(_scored_judges()) < 2:
                for judge in panel:
                    if len(_scored_judges()) >= 2:
                        break
                    if judge in _scored_judges():
                        continue
                    rec, v = _one_judge(judge)
                    verdicts[judge + "#retry"] = v
                    if v and isinstance(v.get("scores"), dict):
                        for pid, sv in v["scores"].items():
                            scores.setdefault(pid, {})[judge] = sv.get("total")
            # §3.4 抽检：按任务序轮换抽 1 份（覆盖两臂），同裁判匿名复评
            sampled_pid = staged[(ti - 1) % len(staged)]["pid"]
            judge0 = panel[0]
            wdir2 = os.path.join(RUN, "judge-work", task, judge0.replace("/", "_") + "-resample")
            os.makedirs(wdir2, exist_ok=True)
            staged2 = [{"pid": "Q1", "files": next(s["files"] for s in staged if s["pid"] == sampled_pid)}]
            brief2 = build_brief(task, staged2, is_ab)
            STAGING_ROOTS["%s-resample" % task] = os.path.join(staging, sampled_pid + "_resample")
            os.makedirs(STAGING_ROOTS["%s-resample" % task], exist_ok=True)
            import shutil as _sh
            _sh.copytree(os.path.join(staging, sampled_pid),
                         STAGING_ROOTS["%s-resample" % task], dirs_exist_ok=True)
            rec2, v2 = judge_call(judge0, wdir2, brief2, staged2, "%s-resample" % task)
            first = scores.get(sampled_pid, {}).get(judge0)
            second = v2.get("scores", {}).get("Q1", {}).get("total") if v2 else None
            action, jc_delta = "pass", (abs(first - second) if (first is not None and second is not None) else None)
            final_scores = {pid: (sum([s for s in d.values() if s is not None]) /
                                  len([s for s in d.values() if s is not None])
                                  if d and any(s is not None for s in d.values()) else None)
                            for pid, d in scores.items()}
            if jc_delta is not None and jc_delta > 2:
                action = "re-judged"
                judge1 = panel[1]
                if verdicts.get(judge1) and isinstance(verdicts[judge1].get("scores"), dict):
                    news = {pid: sv.get("total") for pid, sv in verdicts[judge1]["scores"].items()}
                    final_scores = {pid: news.get(pid) for pid in final_scores}
                    if news.get(sampled_pid) is not None and first is not None \
                            and abs(news[sampled_pid] - first) > 2:
                        action = "unjudgeable"
                else:
                    action = "unjudgeable"
            J.write_json(os.path.join(RUN, "records", "judging", "%s.json" % task),
                         {"tested_family": fam, "panel": panel,
                          "panel_rotation": "task_idx=%d parity=%s" % (ti, ti % 2),
                          "verdicts": {k: v for k, v in verdicts.items()},
                          "per_judge_scores": scores, "final_scores": final_scores,
                          "judge_consistency": {"task": task, "resample_of": sampled_pid,
                                                "first_score": first, "second_score": second,
                                                "delta": jc_delta, "action": action},
                          "judge_prompt_version": J.JUDGE_PROMPT_VERSION,
                          "protocol": "ab-item" if is_ab else "weighted-dims"})
            results[task] = {"status": "judged" if action != "unjudgeable" else "unjudgeable",
                             "tested_family": fam, "panel": panel,
                             "final_scores": final_scores, "mapping": mapping,
                             "judge_consistency": {"task": task, "resample_of": sampled_pid,
                                                   "first_score": first, "second_score": second,
                                                   "delta": jc_delta, "action": action}}
            J.write_json(jp, results)
            print("JUDGED %s fam=%s action=%s" % (task, fam, action), flush=True)
        except Exception as exc:
            import traceback
            results[task] = {"status": "infra_error",
                             "error": type(exc).__name__ + ":" + str(exc)[:300],
                             "trace": traceback.format_exc()[-1500:]}
            J.write_json(jp, results)
            print("INFRA_ERROR %s: %s" % (task, exc), flush=True)
            continue
    J.write_json(os.path.join(RUN, "records", "judge-results.json"),
                 {"office-templates": results})
    bad = [1 for t in results.values() if t.get("status") in ("unjudgeable", "infra_error")]
    return 1 if bad else 0


if __name__ == "__main__":
    sys.exit(main())
