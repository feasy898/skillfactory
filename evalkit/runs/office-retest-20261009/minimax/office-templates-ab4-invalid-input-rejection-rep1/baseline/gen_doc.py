#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""文书生成入口。

用法:
    python3 gen_doc.py --template <模板名> --data <数据文件.json> [--outdir out]

行为契约:
  * 非法输入（模板名非法 / 数据文件不存在 / JSON 解析失败 / JSON 顶层非对象）
    一律以退出码 2 拒绝，stderr 输出中文错误信息，且**零产物残留**：
    不创建输出目录、不写出任何文件、不留半成品。
  * 必填字段缺失**不算非法输入**：以占位符补齐并在 stderr 给出提示，
    正常生成产物，退出码 0。
  * 合法输入退出码 0，并在 out/<模板名>/<数据文件名>/ 下生成
    文书.docx 与 fields.json 两产物。
"""

import argparse
import io
import json
import os
import sys


def _ensure_docx():
    """确保可导入 python-docx：部分环境下依赖装在用户 site-packages。"""
    try:
        import docx  # noqa: F401
        return
    except ImportError:
        pass
    import glob
    import site

    candidates = []
    try:
        candidates.extend(site.getusersitepackages())
    except Exception:
        pass
    candidates.extend(glob.glob(os.path.expanduser("~/.local/lib/python3.*/site-packages")))
    for path in candidates:
        if os.path.isdir(os.path.join(path, "docx")) and path not in sys.path:
            sys.path.append(path)
            break

# 退出码约定
EXIT_OK = 0
EXIT_INVALID_INPUT = 2

# 占位符：必填字段缺失时使用（缺失不拒绝，仅提示）
PLACEHOLDER = "待定"

# 四个标准模板及其必填字段
TEMPLATES = {
    "周报": ["部门", "填报人", "周期", "本周工作内容", "下周工作计划", "报送日期"],
    "工作总结": ["部门", "填报人", "周期", "本周工作内容", "报送日期"],
    "会议通知": ["会议名称", "会议时间", "会议地点", "参会人员", "报送日期"],
    "请示函": ["请示事由", "请示部门", "请求事项", "报送日期"],
}


def _fail(message):
    """打印中文错误信息到 stderr 并以约定退出码 2 结束。"""
    sys.stderr.write("错误：%s\n" % message)
    sys.stderr.flush()
    sys.exit(EXIT_INVALID_INPUT)


def load_data(path):
    """读取并解析数据文件。所有失败均归类为非法输入（退出码 2）。"""
    if not os.path.exists(path):
        _fail("数据文件不存在：%s，请检查 --data 指定的路径。" % path)
    if os.path.isdir(path):
        _fail("数据文件路径是一个目录而非文件：%s。" % path)
    try:
        with io.open(path, "r", encoding="utf-8") as fh:
            raw = fh.read()
    except OSError as exc:
        _fail("数据文件无法读取：%s（%s）。" % (path, exc))
    except UnicodeDecodeError:
        _fail("数据文件不是 UTF-8 编码的文本：%s。" % path)

    try:
        data = json.loads(raw)
    except ValueError as exc:
        _fail("数据文件不是合法 JSON，解析失败：%s（%s）。" % (path, exc))

    if not isinstance(data, dict):
        _fail(
            "数据文件 JSON 顶层必须是对象（键值对），实际为 %s：%s。"
            % (type(data).__name__, path)
        )
    return data


def normalize(data, required, template):
    """补齐缺失的必填字段。缺失不视为非法输入，只做提示。"""
    normalized = {}
    missing = []
    for key in required:
        if key in data and data[key] not in (None, ""):
            normalized[key] = data[key]
        else:
            normalized[key] = PLACEHOLDER
            missing.append(key)
    # 非必填但存在的数据一并保留
    for key, value in data.items():
        if key not in normalized:
            normalized[key] = value
    if missing:
        sys.stderr.write(
            "提示：模板「%s」的必填字段缺失 %s，已用占位符「%s」补齐。"
            % (template, "、".join(missing), PLACEHOLDER)
        )
        sys.stderr.write("\n")
        sys.stderr.flush()
    return normalized


def build_docx(template, fields):
    """在内存中构造 docx 字节流，构造成功后才落盘，避免半成品。"""
    _ensure_docx()
    from docx import Document
    from docx.shared import Pt

    doc = Document()
    doc.add_heading("%s" % template, level=1)
    if fields.get("报送日期"):
        doc.add_paragraph("报送日期：%s" % fields["报送日期"])
    for key, value in fields.items():
        if key == "报送日期":
            continue
        if isinstance(value, list):
            lines = ["%s：" % key]
            lines.extend("- %s" % item for item in value)
            doc.add_paragraph("\n".join(lines))
        else:
            doc.add_paragraph("%s：%s" % (key, value))
    for para in doc.paragraphs:
        for run in para.runs:
            run.font.size = Pt(11)
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def main(argv=None):
    parser = argparse.ArgumentParser(
        prog="gen_doc.py", description="按模板与数据文件生成 Word 文书"
    )
    parser.add_argument("--template", required=True, help="模板名（周报/工作总结/会议通知/请示函）")
    parser.add_argument("--data", required=True, help="数据文件路径（JSON）")
    parser.add_argument("--outdir", default="out", help="输出根目录，默认 out")
    args = parser.parse_args(argv)

    # —— 校验 1：模板名（先于任何 IO 之外的检查，且先于建目录）——
    if args.template not in TEMPLATES:
        _fail(
            "模板名非法：%s，可用模板为 %s。"
            % (args.template, "、".join(sorted(TEMPLATES)))
        )

    # —— 校验 2/3/4：数据文件存在性、JSON 可解析性、顶层为对象 ——
    data = load_data(args.data)

    # —— 校验通过后才做缺省补齐（缺失不算非法）——
    fields = normalize(data, TEMPLATES[args.template], args.template)

    # —— 全部校验通过后才开始产生产物 ——
    case = os.path.splitext(os.path.basename(args.data))[0]
    outdir = os.path.join(args.outdir, args.template, case)
    docx_bytes = build_docx(args.template, fields)          # 内存构造，失败不留痕
    fields_json = json.dumps(fields, ensure_ascii=False, indent=2)

    os.makedirs(outdir, exist_ok=True)
    with open(os.path.join(outdir, "fields.json"), "w", encoding="utf-8") as fh:
        fh.write(fields_json)
    with open(os.path.join(outdir, "文书.docx"), "wb") as fh:
        fh.write(docx_bytes)

    sys.stdout.write("已生成：%s\n" % outdir)
    sys.stdout.flush()
    return EXIT_OK


if __name__ == "__main__":
    sys.exit(main())
