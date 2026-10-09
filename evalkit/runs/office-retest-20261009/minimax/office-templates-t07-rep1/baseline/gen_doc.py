#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""周报成文引擎（baseline 臂自建实现）。

把「工作总结」格式的季度素材映射进「周报」字段表，产出 文书.docx + fields.json。

纪律：缺数据不编造。无逐字来源的字段一律记 missing 并以 ____ 占位渲染，
不以「看起来合理」的值填充。
"""
import argparse
import datetime
import json
import os
import re
import sys
import zipfile

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.oxml.ns import qn
from docx.shared import Cm, Pt

PLACEHOLDER = "____"

# 周报字段表：(字段名, 必填, kind)
WEEKLY_FIELDS = [
    ("部门", True, "text"),
    ("填报人", True, "text"),
    ("周期", True, "text"),
    ("本周工作内容", True, "list"),
    ("下周工作计划", True, "list"),
    ("问题与需协调事项", False, "list"),
    ("报送日期", False, "text"),
]

# 素材键 -> 周报字段名。映射只允许搬运原文，不允许改写。
SOURCE_KEYS = ["总结主体", "总结时段", "工作回顾", "主要成绩", "存在问题", "下一步工作打算", "成文日期"]

# 工作总结素材键 -> 周报字段 的映射表（逐字搬运）
MAPPING = {
    "部门": "总结主体",
    "周期": "总结时段",
    "本周工作内容": "工作回顾",
    "下周工作计划": "下一步工作打算",
    "问题与需协调事项": "存在问题",
    "报送日期": "成文日期",
    # 「填报人」在季度素材中无对应键 —— 无来源，保持缺失，不编造。
}


def die(msg):
    sys.stderr.write(msg + "\n")
    raise SystemExit(2)


def is_blank(v):
    """F1 缺失判定：键不存在 / 空白串 / 空数组 / 全空白数组。"""
    if v is None:
        return True
    if isinstance(v, str):
        return v.strip() == ""
    if isinstance(v, (list, tuple)):
        if len(v) == 0:
            return True
        return all(is_blank(x) for x in v)
    return False


def to_text(v):
    """F3/F5：宽容取值 + strip。"""
    if isinstance(v, str):
        return v.strip()
    if isinstance(v, (list, tuple)):
        return str(v[0]).strip() if v else ""
    return str(v).strip()


def to_list(v):
    """F3/F5：字符串→单项；数组→逐项 str().strip() 去空。"""
    if isinstance(v, str):
        s = v.strip()
        return [s] if s else []
    if isinstance(v, (list, tuple)):
        out = []
        for x in v:
            if isinstance(x, str):
                s = x.strip()
            elif x is None:
                continue
            else:
                s = str(x).strip()
            if s:
                out.append(s)
        return out
    if v is None:
        return []
    s = str(v).strip()
    return [s] if s else []


def set_font(run, east_asia, size_pt, ascii_font="Times New Roman"):
    run.font.size = Pt(size_pt)
    run.font.name = ascii_font
    rpr = run._element.get_or_add_rPr()
    rfonts = rpr.find(qn("w:rFonts"))
    if rfonts is None:
        rfonts = rpr.makeelement(qn("w:rFonts"), {})
        rpr.append(rfonts)
    rfonts.set(qn("w:ascii"), ascii_font)
    rfonts.set(qn("w:hAnsi"), ascii_font)
    rfonts.set(qn("w:eastAsia"), east_asia)


def set_first_line_chars(par, chars=200):
    """V4：首行缩进两字符（w:firstLineChars=200）。"""
    ppr = par._element.get_or_add_pPr()
    ind = ppr.find(qn("w:ind"))
    if ind is None:
        ind = ppr.makeelement(qn("w:ind"), {})
        ppr.append(ind)
    ind.set(qn("w:firstLineChars"), str(chars))


def body_par(doc, text, east_asia="仿宋", size=16, align=None, indent_chars=200):
    p = doc.add_paragraph()
    pf = p.paragraph_format
    pf.space_before = Pt(0)
    pf.space_after = Pt(0)
    pf.line_spacing_rule = WD_LINE_SPACING.EXACTLY
    pf.line_spacing = Pt(28)
    if align is not None:
        p.alignment = align
    if indent_chars:
        set_first_line_chars(p, indent_chars)
    r = p.add_run(text)
    set_font(r, east_asia, size)
    return p


def normalize_docx(path):
    """V8 确定性：归一化 zip 条目时间戳与文档属性时间源，整文件逐字节可复现。"""
    with zipfile.ZipFile(path) as zf:
        names = zf.namelist()
        blobs = {n: zf.read(n) for n in names}
    tmp = path + ".tmp"
    with zipfile.ZipFile(tmp, "w", zipfile.ZIP_DEFLATED) as out:
        for n in names:
            zi = zipfile.ZipInfo(n, date_time=(1980, 1, 1, 0, 0, 0))
            zi.compress_type = zipfile.ZIP_DEFLATED
            zi.external_attr = 0o600 << 16
            out.writestr(zi, blobs[n])
    os.replace(tmp, path)


def build_doc(fields, title, out_docx):
    doc = Document()

    sec = doc.sections[0]
    sec.page_width = Cm(21.0)
    sec.page_height = Cm(29.7)
    sec.top_margin = Cm(3.7)
    sec.bottom_margin = Cm(3.5)
    sec.left_margin = Cm(2.8)
    sec.right_margin = Cm(2.6)

    val = {f["name"]: f for f in fields}
    filled = lambda n: (val[n]["status"] == "filled")

    def text_of(n):
        if filled(n):
            return val[n]["value"]
        return PLACEHOLDER if val[n]["required"] else ""

    # V2 标题：居中，黑体二号
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(0)
    r = p.add_run(title)
    set_font(r, "黑体", 22)

    # 基本信息行：填报人：____　　周期：...
    body_par(
        doc,
        "填报人：%s　　周期：%s" % (text_of("填报人"), text_of("周期")),
    )

    sections = [
        ("一、本周工作内容", "本周工作内容"),
        ("二、下周工作计划", "下周工作计划"),
        ("三、问题与需协调事项", "问题与需协调事项"),
    ]
    for head, fname in sections:
        f = val[fname]
        # F4 选填列表缺失 → 整节省略，文中不出现节头
        if f["status"] == "missing" and not f["required"]:
            continue
        body_par(doc, head, east_asia="黑体")
        items = f["value"] if f["status"] == "filled" else None
        if not items:
            items = [PLACEHOLDER]  # 必填缺失 → 占位条目
        for i, it in enumerate(items, 1):
            body_par(doc, "%d．%s" % (i, it))

    # V5 落款：右对齐、右缩进两字（≈32pt）
    p = body_par(
        doc,
        "报送日期：%s" % text_of("报送日期"),
        align=WD_ALIGN_PARAGRAPH.RIGHT,
        indent_chars=0,
    )
    p.paragraph_format.right_indent = Pt(32)

    # 文档属性去时间源：钉死为固定纪元，不取当前时间
    cp = doc.core_properties
    epoch = datetime.datetime(1980, 1, 1, 0, 0, 0)
    cp.created = epoch
    cp.modified = epoch
    cp.last_printed = epoch
    cp.revision = 1
    cp.author = ""
    cp.last_modified_by = ""
    cp.title = title

    doc.save(out_docx)
    normalize_docx(out_docx)


def main():
    ap = argparse.ArgumentParser(description="周报成文引擎")
    ap.add_argument("--data", required=True)
    ap.add_argument("--outdir", required=True)
    args = ap.parse_args()

    if not os.path.isfile(args.data):
        die("数据文件不存在：%s" % args.data)
    try:
        with open(args.data, encoding="utf-8") as fh:
            raw = fh.read()
    except OSError as exc:
        die("数据文件读取失败：%s" % exc)
    try:
        data = json.loads(raw)
    except json.JSONDecodeError:
        die("数据文件不是合法 JSON")
    if not isinstance(data, dict):
        die("顶层必须是对象")

    # ---- 字段台账：逐字搬运，无来源即缺失 ----
    fields = []
    consumed = set()
    for name, required, kind in WEEKLY_FIELDS:
        src = MAPPING.get(name)
        if src is None:
            # 模板字段在素材中无对应键 —— 无真实来源，保持缺失
            fields.append({
                "name": name, "required": required, "kind": kind,
                "status": "missing", "value": None,
                "rendered_as": "以____占位" if required else "留空（渲染为空字符串）",
                "source_key": None,
                "note": "季度素材中无对应键，无真实来源，未编造",
            })
            continue
        consumed.add(src)
        if src not in data:
            fields.append({
                "name": name, "required": required, "kind": kind,
                "status": "missing", "value": None,
                "rendered_as": "以____占位" if required else "留空（渲染为空字符串）",
                "source_key": src, "note": "素材键缺失",
            })
            continue
        v = data[src]
        if is_blank(v):
            fields.append({
                "name": name, "required": required, "kind": kind,
                "status": "missing", "value": None,
                "rendered_as": "以____占位条目" if (required and kind == "list") else
                               ("以____占位" if required else "留空（渲染为空字符串）"),
                "source_key": src, "note": "素材值为空",
            })
            continue
        val = to_list(v) if kind == "list" else to_text(v)
        fields.append({
            "name": name, "required": required, "kind": kind,
            "status": "filled", "value": val,
            "source_key": src,
        })

    # F6 模板外键：未被消费的工作总结键
    unknown_keys = sorted(k for k in data if k not in consumed)

    filled_fields = [f["name"] for f in fields if f["status"] == "filled"]
    missing_fields = [f["name"] for f in fields if f["status"] == "missing"]
    missing_required = [f["name"] for f in fields
                        if f["status"] == "missing" and f["required"]]

    v = {f["name"]: f for f in fields}
    title = "%s工作周报" % (v["部门"]["value"] if v["部门"]["status"] == "filled"
                            else PLACEHOLDER)

    summary = {
        "total": len(fields),
        "filled": len(filled_fields),
        "missing": len(missing_fields),
        "missing_required": len(missing_required),
        "missing_required_names": missing_required,
        "unknown_keys": unknown_keys,
    }

    os.makedirs(args.outdir, exist_ok=True)
    docx_path = os.path.join(args.outdir, "文书.docx")
    build_doc(fields, title, docx_path)

    ledger = {
        "template": "周报",
        "title": title,
        "filled_fields": filled_fields,
        "missing_fields": missing_fields,
        "fields": fields,
        "summary": summary,
        "outputs": {"文书.docx": "文书.docx"},
    }
    with open(os.path.join(args.outdir, "fields.json"), "w", encoding="utf-8") as fh:
        json.dump(ledger, fh, ensure_ascii=False, indent=2)
        fh.write("\n")

    sys.stdout.write(
        "周报已生成：%s（filled %d/%d，missing %d，unknown_keys %s）\n"
        % (args.outdir, summary["filled"], summary["total"],
           summary["missing"], unknown_keys or "无")
    )
    return 0


if __name__ == "__main__":
    sys.exit(main())
