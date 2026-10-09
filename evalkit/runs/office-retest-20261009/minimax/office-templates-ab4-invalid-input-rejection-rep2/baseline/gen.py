#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""文书生成入口（baseline 臂自研实现，ab4-invalid-input-rejection）。

用法::

    python3 gen.py --template 周报 --data fixtures/周报/case1.json [--outdir out]

行为约定（对应本臂任务书）：

* 四类非法输入 —— 模板名非法 / 数据文件不存在 / JSON 解析失败 / JSON 顶层非对象 ——
  一律以退出码 2 拒绝，错误信息写 stderr（中文），并且在校验全部通过之前
  不触碰文件系统（不建输出目录、不写半成品），因此非法调用零产物残留。
* 字段缺失不算非法输入：缺失字段按「（未填写）」占位，正常出文书、退出码 0。
* 合法输入退出码 0，产物落在 ``<outdir>/<模板名>/<数据文件名>/`` 下的
  ``文书.docx`` 与 ``fields.json`` 两件。
"""

import argparse
import json
import os
import sys
from pathlib import Path

EXIT_OK = 0
EXIT_INVALID_INPUT = 2

PLACEHOLDER = "（未填写）"

# 模板登记表：模板名 -> 标题 + 有序字段（键, 中文标签）。
# 「证书」刻意不在表内，用作模板名非法的探针。
TEMPLATES = {
    "周报": {
        "title": "周报",
        "fields": [
            ("部门", "部门"),
            ("填报人", "填报人"),
            ("周期", "周期"),
            ("本周工作内容", "本周工作内容"),
            ("下周工作计划", "下周工作计划"),
            ("问题与需协调事项", "问题与需协调事项"),
            ("报送日期", "报送日期"),
        ],
    },
    "月报": {
        "title": "月报",
        "fields": [
            ("部门", "部门"),
            ("填报人", "填报人"),
            ("周期", "周期"),
            ("本月工作内容", "本月工作内容"),
            ("下月工作计划", "下月工作计划"),
            ("问题与需协调事项", "问题与需协调事项"),
            ("报送日期", "报送日期"),
        ],
    },
    "会议纪要": {
        "title": "会议纪要",
        "fields": [
            ("会议主题", "会议主题"),
            ("会议时间", "会议时间"),
            ("会议地点", "会议地点"),
            ("主持人", "主持人"),
            ("参会人员", "参会人员"),
            ("会议内容", "会议内容"),
            ("待办事项", "待办事项"),
        ],
    },
}


class InvalidInput(Exception):
    """输入非法：统一映射到退出码 2 + 中文 stderr。"""


def _ensure_docx_importable():
    """在 venv（user site 被禁用）里也能 import 到 python-docx。"""
    try:
        import docx  # noqa: F401

        return
    except ImportError:
        pass
    for extra in (
        Path.home() / ".local/lib/python3.12/site-packages",
        Path.home() / ".local/lib/python3.11/site-packages",
        Path.home() / ".local/lib/python3.10/site-packages",
    ):
        if extra.is_dir() and str(extra) not in sys.path:
            sys.path.append(str(extra))
    try:
        import docx  # noqa: F401
    except ImportError as exc:  # pragma: no cover - 环境缺依赖时的兜底
        raise SystemExit(
            "错误：运行环境缺少 python-docx 依赖，无法生成文书（%s）" % exc
        )


class ChineseArgumentParser(argparse.ArgumentParser):
    """argparse 默认报错是英文，这里保证 stderr 一定有中文提示。"""

    def error(self, message):  # noqa: D401 - argparse 钩子
        sys.stderr.write(
            "错误：命令行参数不合法（缺少必要参数或取值无法识别）。"
            "用法：--template <模板名> --data <JSON 数据文件路径> [--outdir <输出根目录>]\n"
            "（argparse 原始提示：%s）\n" % message
        )
        raise SystemExit(EXIT_INVALID_INPUT)


def build_parser():
    parser = ChineseArgumentParser(
        prog="gen.py", description="文书生成入口：按模板把 JSON 数据渲染成 文书.docx"
    )
    parser.add_argument("--template", required=True, help="模板名，如 周报")
    parser.add_argument("--data", required=True, help="JSON 数据文件路径")
    parser.add_argument("--outdir", default="out", help="输出根目录（默认 out）")
    return parser


def load_template(name):
    spec = TEMPLATES.get(name)
    if spec is None:
        raise InvalidInput(
            "未知模板「%s」。可用模板：%s。"
            % (name, "、".join(sorted(TEMPLATES)))
        )
    return spec


def load_data(path_text):
    """读取并解析数据文件；任何一步失败都算非法输入。"""
    path = Path(path_text)
    if not path.exists():
        raise InvalidInput("数据文件不存在：%s" % path_text)
    if not path.is_file():
        raise InvalidInput("数据路径不是普通文件：%s" % path_text)
    try:
        raw = path.read_text(encoding="utf-8")
    except UnicodeDecodeError as exc:
        raise InvalidInput(
            "数据文件不是合法的 UTF-8 文本：%s（%s）" % (path_text, exc.reason)
        )
    except OSError as exc:
        raise InvalidInput("数据文件读取失败：%s（%s）" % (path_text, exc))
    try:
        payload = json.loads(raw)
    except json.JSONDecodeError as exc:
        raise InvalidInput(
            "数据文件不是合法 JSON：%s（第 %d 行第 %d 列：%s）"
            % (path_text, exc.lineno, exc.colno, exc.msg)
        )
    if not isinstance(payload, dict):
        raise InvalidInput(
            "数据文件顶层必须是 JSON 对象，实际为 %s：%s"
            % (_json_type_cn(payload), path_text)
        )
    return path, payload


def _json_type_cn(value):
    if isinstance(value, list):
        return "数组"
    if isinstance(value, str):
        return "字符串"
    if isinstance(value, bool):
        return "布尔值"
    if isinstance(value, (int, float)):
        return "数值"
    if value is None:
        return "null"
    return type(value).__name__


def render_value(value):
    if value is None or value == "":
        return PLACEHOLDER
    if isinstance(value, list):
        items = [render_value(v) for v in value]
        return "；".join(i for i in items if i) or PLACEHOLDER
    if isinstance(value, dict):
        return "；".join("%s：%s" % (k, render_value(v)) for k, v in value.items())
    return str(value)


def resolve_fields(spec, payload):
    """逐字段取值；缺失或空值一律占位，不视为非法输入。"""
    resolved = []
    missing = []
    for key, label in spec["fields"]:
        present = key in payload and payload[key] not in (None, "", [], {})
        value = render_value(payload.get(key)) if present else PLACEHOLDER
        if not present:
            missing.append(key)
        resolved.append(
            {"key": key, "label": label, "present": present, "value": value}
        )
    return resolved, missing


def write_atomic(target, writer):
    """先写同目录临时文件再 os.replace，保证不留半成品。"""
    target.parent.mkdir(parents=True, exist_ok=True)
    tmp = target.parent / (".{}.tmp".format(target.name))
    try:
        writer(str(tmp))
        os.replace(str(tmp), str(target))
    finally:
        if tmp.exists():
            tmp.unlink()


def build_docx(path, spec, resolved):
    from docx import Document
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.oxml.ns import qn
    from docx.shared import Pt

    doc = Document()
    normal = doc.styles["Normal"]
    normal.font.name = "SimSun"
    normal.font.size = Pt(11)
    rpr = normal.element.get_or_add_rPr()
    rpr.rFonts.set(qn("w:eastAsia"), "宋体")

    heading = doc.add_heading(spec["title"], level=0)
    heading.alignment = WD_ALIGN_PARAGRAPH.CENTER

    for item in resolved:
        para = doc.add_paragraph()
        label_run = para.add_run("%s：" % item["label"])
        label_run.bold = True
        para.add_run(item["value"])

    doc.save(path)


def run(args):
    spec = load_template(args.template)
    data_path, payload = load_data(args.data)
    resolved, missing = resolve_fields(spec, payload)

    # 校验全部通过后，才开始落盘。
    case_dir = Path(args.outdir) / args.template / data_path.stem
    docx_path = case_dir / "文书.docx"
    fields_path = case_dir / "fields.json"

    _ensure_docx_importable()
    write_atomic(docx_path, lambda p: build_docx(p, spec, resolved))

    fields_doc = {
        "template": args.template,
        "title": spec["title"],
        "source": str(data_path),
        "case": data_path.stem,
        "artifacts": ["文书.docx", "fields.json"],
        "fields": resolved,
        "missing_fields": missing,
    }
    def _write_fields(tmp_path):
        Path(tmp_path).write_text(
            json.dumps(fields_doc, ensure_ascii=False, indent=2) + "\n",
            encoding="utf-8",
        )

    write_atomic(fields_path, _write_fields)

    sys.stdout.write(
        "已生成：%s、%s（缺失字段 %d 个：%s）\n"
        % (
            docx_path,
            fields_path,
            len(missing),
            "、".join(missing) if missing else "无",
        )
    )
    return EXIT_OK


def main(argv=None):
    args = build_parser().parse_args(argv)
    try:
        return run(args)
    except InvalidInput as exc:
        sys.stderr.write("错误：%s\n" % exc)
        return EXIT_INVALID_INPUT


if __name__ == "__main__":
    sys.exit(main())