#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""交付自检：核对两个单元目录的产物齐全、docx 可开、fields.json 自洽。

    /usr/bin/python3 verify.py
退出码 0 = 全绿；非 0 = 有断言失败。
"""

import json
import os
import sys

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH

BASE = os.path.dirname(os.path.abspath(__file__))
CASES = [
    {
        "dir": "out/请示函-empty",
        "template": "请示函",
        "missing_required": ["发文机关", "主送机关", "请示事由", "请示事项"],
        "missing_optional": ["成文日期"],
        "docx_must_contain": ["关于____的请示", "____：", "一、____", "妥否，请批示。"],
        "docx_must_not_contain": ["三、", "存在问题"],
    },
    {
        "dir": "out/工作总结-blank",
        "template": "工作总结",
        "missing_required": ["总结主体", "总结时段", "工作回顾", "主要成绩", "下一步工作打算"],
        "missing_optional": ["存在问题", "成文日期"],
        "docx_must_contain": ["工作总结", "一、工作回顾", "1．____", "二、主要成绩",
                              "三、下一步工作打算"],
        "docx_must_not_contain": ["存在问题"],
    },
]

fails = []


def ck(cond, msg):
    if not cond:
        fails.append(msg)


for case in CASES:
    d = os.path.join(BASE, case["dir"])
    docx_p = os.path.join(d, "文书.docx")
    json_p = os.path.join(d, "fields.json")

    ck(os.path.isdir(d), "缺目录：%s" % case["dir"])
    ck(os.path.isfile(docx_p), "缺产物：%s/文书.docx" % case["dir"])
    ck(os.path.isfile(json_p), "缺产物：%s/fields.json" % case["dir"])
    ck(sorted(os.listdir(d)) == ["fields.json", "文书.docx"] or
       sorted(os.listdir(d)) == sorted(["文书.docx", "fields.json"]),
       "目录内产物不是恰好两个：%s -> %s" % (case["dir"], sorted(os.listdir(d))))
    if not (os.path.isfile(docx_p) and os.path.isfile(json_p)):
        continue

    doc = Document(docx_p)
    text = "\n".join(p.text for p in doc.paragraphs)
    for frag in case["docx_must_contain"]:
        ck(frag in text, "%s 正文缺片段：%r" % (case["dir"], frag))
    for frag in case["docx_must_not_contain"]:
        ck(frag not in text, "%s 正文出现不该有的片段：%r" % (case["dir"], frag))

    with open(json_p, encoding="utf-8") as fh:
        led = json.load(fh)
    for key in ("title", "fields", "missing_required_names", "missing_optional",
                "unknown_keys", "summary"):
        ck(key in led, "%s fields.json 缺键：%s" % (case["dir"], key))
    ck(led.get("template") == case["template"],
       "%s template 不符：%r" % (case["dir"], led.get("template")))
    ck(led.get("missing_required_names") == case["missing_required"],
       "%s missing_required_names 不符：%r" % (case["dir"], led.get("missing_required_names")))
    ck(led.get("missing_optional") == case["missing_optional"],
       "%s missing_optional 不符：%r" % (case["dir"], led.get("missing_optional")))
    ck(led.get("unknown_keys") == [], "%s unknown_keys 非空：%r" % (case["dir"], led.get("unknown_keys")))

    names = [e["name"] for e in led["fields"]]
    ck(len(names) == len(set(names)), "%s 字段台账有重名" % case["dir"])
    for e in led["fields"]:
        ck(e["state"] in ("rendered", "missing_required", "missing_optional"),
           "%s 字段 %s 状态非法：%r" % (case["dir"], e["name"], e["state"]))
        if e["state"].startswith("missing"):
            ck(e["value"] in (None, []), "%s 字段 %s 缺失但有值" % (case["dir"], e["name"]))
            ck(bool(e["rendered_as"]), "%s 字段 %s 缺失未记占位语义" % (case["dir"], e["name"]))
    s = led["summary"]
    ck(s["required_missing"] == len(led["missing_required_names"]),
       "%s summary.required_missing 与列表不符" % case["dir"])
    ck(s["optional_missing"] == len(led["missing_optional"]),
       "%s summary.optional_missing 与列表不符" % case["dir"])

    # 落款右对齐不因空值破坏
    aligns = [p.alignment for p in doc.paragraphs]
    ck(sum(1 for a in aligns if a == WD_ALIGN_PARAGRAPH.RIGHT) >= 2,
       "%s 落款右对齐段落不足" % case["dir"])
    ck(sum(1 for a in aligns if a == WD_ALIGN_PARAGRAPH.CENTER) >= 1,
       "%s 标题未居中" % case["dir"])

if fails:
    print("FAIL (%d)" % len(fails))
    for f in fails:
        print("  - %s" % f)
    sys.exit(1)

print("PASS：2 个单元目录 × 2 产物齐全，docx 可开、版式占位符合预期、fields.json 自洽。")
