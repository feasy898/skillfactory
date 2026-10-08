#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""runs/office-retest-20261009/x_ab5_replay.py — ab5 rubric ③④ 裁判侧臂外重放取证工具（零模型）。

依据 planner 裁决 planning/test-cards-20261008/ab5-protocol-adjudication.md（提交 9827318）：
ab5 的 ③④（确定性对拍、红绿自校验）改由裁判在臂外对交付产物重放计分；①②⑤ 仍按原口径
（产物自含 + run_log 留痕可作辅助证据）。本工具为重放的机械执行层（判据面不进臂、裁判盲态
由上层 x_judges.py 匿名化保证），产出纯事实 JSON 供裁判按 rubric 档位计分：

  A. 确定性对拍重放：从臂交付 run_log.json 提取生成命令，在沙盒（臂树副本）以同命令连跑
     两遍到同一 outdir（第二遍覆盖复写），逐遍 sha256 快照并双跑对拍——
     文书.docx 整文件 sha256 是否一致 + python-docx 解析级（段落文本/对齐/缩进/字体/行距）
     是否一致；fields.json 整文件/除 outputs 键外是否一致；并对照臂交付产物 sha256。
     命令不可提取时如实 replay_skipped（通道/装置降级留痕，不假绿）。
  B. runner 红绿复检（reference 口径，重打包后 dist 包内）：
     绿=eval/runner.py reference/out reference/out（期望 exit 0 ok=true 68 检查）；
     红=eval/runner.py <空目录> reference/out（期望 exit 1 且 units_found 不通过）。
  C. 交付单元对拍（rubric ③ 端态等价）：eval/runner.py <臂交付out根> reference/out，
     提取交付单元的实质检查结果与该单元限定的字段填充一致率（units_found/products_exact
     在单任务臂上下文按 N/A 记录，不以 8 单元全集口径错判交付单元）。

用法：python x_ab5_replay.py <arm_dir> <out_json> [--sandbox <dir>]
退出码：0=重放完成（无论事实红绿，事实如实记录）；2=装置错误（臂目录不存在等）。
"""

import argparse
import hashlib
import json
import os
import re
import shutil
import subprocess
import sys
import tempfile
import time

ROOT = "/home/pan-ding/workspace/skill工厂/skillfactory"
DIST = os.path.join(ROOT, "dist", "office-templates-skill")
VPY = os.path.join(ROOT, ".venv-dev", "bin", "python")
DOCX_NAME = "文书.docx"
FIELDS_NAME = "fields.json"
REPLAY_TIMEOUT = 120


def sha256_file(p):
    h = hashlib.sha256()
    with open(p, "rb") as f:
        for c in iter(lambda: f.read(1 << 20), b""):
            h.update(c)
    return h.hexdigest()


def find_out_units(arm_dir):
    """交付产物单元：arm 树内含 文书.docx+fields.json 的目录（records/ 之外）。"""
    units = []
    for dp, dns, fns in os.walk(arm_dir):
        if os.path.sep + "records" in dp or os.path.sep + ".zcode" in dp:
            continue
        if DOCX_NAME in fns and FIELDS_NAME in fns:
            units.append(os.path.relpath(dp, arm_dir))
    return sorted(units)


def extract_commands(run_log_path, arm_dir):
    """从 run_log.json 递归提取生成命令。两种形态：字符串数组（engine_invocation）或
    单条命令字符串（如「<python> generate.py --template … --outdir …」，shlex 切分）。
    硬性要求：命令须同时含 --template 与 --outdir（防把自评/对拍脚本误当生成命令）。"""
    try:
        import shlex
        data = json.load(open(run_log_path, encoding="utf-8"))
    except Exception:
        return []
    cmds = []
    seen = set()

    def consider(argv):
        if not (isinstance(argv, list) and len(argv) >= 3
                and all(isinstance(s, str) for s in argv)):
            return
        joined = " ".join(argv)
        if "--template" in joined and "--outdir" in joined and joined not in seen:
            seen.add(joined)
            cmds.append(list(argv))

    def walk(x):
        if isinstance(x, list):
            if x and all(isinstance(s, str) for s in x):
                consider(list(x))
            for v in x:
                walk(v)
        elif isinstance(x, dict):
            for v in x.values():
                walk(v)
        elif isinstance(x, str):
            if "--template" in x and "--outdir" in x:
                try:
                    consider(shlex.split(x))
                except ValueError:
                    consider(x.split())
    walk(data)
    return cmds


def docx_layout_facts(path):
    """python-docx 解析级事实（rubric ① 的 docx 口径：段落文本/对齐/缩进/字体/行距）。"""
    try:
        from docx import Document
        doc = Document(path)
        facts = []
        for p in doc.paragraphs:
            pf = p.paragraph_format
            fonts = []
            for r in p.runs[:4]:
                east = None
                rPr = r._element.rPr
                if rPr is not None and rPr.rFonts is not None:
                    east = rPr.rFonts.get(
                        "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}eastAsia")
                fonts.append([r.text, east, r.font.size.pt if r.font.size else None])
            facts.append({"text": p.text, "align": str(p.alignment), "first_line_indent":
                          str(pf.first_line_indent), "line_spacing": str(pf.line_spacing),
                          "right_indent": str(pf.right_indent), "runs": fonts})
        return {"ok": True, "paragraphs": facts,
                "tables": [[[c.text for c in row.cells] for row in t.rows] for t in doc.tables]}
    except Exception as e:  # noqa: BLE001
        return {"ok": False, "error": str(e)[:300]}


def fields_modulo_outputs(path):
    """fields.json 内容，剥 outputs 键（路径键随平台写法变化属 spec S4 豁免面）。"""
    d = json.load(open(path, encoding="utf-8"))
    if isinstance(d, dict):
        d = {k: v for k, v in d.items() if k != "outputs"}
    return json.dumps(d, ensure_ascii=False, sort_keys=False)


def run_runner(args_list):
    """在重打包后 dist 包内跑 runner（reference 口径），返回 (exit_code, json_dict|None)。"""
    p = subprocess.run([VPY, "eval/runner.py"] + args_list, capture_output=True,
                       cwd=DIST, timeout=REPLAY_TIMEOUT)
    try:
        return p.returncode, json.loads(p.stdout.decode("utf-8"))
    except Exception:
        return p.returncode, None


def replay_command(cmd, arm_dir, sandbox, outdir_rel, interpreter):
    """沙盒内单遍执行：返回 (exit_code, {file: sha256}, stderr_head)。"""
    p = subprocess.run([interpreter] + cmd[1:] if _looks_interpreter(cmd) else cmd,
                       capture_output=True, cwd=sandbox, timeout=REPLAY_TIMEOUT)
    out = os.path.join(sandbox, outdir_rel)
    shas = {}
    if os.path.isdir(out):
        for fn in sorted(os.listdir(out)):
            fp = os.path.join(out, fn)
            if os.path.isfile(fp):
                shas[fn] = sha256_file(fp)
    return p.returncode, shas, p.stderr.decode("utf-8", errors="replace")[:400]


def _looks_interpreter(cmd):
    first = os.path.basename(cmd[0] or "")
    return first.startswith("python") or first.endswith("python") or first in ("python3", "python")


def main(argv=None):
    # docx 解析级事实需要 python-docx：当前解释器缺依赖时以装置 venv 重执行自身（如实留痕）
    try:
        import docx  # noqa: F401
    except ImportError:
        if os.path.abspath(sys.executable) != os.path.abspath(VPY) and os.path.exists(VPY):
            os.execv(VPY, [VPY, os.path.abspath(__file__)] + list(sys.argv[1:]))
        print("警告：python-docx 不可用且 venv 重执行失败，docx 解析级事实将缺失（如实降级）",
              file=sys.stderr)
    ap = argparse.ArgumentParser(add_help=False)
    ap.add_argument("arm_dir")
    ap.add_argument("out_json")
    ap.add_argument("--sandbox", default=None)
    a = ap.parse_args(argv)
    arm_dir = os.path.abspath(a.arm_dir)
    if not os.path.isdir(arm_dir):
        print("arm 目录不存在: %s" % arm_dir, file=sys.stderr)
        return 2
    facts = {"tool": "x_ab5_replay.py", "adjudication":
             "planning/test-cards-20261008/ab5-protocol-adjudication.md（9827318）",
             "started_at": time.strftime("%Y-%m-%dT%H:%M:%S%z"),
             "interpreter": VPY, "dist_package": os.path.relpath(DIST, ROOT)}
    # ---- B) runner 红绿复检（reference 口径；重打包后包内）
    rc_green, j_green = run_runner(["reference/out", "reference/out"])
    empty_dir = tempfile.mkdtemp(prefix="ab5red-")
    try:
        rc_red, j_red = run_runner([empty_dir, "reference/out"])
    finally:
        shutil.rmtree(empty_dir, ignore_errors=True)
    red_units_found = None
    if j_red:
        red_units_found = next((c for c in j_red.get("checks", [])
                                if c["name"] == "units_found"), {})
    facts["runner_redgreen"] = {
        "green": {"exit_code": rc_green, "ok": (j_green or {}).get("ok"),
                  "checks_total": len((j_green or {}).get("checks", [])),
                  "checks_passed": sum(1 for c in (j_green or {}).get("checks", []) if c["pass"])},
        "red": {"exit_code": rc_red, "units_found_pass": (red_units_found or {}).get("pass"),
                "units_found_detail": (red_units_found or {}).get("detail", "")[:300]},
    }
    # ---- A) 确定性对拍重放
    out_units = find_out_units(arm_dir)
    facts["delivered_units"] = out_units
    run_log = os.path.join(arm_dir, "run_log.json")
    cmds = extract_commands(run_log, arm_dir) if os.path.isfile(run_log) else []
    det = {"run_log_present": os.path.isfile(run_log), "commands_extracted": len(cmds)}
    if cmds:
        cmd = cmds[-1]  # 生成命令取 run_log 中最后一条带模板/outdir 的命令
        m = re.search(r"(?:--outdir(?:=|\s+))([^\s]+)", " ".join(cmd))
        outdir_rel = m.group(1) if m else (out_units[0] if out_units else "out")
        sandbox = a.sandbox or tempfile.mkdtemp(prefix="ab5replay-")
        made_sandbox = a.sandbox is None
        try:
            if os.path.isdir(sandbox):
                shutil.rmtree(sandbox)
            shutil.copytree(arm_dir, sandbox,
                            ignore=shutil.ignore_patterns("records", "__pycache__", ".git"))
            for stub in (".venv", "venv"):  # 臂内解释器桩 → 沙盒统一 venv 解释器（如实留痕）
                p = os.path.join(sandbox, stub)
                if os.path.isdir(p):
                    shutil.rmtree(p)
            # 交付产物树先清空：双跑必须真实再生成（防把臂交付当跑出结果的自证）
            for stale in (outdir_rel, "out"):
                sp = os.path.join(sandbox, stale)
                if os.path.isdir(sp):
                    shutil.rmtree(sp)
            i0 = VPY if _looks_interpreter(cmd) else cmd[0]
            r1 = replay_command(cmd, arm_dir, sandbox, outdir_rel, i0)
            s1 = os.path.join(sandbox, outdir_rel)
            snap1 = {fn: sha256_file(os.path.join(s1, fn)) for fn in os.listdir(s1)} \
                if os.path.isdir(s1) else {}
            f1 = {fn: docx_layout_facts(os.path.join(s1, fn)) if fn.endswith(".docx")
                  else open(os.path.join(s1, fn), encoding="utf-8").read()
                  for fn in os.listdir(s1)} if os.path.isdir(s1) else {}
            # 隔 5 秒再跑第二遍：跨过 python-docx zip 时间戳 2 秒粒度（照补强回归「隔 5 秒双跑」先例）
            time.sleep(5)
            r2 = replay_command(cmd, arm_dir, sandbox, outdir_rel, i0)
            s2 = os.path.join(sandbox, outdir_rel)
            snap2 = {fn: sha256_file(os.path.join(s2, fn)) for fn in os.listdir(s2)} \
                if os.path.isdir(s2) else {}
            f2 = {fn: docx_layout_facts(os.path.join(s2, fn)) if fn.endswith(".docx")
                  else open(os.path.join(s2, fn), encoding="utf-8").read()
                  for fn in os.listdir(s2)} if os.path.isdir(s2) else {}
            cmp_d = {}
            for fn in sorted(set(snap1) | set(snap2)):
                ent = {"sha_run1": snap1.get(fn), "sha_run2": snap2.get(fn),
                       "byte_identical": snap1.get(fn) == snap2.get(fn) and snap1.get(fn)}
                if fn.endswith(".docx") and f1.get(fn, {}).get("ok") and f2.get(fn, {}).get("ok"):
                    ent["parsed_layout_identical"] = f1[fn] == f2[fn]
                if fn.endswith(".json"):
                    try:
                        ent["equal_modulo_outputs"] = (fields_modulo_outputs(os.path.join(s1, fn))
                                                       == fields_modulo_outputs(os.path.join(s2, fn)))
                    except Exception:
                        pass
                cmp_d[fn] = ent
            delivered_sha = {}
            for rel in out_units:
                du = os.path.join(arm_dir, rel)
                delivered_sha[rel] = {fn: sha256_file(os.path.join(du, fn))
                                      for fn in sorted(os.listdir(du))}
            det.update({"replayed_command": cmd, "outdir": outdir_rel,
                        "run1_exit": r1[0], "run2_exit": r2[0],
                        "run1_files": r1[1], "run2_files": r2[1],
                        "compare_run1_vs_run2": cmp_d,
                        "delivered_sha256": delivered_sha})
        finally:
            if made_sandbox:
                shutil.rmtree(sandbox, ignore_errors=True)
    else:
        det["replay_skipped"] = ("run_log.json 缺失或未提取到生成命令——按裁决降级：① 依据臂交付产物"
                                 "与 run_log 自录留痕判，重放缺考如实留痕（不假绿）")
    facts["determinism_replay"] = det
    # ---- C) 交付单元 runner 对拍（reference 口径）
    if out_units:
        # 单元名 = 相对路径末两段（<模板>/caseN，与 runner 检查名前缀同口径）
        delivered_names = {"/".join(rel.replace(os.sep, "/").split("/")[-2:]) for rel in out_units}
        out_root = os.path.dirname(os.path.join(arm_dir, out_units[0]))
        rc_cmp, j_cmp = run_runner([out_root, "reference/out"])
        unit_checks, all_unit_pass = {}, True
        for c in (j_cmp or {}).get("checks", []):
            if any(c["name"].startswith(u + "/") for u in delivered_names):
                unit_checks[c["name"]] = c["pass"]
                if not c["pass"]:
                    all_unit_pass = False
        # 单元限定的字段填充一致率：按 runner §10 同语义独立复算（status 相等且 filled 值相等）
        ag = next((c["detail"] for c in (j_cmp or {}).get("checks", [])
                   if c["name"] == "field_fill_agreement_ge_90pct"), "")
        unit_agree = {}
        for u in sorted(delivered_names):
            tpl, case = u.split("/")
            ref_p = os.path.join(DIST, "reference", "out", tpl, case, FIELDS_NAME)
            del_p = os.path.join(arm_dir, tpl, case, FIELDS_NAME)
            if not (os.path.isfile(ref_p) and os.path.isfile(del_p)):
                # 交付单元相对路径可能不止两层（如 out/<模板>/caseN）——按 out_units 原始相对路径找
                for rel in out_units:
                    if rel.replace(os.sep, "/").endswith(u):
                        del_p = os.path.join(arm_dir, rel)
                        break
            try:
                if os.path.isdir(del_p):
                    del_p = os.path.join(del_p, FIELDS_NAME)
                rf = json.load(open(ref_p, encoding="utf-8")).get("fields", [])
                tf = json.load(open(del_p, encoding="utf-8")).get("fields", [])
                t_by = {i.get("name"): i for i in tf if isinstance(i, dict)}
                matched = sum(1 for r in rf if isinstance(r, dict)
                              and (t := t_by.get(r.get("name"))) is not None
                              and r.get("status") == t.get("status")
                              and (r.get("status") != "filled" or r.get("value") == t.get("value")))
                total = len(rf)
                unit_agree[u] = {"matched": matched, "total": total,
                                 "rate": round(matched / total, 4) if total else 0.0,
                                 "ge_90pct": total > 0 and matched / total >= 0.90,
                                 "mismatched_fields": sorted(
                                     r.get("name") for r in rf if isinstance(r, dict)
                                     and not (t_by.get(r.get("name"))
                                              and r.get("status") == t_by[r.get("name")].get("status")
                                              and (r.get("status") != "filled"
                                                   or r.get("value") == t_by[r.get("name")].get("value"))))}
            except Exception as e:  # noqa: BLE001
                unit_agree[u] = {"error": str(e)[:200]}
        facts["runner_delivered"] = {
            "exit_code": rc_cmp, "ok": (j_cmp or {}).get("ok"),
            "delivered_unit_ids": sorted(delivered_names),
            "unit_scoped_checks": unit_checks, "unit_scoped_all_pass": all_unit_pass,
            "units_found_note": ("单任务臂交付恰含本任务单元；runner 的 units_found/products_exact "
                                 "按 8 单元全集口径在臂上下文记 N/A（裁决端态口径：以交付单元实质"
                                 "检查+单元限定一致率为准，不放宽、不错判）"),
            "agreement_detail": ag,
            "unit_scoped_agreement": unit_agree,
        }
    facts["finished_at"] = time.strftime("%Y-%m-%dT%H:%M:%S%z")
    os.makedirs(os.path.dirname(os.path.abspath(a.out_json)), exist_ok=True)
    with open(a.out_json, "w", encoding="utf-8", newline="\n") as f:
        json.dump(facts, f, ensure_ascii=False, indent=2)
        f.write("\n")
    print(json.dumps({"exit": 0, "out": a.out_json,
                      "replay": "done" if det.get("commands_extracted") else "skipped",
                      "green_exit": facts["runner_redgreen"]["green"]["exit_code"],
                      "red_exit": facts["runner_redgreen"]["red"]["exit_code"]}, ensure_ascii=False))
    return 0


if __name__ == "__main__":
    sys.exit(main())
