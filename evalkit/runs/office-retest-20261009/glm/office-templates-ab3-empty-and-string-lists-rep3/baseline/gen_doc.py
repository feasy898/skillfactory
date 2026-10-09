#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""ab3-empty-and-string-lists 基线文书生成器。

命令行契约：
    python gen_doc.py --template <周报|工作总结> --data <数据json路径> --outdir <输出目录>
每次调用在 outdir 生成两件产物：文书.docx 与 fields.json。

核心行为：
- 空列表按必填缺失处理：渲染占位条目（如「1．____」），fields.json 记占位语义；
- 字符串型列表按单项化处理：渲染为单项条目，fields.json 中该字段值为单元素数组；
- 选填小节缺失：整节省略（正文不出现该节头），fields.json 仍记 missing（required=false）；
- 必填缺失以 ____ 占位（如标题「____工作周报」「____2026年9月工作总结」）；
- 小节编号「一、二、三、」条目编号「1．2．3．」；
- 落款右对齐，不因空日期/缺主体破坏。
"""

import argparse
import json
import sys
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Cm, Pt

PLACEHOLDER = "____"
FULLWIDTH_DOT = "\uff0e"        # 条目编号全角句点：1．
IDEOGRAPHIC_COMMA = "\u3001"    # 小节编号顿号：一、
CN_NUMERALS = ["一", "二", "三", "四", "五", "六", "七", "八", "九", "十"]

BODY_PT = 16        # 三号
LINE_PT = 28        # 公文常用固定行距

# 模板字段定义：kind = scalar|list；required 必填/选填
SCHEMA = {
    "周报": {
        "fields": {
            "填报人": {"kind": "scalar", "required": True},
            "周期": {"kind": "scalar", "required": True},
            "本周工作内容": {"kind": "list", "required": True},
            "下周工作计划": {"kind": "list", "required": True},
            "问题与需协调事项": {"kind": "list", "required": False},
            "报送日期": {"kind": "scalar", "required": True},
        },
        "sections": ["本周工作内容", "下周工作计划", "问题与需协调事项"],
        # 落款行：(字段名, 行内标签)；标签 None 表示直接渲染字段文本
        "signature": [("填报人", "填报人"), ("报送日期", "报送日期")],
    },
    "工作总结": {
        "fields": {
            "总结主体": {"kind": "scalar", "required": True},
            "总结时段": {"kind": "scalar", "required": True},
            "工作回顾": {"kind": "list", "required": True},
            "主要成绩": {"kind": "list", "required": True},
            "存在问题": {"kind": "list", "required": False},
            "下一步工作打算": {"kind": "list", "required": True},
            "成文日期": {"kind": "scalar", "required": False},
        },
        "sections": ["工作回顾", "主要成绩", "存在问题", "下一步工作打算"],
        "signature": [("总结主体", None), ("成文日期", None)],
    },
}


def item_label(i):
    return f"{i}{FULLWIDTH_DOT}"


def section_label(i):
    return f"{CN_NUMERALS[i]}{IDEOGRAPHIC_COMMA}"


def _is_blank(v):
    return v is None or (isinstance(v, str) and v.strip() == "")


# ---------------------------------------------------------------- 字段解析
def resolve_scalar(raw, required):
    """标量字段：filled / placeholder（必填缺失）/ missing-optional-omitted（选填缺失）。"""
    if _is_blank(raw):
        if required:
            return {
                "status": "placeholder",
                "value": raw if isinstance(raw, str) else None,
                "text": PLACEHOLDER,
                "placeholder": PLACEHOLDER,
                "note": "必填缺失，以____占位",
            }
        return {
            "status": "missing-optional-omitted",
            "value": None,
            "text": None,
            "note": "选填缺失，整处省略不渲染（required=false）",
        }
    return {"status": "filled", "value": str(raw), "text": str(raw)}


def resolve_list(raw, required):
    """列表字段：filled / empty-list-placeholder / normalized-single-item / missing-optional-omitted。"""
    if isinstance(raw, str) and not _is_blank(raw):
        # 字符串型列表：按单项化处理
        return {
            "status": "normalized-single-item",
            "value": [raw],
            "raw": raw,
            "items": [f"{item_label(1)}{raw}"],
            "note": "字符串型列表按单项化处理，字段值为单元素数组",
        }
    if isinstance(raw, list) and raw:
        values = [str(x) for x in raw]
        return {
            "status": "filled",
            "value": values,
            "items": [f"{item_label(i)}{t}" for i, t in enumerate(values, 1)],
        }
    # 空列表 / 缺失
    if required:
        # 空列表按必填缺失处理：渲染占位条目
        placeholder_item = f"{item_label(1)}{PLACEHOLDER}"
        return {
            "status": "empty-list-placeholder",
            "value": [] if isinstance(raw, list) else None,
            "items": [placeholder_item],
            "placeholder_item": placeholder_item,
            "note": "空列表按必填缺失处理，渲染占位条目" if isinstance(raw, list)
            else "必填列表缺失，渲染占位条目",
        }
    return {
        "status": "missing-optional-omitted",
        "value": None,
        "items": None,
        "note": "选填小节缺失，整节省略（正文不出现该节头），required=false",
    }


def resolve_all(template, data):
    schema = SCHEMA[template]
    resolved = {}
    for name, spec in schema["fields"].items():
        raw = data.get(name)
        if spec["kind"] == "scalar":
            resolved[name] = resolve_scalar(raw, spec["required"])
        else:
            resolved[name] = resolve_list(raw, spec["required"])
    return resolved


def build_title(template, resolved):
    if template == "周报":
        return f"{resolved['周期']['text']}工作周报"
    return f"{resolved['总结主体']['text']}{resolved['总结时段']['text']}工作总结"


def build_sections(template, resolved):
    """按模板小节顺序渲染实际出现的小节，编号连续「一、二、三、」；选填缺失整节省略。"""
    rendered, omitted = [], []
    for name in SCHEMA[template]["sections"]:
        r = resolved[name]
        if r["items"] is None:
            omitted.append({
                "field": name,
                "header": name,
                "reason": "选填缺失，整节省略（正文不出现该节头）",
            })
        else:
            label = section_label(len(rendered))
            rendered.append({
                "field": name,
                "number": label.rstrip(IDEOGRAPHIC_COMMA),
                "header": f"{label}{name}",
                "items": r["items"],
            })
    return rendered, omitted


def build_signature(template, resolved):
    """落款行：字段缺失时必填以____占位、选填省略，右对齐不破坏。"""
    lines = []
    for field, label in SCHEMA[template]["signature"]:
        r = resolved[field]
        if r["text"] is None:
            continue
        lines.append(f"{label}：{r['text']}" if label else r["text"])
    return lines


# ---------------------------------------------------------------- docx 排版
def _set_ind(paragraph, first_line_chars=None, right_chars=None):
    pPr = paragraph._p.get_or_add_pPr()
    ind = pPr.find(qn("w:ind"))
    if ind is None:
        ind = pPr.makeelement(qn("w:ind"), {})
        pPr.append(ind)
    if first_line_chars is not None:
        ind.set(qn("w:firstLineChars"), str(first_line_chars))
        ind.set(qn("w:firstLine"), str(int(first_line_chars / 100 * BODY_PT * 20)))
    if right_chars is not None:
        ind.set(qn("w:rightChars"), str(right_chars))
        ind.set(qn("w:right"), str(int(right_chars / 100 * BODY_PT * 20)))


def add_paragraph(doc, text, cjk_font, size_pt, *, bold=False, align=None,
                  first_line_chars=None, right_chars=None, line_pt=LINE_PT,
                  space_before=0, space_after=0):
    p = doc.add_paragraph()
    pf = p.paragraph_format
    if align is not None:
        pf.alignment = align
    pf.line_spacing = Pt(line_pt)
    pf.space_before = Pt(space_before)
    pf.space_after = Pt(space_after)
    if first_line_chars is not None or right_chars is not None:
        _set_ind(p, first_line_chars, right_chars)
    run = p.add_run(text)
    run.font.name = "Times New Roman"
    run.font.size = Pt(size_pt)
    run.font.bold = bold
    rPr = run._element.get_or_add_rPr()
    rPr.get_or_add_rFonts().set(qn("w:eastAsia"), cjk_font)
    return p


def render_docx(template, title, sections, signature_lines, out_path):
    doc = Document()
    section = doc.sections[0]
    section.page_width, section.page_height = Cm(21.0), Cm(29.7)   # A4
    section.top_margin, section.bottom_margin = Cm(3.7), Cm(3.5)
    section.left_margin, section.right_margin = Cm(2.8), Cm(2.6)

    normal = doc.styles["Normal"]
    normal.font.name = "Times New Roman"
    normal.font.size = Pt(BODY_PT)
    normal.element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:eastAsia"), "仿宋")

    doc.core_properties.title = title

    # 标题：二号宋体加粗居中
    add_paragraph(doc, title, "宋体", 22, bold=True,
                  align=WD_ALIGN_PARAGRAPH.CENTER, line_pt=33, space_after=12)
    # 正文小节：小节头黑体三号，条目仿宋三号，均首行缩进两字
    for sec in sections:
        add_paragraph(doc, sec["header"], "黑体", BODY_PT,
                      first_line_chars=200, space_before=6)
        for item in sec["items"]:
            add_paragraph(doc, item, "仿宋", BODY_PT, first_line_chars=200)
    # 落款：空一行后右对齐（右缩进两字），不因空日期/缺主体破坏
    add_paragraph(doc, "", "仿宋", BODY_PT)
    for line in signature_lines:
        add_paragraph(doc, line, "仿宋", BODY_PT,
                      align=WD_ALIGN_PARAGRAPH.RIGHT, right_chars=200)

    doc.save(out_path)


# ---------------------------------------------------------------- 主流程
def main():
    ap = argparse.ArgumentParser(description="周报/工作总结 文书生成（列表边界行为）")
    ap.add_argument("--template", required=True, choices=sorted(SCHEMA))
    ap.add_argument("--data", required=True, type=Path)
    ap.add_argument("--outdir", required=True, type=Path)
    args = ap.parse_args()

    data = json.loads(args.data.read_text(encoding="utf-8"))
    if not isinstance(data, dict):
        sys.exit(f"数据 json 顶层必须是对象：{args.data}")

    template = args.template
    resolved = resolve_all(template, data)
    title = build_title(template, resolved)
    sections, omitted = build_sections(template, resolved)
    signature_lines = build_signature(template, resolved)

    unused = {k: v for k, v in data.items() if k not in SCHEMA[template]["fields"]}

    fields_record = []
    for name, spec in SCHEMA[template]["fields"].items():
        r = resolved[name]
        entry = {
            "name": name,
            "kind": spec["kind"],
            "required": spec["required"],
            "status": r["status"],
            "value": r["value"],
        }
        for extra in ("placeholder", "placeholder_item", "raw", "note"):
            if extra in r:
                entry[extra] = r[extra]
        fields_record.append(entry)

    record = {
        "template": template,
        "data": str(args.data),
        "title": title,
        "fields": fields_record,
        "sections": sections,
        "omitted_sections": omitted,
        "signature": {"alignment": "right", "lines": signature_lines},
    }
    if unused:
        record["unused_data"] = unused

    args.outdir.mkdir(parents=True, exist_ok=True)
    render_docx(template, title, sections, signature_lines, args.outdir / "文书.docx")
    (args.outdir / "fields.json").write_text(
        json.dumps(record, ensure_ascii=False, indent=2) + "\n", encoding="utf-8"
    )
    print(f"[ok] {template} -> {args.outdir}/{{文书.docx,fields.json}} 标题「{title}」")


if __name__ == "__main__":
    main()
