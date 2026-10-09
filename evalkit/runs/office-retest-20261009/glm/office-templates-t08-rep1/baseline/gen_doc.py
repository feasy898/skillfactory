#!/usr/bin/env python
# -*- coding: utf-8 -*-
"""gen_doc.py — 工作总结文书生成（本目录 brief 的范围内任务）

读入数据 JSON，按 GB/T 9704 公文风格生成 `文书.docx`，并输出字段填充台账
`fields.json`（哪些字段填了、哪些缺了）。缺数据不编造：必填缺失以 ____ 占位
并如实记账，选填缺失留空或整节省略。

版式（对齐公文惯例）：
  V1  A4，版心 上3.7/下3.5/左2.8/右2.6 cm
  V2  标题居中，黑体二号（22pt），文案 = 总结主体+总结时段+工作总结
  V4  正文仿宋三号（16pt），西文 Times New Roman，固定行距 28 磅，首行缩进两字符
  V5  落款两行（总结主体、成文日期）右对齐，右缩进两字
  V6  小节节头黑体三号，条目以 1．2．3．（全角点）编号

用法：
  /usr/bin/python3 gen_doc.py --data data/工作总结/case1-overflow.json --outdir out

退出码：0 成功；2 输入非法（文件不存在 / 非 UTF-8 / 非法 JSON / 顶层非对象）。
确定性：正文无当前时间、无随机源；docx zip 条目时间戳归一化，同输入重跑逐字节一致。
"""

from __future__ import annotations

import argparse
import json
import sys
import zipfile
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Cm, Pt

try:  # Windows 控制台缺省 GBK，强制 UTF-8 输出
    sys.stdout.reconfigure(encoding="utf-8", errors="replace")
    sys.stderr.reconfigure(encoding="utf-8", errors="replace")
except Exception:  # noqa: BLE001 — 非 Windows / 已重配置时忽略
    pass

DOCX_NAME = "文书.docx"
FIELDS_NAME = "fields.json"
TEMPLATE_NAME = "工作总结"
PLACEHOLDER = "____"          # 必填缺失占位符
CN_NUMS = ("一", "二", "三", "四", "五", "六", "七", "八", "九", "十")

TITLE_FONT = "黑体"           # 标题/节头 eastAsia 字体
TITLE_SIZE = Pt(22)           # 二号
BODY_FONT = "仿宋"            # 正文 eastAsia 字体
ASCII_FONT = "Times New Roman"  # 正文西文字体
BODY_SIZE = Pt(16)            # 三号
LINE_SPACING = Pt(28)         # 固定行距
SIGN_RIGHT_INDENT = Pt(32)    # 落款右缩进两字

ZIP_FIXED_DATE = (1980, 1, 1, 0, 0, 0)  # zip 条目时间戳归一化值（DOS 下限）

# 工作总结字段表：(名称, 必填, 类型)，顺序冻结
FIELDS_SPEC = [
    ("总结主体", True, "text"),
    ("总结时段", True, "text"),
    ("工作回顾", True, "list"),
    ("主要成绩", True, "list"),
    ("存在问题", False, "list"),
    ("下一步工作打算", True, "list"),
    ("成文日期", False, "text"),
]
FIELD_NAMES = {n for n, _r, _k in FIELDS_SPEC}

RENDER_REQ_TEXT = "以____占位"
RENDER_OPT_TEXT = "留空（渲染为空字符串）"
RENDER_REQ_LIST = "以____占位条目"
RENDER_OPT_LIST = "省略该条目/小节"


def norm_text(raw):
    """文本字段规整：None→缺失；字符串→strip；其他类型→str().strip()。"""
    if raw is None:
        return False, ""
    s = raw.strip() if isinstance(raw, str) else str(raw).strip()
    return bool(s), s


def norm_list(raw):
    """列表字段规整：字符串→单项；数组→逐项 str().strip() 去空；空/全空→缺失。"""
    if raw is None:
        return False, []
    if isinstance(raw, str):
        s = raw.strip()
        return bool(s), [s] if s else []
    if isinstance(raw, (list, tuple)):
        items = [str(x).strip() for x in raw]
        items = [x for x in items if x]
        return bool(items), items
    s = str(raw).strip()
    return bool(s), [s] if s else []


def build_fields(data):
    """按字段表逐项记账。返回 (明细列表, 已填值字典)。"""
    details, values = [], {}
    for name, req, kind in FIELDS_SPEC:
        ok, val = norm_text(data.get(name)) if kind == "text" else norm_list(data.get(name))
        if ok:
            details.append({"name": name, "required": req, "kind": kind,
                            "status": "filled", "value": val})
            values[name] = val
        else:
            rendered = (RENDER_REQ_TEXT if req else RENDER_OPT_TEXT) if kind == "text" \
                else (RENDER_REQ_LIST if req else RENDER_OPT_LIST)
            details.append({"name": name, "required": req, "kind": kind,
                            "status": "missing", "value": None, "rendered_as": rendered})
    return details, values


def pick(values, name, required):
    """取渲染用文本：必填缺失→____；选填缺失→空字符串。"""
    v = values.get(name)
    if v is None:
        return PLACEHOLDER if required else ""
    return v


def build_title(values):
    """标题文案 = 总结主体 + 总结时段 + 工作总结（缺失必填以 ____ 代入）。"""
    return "%s%s工作总结" % (pick(values, "总结主体", True), pick(values, "总结时段", True))


# ---------------------------------------------------------------- docx 版式

def set_page(section):
    section.page_width = Cm(21.0)
    section.page_height = Cm(29.7)
    section.top_margin = Cm(3.7)
    section.bottom_margin = Cm(3.5)
    section.left_margin = Cm(2.8)
    section.right_margin = Cm(2.6)


def style_run(run, east, size):
    """run 字体声明：eastAsia=east；正文西文 TNR，其余西文与 east 同名。"""
    run.font.size = size
    run.font.name = ASCII_FONT if east == BODY_FONT else east
    run._element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:eastAsia"), east)


FIRST_LINE_INDENT = Pt(32)   # 首行缩进两字符的绝对值（16pt 字号 × 2）


def set_para_base(p, indent=True):
    """段落基础版式：固定行距 28 磅、段前段后 0；indent 时首行缩进两字符。
    同时写 w:firstLineChars=200（Word 按字号自适应，优先级更高）与
    w:firstLine=32pt（LibreOffice 等不识别 chars 变体的渲染器回退）。"""
    pf = p.paragraph_format
    pf.line_spacing = LINE_SPACING
    pf.space_before = Pt(0)
    pf.space_after = Pt(0)
    if indent:
        pf.first_line_indent = FIRST_LINE_INDENT
        p._p.get_or_add_pPr().get_or_add_ind().set(qn("w:firstLineChars"), "200")


def add_title(doc, text):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    style_run(p.add_run(text), TITLE_FONT, TITLE_SIZE)


def add_body(doc, text, indent=True):
    p = doc.add_paragraph()
    set_para_base(p, indent)
    if text:
        style_run(p.add_run(text), BODY_FONT, BODY_SIZE)


def add_section_header(doc, text):
    p = doc.add_paragraph()
    set_para_base(p, True)
    style_run(p.add_run(text), TITLE_FONT, BODY_SIZE)


def add_signature(doc, text):
    """落款行：右对齐 + 右缩进两字。text 为空时仍产出空段（不破坏行序）。"""
    p = doc.add_paragraph()
    pf = p.paragraph_format
    pf.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    pf.right_indent = SIGN_RIGHT_INDENT
    pf.line_spacing = LINE_SPACING
    pf.space_before = Pt(0)
    pf.space_after = Pt(0)
    if text:
        style_run(p.add_run(text), BODY_FONT, BODY_SIZE)


def emit_list(doc, values, name):
    """条目以 1．2．3．编号；必填列表缺失→单个占位条目 1．____。"""
    items = values.get(name) or [PLACEHOLDER]
    for i, item in enumerate(items, 1):
        add_body(doc, "%d．%s" % (i, item))


def build_doc(title, values):
    """段落顺序：标题 → 各小节 → 落款两行。"""
    doc = Document()
    set_page(doc.sections[0])
    add_title(doc, title)
    add_section_header(doc, "%s、工作回顾" % CN_NUMS[0])
    emit_list(doc, values, "工作回顾")
    add_section_header(doc, "%s、主要成绩" % CN_NUMS[1])
    emit_list(doc, values, "主要成绩")
    if "存在问题" in values:          # 选填小节缺失整节省略
        add_section_header(doc, "%s、存在问题" % CN_NUMS[2])
        emit_list(doc, values, "存在问题")
    add_section_header(doc, "%s、下一步工作打算" % CN_NUMS[3])
    emit_list(doc, values, "下一步工作打算")
    add_signature(doc, pick(values, "总结主体", True))
    add_signature(doc, pick(values, "成文日期", False))
    return doc


# ---------------------------------------------------------------- 台账与 zip 归一化

def build_meta(data, title, docx_path, fields_path):
    details, values = build_fields(data)
    filled = [d["name"] for d in details if d["status"] == "filled"]
    missing = [d["name"] for d in details if d["status"] == "missing"]
    missing_required = [d["name"] for d in details
                        if d["status"] == "missing" and d["required"]]
    unknown = sorted(k for k in data.keys() if k not in FIELD_NAMES)
    meta = {
        "template": TEMPLATE_NAME,
        "title": title,
        "filled_fields": filled,
        "missing_fields": missing,
        "fields": details,
        "summary": {
            "total": len(details),
            "filled": len(filled),
            "missing": len(missing),
            "missing_required": len(missing_required),
            "missing_required_names": missing_required,
            "unknown_keys": unknown,
        },
        "outputs": {"docx": str(docx_path), "fields_json": str(fields_path)},
    }
    return meta, values


def canonicalize_docx_zip(path):
    """把 docx zip 条目时间戳归一化，使整文件逐字节可复现；失败保留原文件。"""
    tmp = path.with_name(path.name + ".canon.tmp")
    try:
        with zipfile.ZipFile(path) as zin:
            items = [(info, zin.read(info.filename)) for info in zin.infolist()]
        with zipfile.ZipFile(tmp, "w") as zout:
            for info, data in items:
                zi = zipfile.ZipInfo(info.filename, date_time=ZIP_FIXED_DATE)
                zi.compress_type = info.compress_type
                zi.external_attr = info.external_attr
                zi.internal_attr = info.internal_attr
                zi.create_system = info.create_system
                zout.writestr(zi, data)
        tmp.replace(path)
        return True
    except OSError:
        try:
            tmp.unlink()
        except OSError:
            pass
        return False


def fail(msg):
    print("错误：%s" % msg, file=sys.stderr)
    return 2


def main(argv=None):
    ap = argparse.ArgumentParser(description="工作总结文书生成（docx + 字段台账）")
    ap.add_argument("--data", required=True, help="数据 JSON 文件路径（UTF-8，顶层为对象）")
    ap.add_argument("--outdir", required=True, help="输出目录（不存在则递归创建）")
    args = ap.parse_args(argv)

    # 先完成全部输入校验，再触碰输出文件系统（零产物残留）
    data_path = Path(args.data)
    if not data_path.is_file():
        return fail("数据文件不存在：%s" % args.data)
    try:
        text = data_path.read_text(encoding="utf-8-sig")
    except OSError as e:
        return fail("数据文件不可读：%s（%s）" % (args.data, e))
    except UnicodeDecodeError as e:
        return fail("数据文件不是 UTF-8 编码：%s（%s）" % (args.data, e))
    try:
        data = json.loads(text)
    except ValueError as e:
        return fail("数据文件不是合法 JSON：%s（%s）" % (args.data, e))
    if not isinstance(data, dict):
        return fail("数据 JSON 顶层必须是对象（键值对），实际为 %s" % type(data).__name__)

    outdir = Path(args.outdir)
    docx_path = outdir / DOCX_NAME
    fields_path = outdir / FIELDS_NAME

    details, _ = build_fields(data)
    title = build_title({d["name"]: d["value"] for d in details if d["status"] == "filled"})
    meta, values = build_meta(data, title, docx_path, fields_path)
    doc = build_doc(title, values)

    try:
        outdir.mkdir(parents=True, exist_ok=True)
        doc.save(str(docx_path))
        canonicalize_docx_zip(docx_path)
        fields_path.write_text(
            json.dumps(meta, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    except OSError as e:
        return fail("写入产物失败：%s（%s）" % (args.outdir, e))

    s = meta["summary"]
    print("已生成：%s（模板=%s，标题=%s，字段 filled/total=%d/%d，"
          "必填缺失=%s，模板外键=%s）"
          % (docx_path, TEMPLATE_NAME, title, s["filled"], s["total"],
             s["missing_required_names"] or "无", s["unknown_keys"] or "无"))
    return 0


if __name__ == "__main__":
    sys.exit(main())
