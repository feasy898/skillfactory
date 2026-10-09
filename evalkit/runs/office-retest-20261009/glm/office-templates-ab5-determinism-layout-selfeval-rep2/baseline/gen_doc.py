#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""会议通知文书生成器（确定性实现）。

命令行契约:
    python3 gen_doc.py --template 会议通知 --data <数据json路径> --outdir <输出目录>

每次调用在 --outdir 下生成两个文件:
    文书.docx    按「会议通知」公文版式排版的会议通知
    fields.json  记录本次使用的模板、输入数据及其哈希、产物路径

确定性保证（同输入连跑两遍产物逐字节一致）:
    - 文档内容仅由输入数据字段决定，无当前时间、无随机内容
    - docProps/core.xml 的创建/修改时间固定为 2000-01-01T00:00:00Z
    - 产物 zip 包内所有条目的时间戳/外部属性/压缩方式统一归一化
    - fields.json 以固定序列化参数写出（sort_keys、固定缩进、UTF-8、末尾换行）
"""

import argparse
import hashlib
import json
import os
import sys
import zipfile
from datetime import datetime, timezone
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.oxml.ns import qn
from docx.shared import Cm, Pt

TEMPLATE_NAME = "会议通知"
DOCX_NAME = "文书.docx"
FIELDS_NAME = "fields.json"

# 固定时间常量：用于 zip 条目时间戳与文档核心属性，杜绝运行时刻进入产物
ZIP_EPOCH = (2000, 1, 1, 0, 0, 0)
CORE_EPOCH = datetime(2000, 1, 1, 0, 0, 0, tzinfo=timezone.utc)

# 公文版式常量：二号方正小标宋标题、三号仿宋_GB2312 正文、固定行距 28 磅、首行缩进两字符
TITLE_FONT_EAST_ASIA = "方正小标宋简体"
BODY_FONT_EAST_ASIA = "仿宋_GB2312"
WESTERN_FONT = "Times New Roman"
TITLE_SIZE = Pt(22)          # 二号
BODY_SIZE = Pt(16)           # 三号
TITLE_LINE_SPACING = Pt(33)
BODY_LINE_SPACING = Pt(28)
FIRST_LINE_INDENT = Pt(32)   # 三号字两字符 = 16pt × 2

REQUIRED_FIELDS = [
    "会议名称", "召开单位", "主送对象", "会议时间", "会议地点",
    "参会人员", "会议议题", "会议要求", "联系人", "联系电话", "发文日期",
]

_DIGITS = "零一二三四五六七八九"


def to_chinese_num(n: int) -> str:
    """1..99 的确定性中文序号（用于会议要求条目编号）。"""
    if n < 1 or n > 99:
        raise ValueError(f"序号超出支持范围: {n}")
    if n < 10:
        return _DIGITS[n]
    if n == 10:
        return "十"
    if n < 20:
        return "十" + _DIGITS[n % 10]
    tens, ones = divmod(n, 10)
    return _DIGITS[tens] + "十" + (_DIGITS[ones] if ones else "")


def style_run(run, east_asia, size, bold=False):
    """显式设置 run 的西文字体、中文字体（w:eastAsia）与字号。"""
    run.font.name = WESTERN_FONT
    run.font.size = size
    run.font.bold = bold
    rpr = run._element.get_or_add_rPr()
    rfonts = rpr.get_or_add_rFonts()
    rfonts.set(qn("w:eastAsia"), east_asia)


def add_para(doc, text, *, east_asia=BODY_FONT_EAST_ASIA, size=BODY_SIZE, bold=False,
             align=WD_ALIGN_PARAGRAPH.LEFT, first_line_indent=None,
             space_before=Pt(0), space_after=Pt(0), line_spacing=BODY_LINE_SPACING):
    p = doc.add_paragraph()
    pf = p.paragraph_format
    pf.alignment = align
    pf.space_before = space_before
    pf.space_after = space_after
    pf.line_spacing_rule = WD_LINE_SPACING.EXACTLY
    pf.line_spacing = line_spacing
    if first_line_indent is not None:
        pf.first_line_indent = first_line_indent
    style_run(p.add_run(text), east_asia, size, bold)
    return p


def build_document(data: dict) -> Document:
    doc = Document()

    # Normal 样式兜底：正文默认西文字体 + 中文字体
    normal = doc.styles["Normal"]
    normal.font.name = WESTERN_FONT
    normal.font.size = BODY_SIZE
    normal_rpr = normal.element.get_or_add_rPr()
    normal_rpr.get_or_add_rFonts().set(qn("w:eastAsia"), BODY_FONT_EAST_ASIA)

    # 页面：A4 + 公文页边距（上 3.7 / 下 3.5 / 左 2.8 / 右 2.6 cm）
    section = doc.sections[0]
    section.page_width = Cm(21.0)
    section.page_height = Cm(29.7)
    section.top_margin = Cm(3.7)
    section.bottom_margin = Cm(3.5)
    section.left_margin = Cm(2.8)
    section.right_margin = Cm(2.6)

    title_text = f"关于召开“{data['会议名称']}”的通知"
    add_para(doc, title_text, east_asia=TITLE_FONT_EAST_ASIA, size=TITLE_SIZE,
             align=WD_ALIGN_PARAGRAPH.CENTER, space_after=Pt(18),
             line_spacing=TITLE_LINE_SPACING)

    add_para(doc, f"{data['主送对象']}：")

    add_para(doc,
             f"经研究，定于{data['会议时间']}在{data['会议地点']}召开{data['会议名称']}，"
             f"请{data['参会人员']}准时参加。",
             first_line_indent=FIRST_LINE_INDENT)
    add_para(doc, f"本次会议议题：{data['会议议题']}。", first_line_indent=FIRST_LINE_INDENT)
    add_para(doc, "会议要求如下：", first_line_indent=FIRST_LINE_INDENT)

    requirements = data["会议要求"]
    if isinstance(requirements, str):
        requirements = [requirements]
    for i, item in enumerate(requirements):
        suffix = "。" if i == len(requirements) - 1 else "；"
        add_para(doc, f"{to_chinese_num(i + 1)}、{item}{suffix}",
                 first_line_indent=FIRST_LINE_INDENT)

    add_para(doc, f"会务联系：{data['联系人']}，联系电话：{data['联系电话']}。",
             first_line_indent=FIRST_LINE_INDENT)

    add_para(doc, data["召开单位"], align=WD_ALIGN_PARAGRAPH.RIGHT, space_before=Pt(28))
    add_para(doc, data["发文日期"], align=WD_ALIGN_PARAGRAPH.RIGHT)

    # 核心属性全部固定/由输入派生，不引入运行时刻
    cp = doc.core_properties
    cp.title = title_text
    cp.subject = TEMPLATE_NAME
    cp.author = data["召开单位"]
    cp.last_modified_by = data["召开单位"]
    cp.created = CORE_EPOCH
    cp.modified = CORE_EPOCH
    cp.revision = 1
    cp.category = "通知"
    cp.comments = ""
    return doc


def normalize_zip(path: Path) -> None:
    """重写 zip 容器，把每个条目的时间戳/属性/压缩方式归一化为固定值。"""
    with zipfile.ZipFile(path) as zin:
        entries = [(info.filename, zin.read(info.filename)) for info in zin.infolist()]
    tmp = path.parent / (path.name + ".normalizing")
    with zipfile.ZipFile(tmp, "w") as zout:
        for name, blob in entries:
            info = zipfile.ZipInfo(name, date_time=ZIP_EPOCH)
            info.compress_type = zipfile.ZIP_DEFLATED
            info.create_system = 3
            info.external_attr = 0o644 << 16
            zout.writestr(info, blob)
    os.replace(tmp, path)


def sha256_bytes(blob: bytes) -> str:
    return hashlib.sha256(blob).hexdigest()


def main(argv=None) -> int:
    parser = argparse.ArgumentParser(
        description="会议通知文书生成器：--template 会议通知 --data <数据json> --outdir <输出目录>")
    parser.add_argument("--template", required=True, help="模板名（仅支持：会议通知）")
    parser.add_argument("--data", required=True, help="输入数据 json 路径")
    parser.add_argument("--outdir", required=True, help="输出目录（每次生成 文书.docx 与 fields.json）")
    args = parser.parse_args(argv)

    if args.template != TEMPLATE_NAME:
        print(f"错误：不支持的模板 {args.template!r}（本生成器仅支持「{TEMPLATE_NAME}」）", file=sys.stderr)
        return 2

    data_path = Path(args.data)
    if not data_path.is_file():
        print(f"错误：数据文件不存在：{args.data}", file=sys.stderr)
        return 2
    raw = data_path.read_bytes()
    try:
        data = json.loads(raw.decode("utf-8"))
    except (UnicodeDecodeError, json.JSONDecodeError) as exc:
        print(f"错误：数据文件解析失败：{exc}", file=sys.stderr)
        return 2

    missing = [k for k in REQUIRED_FIELDS if k not in data]
    if missing:
        print(f"错误：数据缺少必需字段：{missing}", file=sys.stderr)
        return 2

    outdir = Path(args.outdir)
    outdir.mkdir(parents=True, exist_ok=True)

    doc = build_document(data)
    docx_path = outdir / DOCX_NAME
    doc.save(str(docx_path))
    normalize_zip(docx_path)

    fields_payload = {
        "template": TEMPLATE_NAME,
        "data": {"path": args.data, "sha256": sha256_bytes(raw)},
        "fields": data,
        "outputs": {
            "docx": os.path.join(args.outdir, DOCX_NAME),
            "fields": os.path.join(args.outdir, FIELDS_NAME),
        },
    }
    fields_text = json.dumps(fields_payload, ensure_ascii=False, indent=2, sort_keys=True) + "\n"
    fields_path = outdir / FIELDS_NAME
    fields_path.write_text(fields_text, encoding="utf-8")

    print(f"OK 模板={args.template} 数据={args.data} 输出目录={args.outdir}")
    print(f"生成 {DOCX_NAME}（{docx_path.stat().st_size} 字节）、{FIELDS_NAME}（{fields_path.stat().st_size} 字节）")
    return 0


if __name__ == "__main__":
    sys.exit(main())
