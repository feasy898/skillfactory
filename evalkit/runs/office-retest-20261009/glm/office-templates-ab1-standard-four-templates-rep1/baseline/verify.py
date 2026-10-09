#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""交付自检：四类模板产物齐全性、fields.json 结构、docx 内容与 GB/T 9704 版式。"""
import json
import sys
from pathlib import Path

sys.path.insert(0, "/home/pan-ding/.local/lib/python3.12/site-packages")
from docx import Document
from docx.shared import Mm, Pt
from docx.oxml.ns import qn

BASE = Path(__file__).resolve().parent
TEMPLATES = ["周报", "请示函", "会议通知", "工作总结"]
fail = []


def check(cond, msg):
    print(("PASS " if cond else "FAIL ") + msg)
    if not cond:
        fail.append(msg)


for t in TEMPLATES:
    d = BASE / "out" / t / "case1"
    files = sorted(p.name for p in d.iterdir())
    check(files == sorted(["文书.docx", "fields.json"]),
          "%s/case1 恰含 文书.docx 与 fields.json: %s" % (t, files))

    fj = json.loads((d / "fields.json").read_text(encoding="utf-8"))
    for key in ("title", "fields", "summary"):
        check(key in fj, "%s fields.json 顶层含 %s" % (t, key))
    for f in fj["fields"]:
        check(all(k in f for k in ("名称", "是否必填", "状态", "值", "渲染方式")),
              "%s 字段 %s 五要素齐全" % (t, f["名称"]))
    s = fj["summary"]
    filled = sum(1 for f in fj["fields"] if f["状态"] == "filled")
    check(s["filled"] == filled and s["missing"] == len(fj["fields"]) - filled
          and s["total"] == len(fj["fields"]),
          "%s summary 计数自洽 filled=%d missing=%d total=%d" % (t, s["filled"], s["missing"], s["total"]))

    doc = Document(str(d / "文书.docx"))
    text = "\n".join(p.text for p in doc.paragraphs)
    flat = text.replace("\n", "")  # 标题回行等换行不影响内容判定

    data = json.loads((BASE / "fixtures" / t / "case1.json").read_text(encoding="utf-8"))
    for k, v in data.items():
        vals = v if isinstance(v, list) else [v]
        for item in vals:
            check(str(item) in flat, "%s docx 含字段 %s 的值 %r" % (t, k, item))

    sec = doc.sections[0]
    check(abs(sec.page_width.mm - 210) < 0.01 and abs(sec.page_height.mm - 297) < 0.01,
          "%s A4 版面" % t)
    check(abs(sec.top_margin.mm - 37) < 0.01 and abs(sec.bottom_margin.mm - 35) < 0.01
          and abs(sec.left_margin.mm - 28) < 0.01 and abs(sec.right_margin.mm - 26) < 0.01,
          "%s 页边距 37/35/28/26mm" % t)

    grid = sec._sectPr.find(qn("w:docGrid"))
    check(grid is not None and grid.get(qn("w:type")) == "lines"
          and grid.get(qn("w:linePitch")) == "579",
          "%s 版式网格 22 行/页" % t)

    title_p = doc.paragraphs[0]
    tr = title_p.runs[0]
    check(abs(tr.font.size.pt - 22) < 0.01, "%s 标题 2号(22pt)" % t)
    rf = tr._element.rPr.find(qn("w:rFonts"))
    check(rf.get(qn("w:eastAsia")) == "方正小标宋简体", "%s 标题小标宋" % t)

    body_runs = [r for p in doc.paragraphs for r in p.runs
                 if r.font.size and abs(r.font.size.pt - 16) < 0.01]
    check(body_runs and all(
        (r._element.rPr.find(qn("w:rFonts")).get(qn("w:eastAsia")) or "") in
        ("仿宋_GB2312", "黑体") for r in body_runs), "%s 正文 3号 仿宋/黑体" % t)

    ls = {p.paragraph_format.line_spacing for p in doc.paragraphs
          if p.paragraph_format.line_spacing and p.text}
    check(ls and all(abs(v.pt - 28.95) < 0.6 or abs(v.pt - 33) < 0.6 for v in ls),
          "%s 行距 28.95pt(标题 33pt): %s" % (t, sorted({round(v.pt, 2) for v in ls})))

    print("  ——— %s 完整正文 ———" % t)
    for p in doc.paragraphs:
        if p.text:
            print("  | " + p.text)
    print()

print("=" * 40)
print("EXIT=%d, FAIL=%d" % (1 if fail else 0, len(fail)))
sys.exit(1 if fail else 0)
