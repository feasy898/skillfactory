#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""公文产物生成器（GB/T 9704 公文风格，A4 版面）。

命令行契约：
    python gen_doc.py --template <周报|请示函|会议通知|工作总结> \
                      --data <数据json路径> --outdir <输出目录>

每次调用在 outdir 下生成两个产物：
    文书.docx   —— GB/T 9704 风格：A4（210×297mm），天头37mm、地脚35mm、
                   订口28mm、翻口26mm；二号小标宋标题居中，三号仿宋_GB2312
                   正文，首行缩进2字，固定行距28pt（每页22行口径），一级
                   小标题用三号黑体，落款右空二字。
    fields.json —— 记账：顶层 title / template / fields / summary；每个字段
                   含 名称、是否必填、状态、值、渲染方式；summary 为
                   filled / missing / total 计数。

缺数据不编造：必填缺失以 ____ 占位，选填缺失留空或整节省略，均在
fields.json 中如实记账（状态 missing、值 null）。
"""

import argparse
import json
import os

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Mm, Pt

# ---- GB/T 9704 版面常量 -------------------------------------------------
PAGE_W, PAGE_H = Mm(210), Mm(297)      # A4
M_TOP, M_BOTTOM = Mm(37), Mm(35)       # 天头 / 地脚
M_LEFT, M_RIGHT = Mm(28), Mm(26)       # 订口 / 翻口

TITLE_SIZE = Pt(22)                    # 二号（标题）
BODY_SIZE = Pt(16)                     # 三号（正文）
LINE_PITCH = Pt(28)                    # 固定行距，每页 22 行口径
INDENT_2CH = Pt(32)                    # 三号字 2 字符首行缩进
TAIL_2CH = Pt(32)                      # 落款右空 2 字

F_XBS = "方正小标宋简体"                # 标题：二号小标宋
F_FS = "仿宋_GB2312"                   # 正文：三号仿宋
F_HT = "黑体"                          # 一级小标题：三号黑体
ASCII_FONT = "Times New Roman"

PLACEHOLDER = "____"
CN_NUM = "一二三四五六七八九十"

# ---- 字段清单（名称, 是否必填），顺序即记账与行文顺序 ---------------------
SCHEMA = {
    "周报": [
        ("部门", True), ("填报人", True), ("周期", True),
        ("本周工作内容", True), ("下周工作计划", False),
        ("问题与需协调事项", False), ("报送日期", True),
    ],
    "请示函": [
        ("请示事由", True), ("主送机关", True), ("请示缘由", True),
        ("请示事项", True), ("请示单位", True), ("联系人", False),
        ("联系电话", False), ("成文日期", True),
    ],
    "会议通知": [
        ("会议名称", True), ("召开单位", True), ("主送对象", True),
        ("会议时间", True), ("会议地点", True), ("参会人员", True),
        ("会议议题", True), ("会议要求", False), ("联系人", False),
        ("联系电话", False), ("发文日期", True),
    ],
    "工作总结": [
        ("总结主体", True), ("总结时段", True), ("工作回顾", True),
        ("主要成绩", False), ("存在问题", False), ("下一步工作打算", False),
        ("成文日期", True),
    ],
}


def is_present(v):
    return v is not None and v != "" and v != []


class Field:
    """单个字段的取值、在场判定与文档渲染值。"""

    def __init__(self, data, name, required):
        self.name = name
        self.required = required
        self.raw = data.get(name)
        self.present = is_present(self.raw)

    @property
    def value(self):
        """渲染值：在场原样；必填缺失给 ____；选填缺失给 None（由调用方省略）。"""
        if self.present:
            return self.raw
        return PLACEHOLDER if self.required else None


def record(field, where):
    """按 fields.json 口径记账一条字段。"""
    if field.present:
        return {"名称": field.name, "是否必填": field.required,
                "状态": "filled", "值": field.raw, "渲染方式": where}
    if field.required:
        return {"名称": field.name, "是否必填": True, "状态": "missing",
                "值": None,
                "渲染方式": f"必填缺失，以{PLACEHOLDER}占位（{where}）"}
    return {"名称": field.name, "是否必填": False, "状态": "missing",
            "值": None, "渲染方式": f"选填缺失，省略（{where}）"}


class GbDoc:
    """GB/T 9704 版式 docx 薄封装：版心、字体、行距统一在此设置。"""

    def __init__(self):
        self.doc = Document()
        sec = self.doc.sections[0]
        sec.page_width, sec.page_height = PAGE_W, PAGE_H
        sec.top_margin, sec.bottom_margin = M_TOP, M_BOTTOM
        sec.left_margin, sec.right_margin = M_LEFT, M_RIGHT
        normal = self.doc.styles["Normal"]
        normal.font.name = ASCII_FONT
        normal.font.size = BODY_SIZE
        normal.element.get_or_add_rPr().get_or_add_rFonts().set(
            qn("w:eastAsia"), F_FS)

    def _apply(self, pf, align, indent, right_indent):
        pf.line_spacing = LINE_PITCH  # Length 赋值即固定行距（exact）
        if align is not None:
            pf.alignment = align
        if indent:
            pf.first_line_indent = INDENT_2CH
        if right_indent is not None:
            pf.right_indent = right_indent

    def para(self, text="", font=F_FS, size=BODY_SIZE, align=None,
             indent=True, bold=False, right_indent=None):
        p = self.doc.add_paragraph()
        self._apply(p.paragraph_format, align, indent, right_indent)
        if text:
            self._run(p, text, font, size, bold)
        return p

    def mixed(self, parts, align=None, indent=True, right_indent=None):
        """同段落多字体：parts = [(text, eastAsia字体), ...]。"""
        p = self.doc.add_paragraph()
        self._apply(p.paragraph_format, align, indent, right_indent)
        for text, font in parts:
            self._run(p, text, font, BODY_SIZE, False)
        return p

    def _run(self, p, text, font, size, bold):
        r = p.add_run(text)
        r.font.name = ASCII_FONT
        r.font.size = size
        r.font.bold = bold
        r._element.rPr.rFonts.set(qn("w:eastAsia"), font)
        return r

    def title(self, text):
        self.para(text, font=F_XBS, size=TITLE_SIZE, bold=False,
                  align=WD_ALIGN_PARAGRAPH.CENTER, indent=False)

    def heading(self, text):
        self.para(text, font=F_HT)

    def item_list(self, items):
        if isinstance(items, str):
            items = [items]
        for i, item in enumerate(items, 1):
            self.para(f"{i}. {item}")

    def sign_off(self, org, date):
        """落款：署名 + 成文日期，右对齐右空二字。"""
        self.para("")
        if org is not None:
            self.para(org, align=WD_ALIGN_PARAGRAPH.RIGHT, indent=False,
                      right_indent=TAIL_2CH)
        if date is not None:
            self.para(date, align=WD_ALIGN_PARAGRAPH.RIGHT, indent=False,
                      right_indent=TAIL_2CH)

    def contact_line(self, f_contact, f_phone):
        """文末附注行：（联系人：…，联系电话：…）；选填各自缺省即略。"""
        parts = []
        if f_contact.present:
            parts.append(f"联系人：{f_contact.raw}")
        if f_phone.present:
            parts.append(f"联系电话：{f_phone.raw}")
        if parts:
            self.para("")
            self.para(f"（{'，'.join(parts)}）")

    def save(self, path):
        self.doc.save(path)


# ---- 四类文种的组装 -------------------------------------------------------
def build_zhoubao(data, d):
    f = {name: Field(data, name, req) for name, req in SCHEMA["周报"]}
    title = f"{f['部门'].value}工作周报"
    d.title(title)
    d.para(f["周期"].value, align=WD_ALIGN_PARAGRAPH.CENTER, indent=False)
    d.para("")

    sections = [("本周工作内容", f["本周工作内容"]),
                ("下周工作计划", f["下周工作计划"]),
                ("问题与需协调事项", f["问题与需协调事项"])]
    visible = [s for s in sections if s[1].present or s[1].required]
    for i, (label, field) in enumerate(visible):
        header = f"{CN_NUM[i]}、{label}"
        if field.present:
            d.heading(header)
            d.item_list(field.raw)
        else:
            d.heading(header)
            d.para(PLACEHOLDER)

    d.sign_off(f"填报人：{f['填报人'].value}",
               f"报送日期：{f['报送日期'].value}")

    where = {"部门": "公文标题", "填报人": "文末右对齐落款「填报人：…」",
             "周期": "标题下居中周期行", "本周工作内容": "正文小标题下逐条编号列示",
             "下周工作计划": "正文小标题下逐条编号列示",
             "问题与需协调事项": "正文小标题下逐条编号列示",
             "报送日期": "文末右对齐落款「报送日期：…」"}
    return title, [record(f[n], where[n]) for n, _ in SCHEMA["周报"]]


def build_qingshi(data, d):
    f = {name: Field(data, name, req) for name, req in SCHEMA["请示函"]}
    title = f"关于{f['请示事由'].value}的请示"
    d.title(title)
    d.para("")
    d.para(f"{f['主送机关'].value}：", indent=False)          # 主送顶格
    d.para(f["请示缘由"].value)
    d.para(f["请示事项"].value)
    d.para("妥否，请批示。")
    d.sign_off(f["请示单位"].value, f["成文日期"].value)
    d.contact_line(f["联系人"], f["联系电话"])

    where = {"请示事由": "公文标题「关于……的请示」",
             "主送机关": "正文首行顶格主送机关称谓",
             "请示缘由": "正文第 1 段（请示缘由）",
             "请示事项": "正文第 2 段（请示事项）",
             "请示单位": "文末右对齐落款署名",
             "联系人": "文末附注行（联系人：…）",
             "联系电话": "文末附注行（联系电话：…）",
             "成文日期": "文末右对齐落款成文日期"}
    return title, [record(f[n], where[n]) for n, _ in SCHEMA["请示函"]]


def build_tongzhi(data, d):
    f = {name: Field(data, name, req) for name, req in SCHEMA["会议通知"]}
    title = f"关于召开{f['会议名称'].value}的通知"
    d.title(title)
    d.para("")
    d.para(f"{f['主送对象'].value}：", indent=False)          # 主送顶格
    d.para("现将会议有关事项通知如下：")

    # 单值条目（必填）：「一、会议时间：……」；会议要求为选填列表条目
    items = [("会议时间", f["会议时间"]), ("会议地点", f["会议地点"]),
             ("参会人员", f["参会人员"]), ("会议议题", f["会议议题"]),
             ("会议要求", f["会议要求"])]
    visible = [s for s in items if s[1].present or s[1].required]
    for i, (label, field) in enumerate(visible):
        header = f"{CN_NUM[i]}、{label}："
        if field.present and isinstance(field.raw, list):
            d.heading(header)
            d.item_list(field.raw)
        else:
            d.mixed([(header, F_HT), (str(field.value), F_FS)])

    d.para("特此通知。")
    d.sign_off(f["召开单位"].value, f["发文日期"].value)
    d.contact_line(f["联系人"], f["联系电话"])

    where = {"会议名称": "公文标题「关于召开……的通知」",
             "召开单位": "文末右对齐落款署名",
             "主送对象": "正文首行顶格主送对象称谓",
             "会议时间": "正文条目「会议时间：…」（黑体条目名）",
             "会议地点": "正文条目「会议地点：…」（黑体条目名）",
             "参会人员": "正文条目「参会人员：…」（黑体条目名）",
             "会议议题": "正文条目「会议议题：…」（黑体条目名）",
             "会议要求": "正文条目「会议要求：」下逐条编号列示",
             "联系人": "文末附注行（联系人：…）",
             "联系电话": "文末附注行（联系电话：…）",
             "发文日期": "文末右对齐落款成文日期"}
    return title, [record(f[n], where[n]) for n, _ in SCHEMA["会议通知"]]


def build_zongjie(data, d):
    f = {name: Field(data, name, req) for name, req in SCHEMA["工作总结"]}
    title = f"{f['总结主体'].value}{f['总结时段'].value}工作总结"
    d.title(title)
    d.para("")

    sections = [("工作回顾", f["工作回顾"]), ("主要成绩", f["主要成绩"]),
                ("存在问题", f["存在问题"]),
                ("下一步工作打算", f["下一步工作打算"])]
    visible = [s for s in sections if s[1].present or s[1].required]
    for i, (label, field) in enumerate(visible):
        header = f"{CN_NUM[i]}、{label}"
        d.heading(header)
        if field.present:
            d.item_list(field.raw)
        else:
            d.para(PLACEHOLDER)

    d.sign_off(f["总结主体"].value, f["成文日期"].value)

    where = {"总结主体": "公文标题及文末右对齐落款署名",
             "总结时段": "公文标题",
             "工作回顾": "正文小标题下逐条编号列示",
             "主要成绩": "正文小标题下逐条编号列示",
             "存在问题": "正文小标题下逐条编号列示",
             "下一步工作打算": "正文小标题下逐条编号列示",
             "成文日期": "文末右对齐落款成文日期"}
    return title, [record(f[n], where[n]) for n, _ in SCHEMA["工作总结"]]


BUILDERS = {"周报": build_zhoubao, "请示函": build_qingshi,
            "会议通知": build_tongzhi, "工作总结": build_zongjie}


def main():
    ap = argparse.ArgumentParser(
        description="GB/T 9704 风格公文生成：--template/--data/--outdir，"
                    "在 outdir 下产出 文书.docx 与 fields.json")
    ap.add_argument("--template", required=True, choices=sorted(BUILDERS))
    ap.add_argument("--data", required=True)
    ap.add_argument("--outdir", required=True)
    args = ap.parse_args()

    with open(args.data, encoding="utf-8") as fh:
        data = json.load(fh)
    if not isinstance(data, dict):
        raise SystemExit(f"数据 JSON 必须是对象：{args.data}")

    os.makedirs(args.outdir, exist_ok=True)
    d = GbDoc()
    title, fields = BUILDERS[args.template](data, d)

    docx_path = os.path.join(args.outdir, "文书.docx")
    d.save(docx_path)

    summary = {
        "filled": sum(1 for x in fields if x["状态"] == "filled"),
        "missing": sum(1 for x in fields if x["状态"] == "missing"),
        "total": len(fields),
    }
    ledger = {"title": title, "template": args.template,
              "fields": fields, "summary": summary}
    with open(os.path.join(args.outdir, "fields.json"), "w",
              encoding="utf-8") as fh:
        json.dump(ledger, fh, ensure_ascii=False, indent=2)
        fh.write("\n")

    print(f"[gen] {args.template} -> {args.outdir}"
          f"（文书.docx + fields.json；filled={summary['filled']} "
          f"missing={summary['missing']} total={summary['total']}）")


if __name__ == "__main__":
    main()
