#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""GB/T 9704 风格公文生成器（office 重测臂：ab1-standard-four-templates）。

命令行契约：
    python3 gen_doc.py --template <周报|请示函|会议通知|工作总结> \
                       --data <数据json路径> --outdir <输出目录>

每次调用在 outdir 下生成两个产物：
    文书.docx    GB/T 9704 公文风格（A4 版面：上37/下35/左28/右26 mm；
                 大标题2号宋体居中，正文3号仿宋，一级小标题3号黑体，
                 固定行距28pt，首行缩进2字，落款右空4字，页脚页码）
    fields.json  字段记账：title / fields（每字段：名称、是否必填、状态、值、渲染方式）
                 / summary（filled/missing/total）

缺数据不编造：必填缺失在正文中以 ____ 占位；选填缺失留空或省略；均在 fields.json 如实记账。
数据 JSON 中的取值一律原样入文（列表逐条编号），不增删改写；
「妥否，请批示。」「现将……有关事项通知如下：」等为公文格式套语，属版式而非数据。
"""

import argparse
import json
import os
import sys

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt

PLACEHOLDER = "____"

FANGSONG = "仿宋"            # 正文：3号仿宋
HEITI = "黑体"                # 一级小标题：3号黑体
SONGTI = "宋体"               # 大标题（代小标宋）与页码
ASCII_FONT = "Times New Roman"

BODY_SIZE = Pt(16)            # 3号
TITLE_SIZE = Pt(22)           # 2号
PAGE_NO_SIZE = Pt(14)         # 4号
LINE_PT = 28                  # 固定行距 28pt
INDENT2 = 32                  # 首行/左空 2 字 = 32pt（3号字）
RIGHT_SPACE4 = 64             # 落款右空 4 字 = 64pt（3号字）


# ---------- 底层排版辅助 ----------

def _set_run_font(r, east, size, bold=False):
    r.font.name = ASCII_FONT
    r.font.size = size
    r.font.bold = bold
    rPr = r._element.get_or_add_rPr()
    rPr.get_or_add_rFonts().set(qn("w:eastAsia"), east)


def para(doc, align=None, indent_chars=0.0, right_pt=0.0, left_pt=0.0,
         line_pt=LINE_PT, before=0, after=0):
    p = doc.add_paragraph()
    pf = p.paragraph_format
    if align is not None:
        pf.alignment = align
    pf.line_spacing_rule = WD_LINE_SPACING.EXACTLY
    pf.line_spacing = Pt(line_pt)
    if indent_chars:
        pf.first_line_indent = Pt(indent_chars * 16)
    if right_pt:
        pf.right_indent = Pt(right_pt)
    if left_pt:
        pf.left_indent = Pt(left_pt)
    pf.space_before = Pt(before)
    pf.space_after = Pt(after)
    return p


def seg(p, text, east=FANGSONG, size=BODY_SIZE, bold=False):
    r = p.add_run(text)
    _set_run_font(r, east, size, bold)
    return r


def add_title(doc, text):
    p = para(doc, align=WD_ALIGN_PARAGRAPH.CENTER, line_pt=34, after=8)
    seg(p, text, east=SONGTI, size=TITLE_SIZE)


def add_subtitle(doc, text):
    p = para(doc, align=WD_ALIGN_PARAGRAPH.CENTER, after=8)
    seg(p, text)


def add_recipient(doc, text):
    p = para(doc)
    seg(p, text)


def add_body(doc, text, **kw):
    p = para(doc, indent_chars=2, **kw)
    seg(p, text)
    return p


def add_section_heading(doc, text):
    p = para(doc, indent_chars=2)
    seg(p, text, east=HEITI)


def add_item(doc, idx, text):
    p = para(doc, indent_chars=2)
    seg(p, "%d.%s" % (idx, text))


def add_signature(doc, lines):
    for i, t in enumerate(lines):
        p = para(doc, align=WD_ALIGN_PARAGRAPH.RIGHT, right_pt=RIGHT_SPACE4,
                 before=(28 if i == 0 else 0))
        seg(p, t)


def add_note(doc, text):
    p = para(doc, left_pt=INDENT2, before=28)
    seg(p, text)


def add_footer_page_number(section):
    p = section.footer.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    seg(p, "— ", east=SONGTI, size=PAGE_NO_SIZE)
    r = p.add_run()
    _set_run_font(r, SONGTI, PAGE_NO_SIZE)
    b = OxmlElement("w:fldChar")
    b.set(qn("w:fldCharType"), "begin")
    t = OxmlElement("w:instrText")
    t.set(qn("xml:space"), "preserve")
    t.text = " PAGE "
    e = OxmlElement("w:fldChar")
    e.set(qn("w:fldCharType"), "end")
    r._element.append(b)
    r._element.append(t)
    r._element.append(e)
    seg(p, " —", east=SONGTI, size=PAGE_NO_SIZE)


def new_document():
    doc = Document()
    sec = doc.sections[0]
    sec.page_width = Cm(21.0)
    sec.page_height = Cm(29.7)
    sec.top_margin = Cm(3.7)
    sec.bottom_margin = Cm(3.5)
    sec.left_margin = Cm(2.8)
    sec.right_margin = Cm(2.6)
    normal = doc.styles["Normal"]
    normal.font.name = ASCII_FONT
    normal.font.size = BODY_SIZE
    normal.element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:eastAsia"), FANGSONG)
    add_footer_page_number(sec)
    return doc


# ---------- 取值与记账 ----------

def val(data, name):
    """取字段值；空串/空列表/缺失/None 一律视为 missing（返回 None），不编造。"""
    v = data.get(name)
    if v is None:
        return None
    if isinstance(v, str):
        return v if v.strip() else None
    if isinstance(v, list):
        items = [str(x).strip() for x in v if str(x).strip()]
        return items or None
    return v


class Ledger(object):
    def __init__(self):
        self.fields = []

    def add(self, name, required, value, how):
        present = value is not None
        self.fields.append({
            "名称": name,
            "必填": bool(required),
            "状态": "filled" if present else "missing",
            "值": value if present else None,
            "渲染方式": how,
        })

    def summary(self):
        filled = sum(1 for r in self.fields if r["状态"] == "filled")
        return {"filled": filled, "missing": len(self.fields) - filled,
                "total": len(self.fields)}


def render_list_section(doc, ledger, heading, name, required, items, how):
    """编号列表节：有值逐条编号；必填缺失以 ____ 占位；选填缺失整节省略。"""
    if items is not None:
        add_section_heading(doc, heading)
        for i, item in enumerate(items, 1):
            add_item(doc, i, item)
        ledger.add(name, required, items, "%s（%d 条）" % (how, len(items)))
    elif required:
        add_section_heading(doc, heading)
        add_body(doc, PLACEHOLDER)
        ledger.add(name, True, None, "%s；必填缺失，以 %s 占位" % (how, PLACEHOLDER))
    else:
        ledger.add(name, False, None, "选填缺失，整节省略（留空）")


def render_contact_note(doc, ledger, contact, phone):
    """附注行：成文日期下一行左空2字加圆括号；选填缺失的部分省略不写。"""
    parts = []
    if contact is not None:
        parts.append("联系人：%s" % contact)
    if phone is not None:
        parts.append("联系电话：%s" % phone)
    if parts:
        add_note(doc, "（%s）" % "，".join(parts))
    ledger.add("联系人", False, contact,
               "附注行（成文日期下一行，左空2字加圆括号）" if contact is not None
               else "选填缺失，附注中省略")
    ledger.add("联系电话", False, phone,
               "附注行（成文日期下一行，左空2字加圆括号）" if phone is not None
               else "选填缺失，附注中省略")


# ---------- 四类模板 ----------

def build_zhoubao(doc, data, ledger):
    dept = val(data, "部门")
    period = val(data, "周期")
    person = val(data, "填报人")
    done = val(data, "本周工作内容")
    plan = val(data, "下周工作计划")
    issues = val(data, "问题与需协调事项")
    date = val(data, "报送日期")

    title = "%s工作周报" % (dept if dept is not None else PLACEHOLDER)
    add_title(doc, title)
    ledger.add("部门", True, dept, "大标题「%s」" % title)

    add_subtitle(doc, "（%s）" % (period if period is not None else PLACEHOLDER))
    ledger.add("周期", True, period, "标题下副题行（居中，3号仿宋）")

    render_list_section(doc, ledger, "一、本周工作内容", "本周工作内容", True,
                        done, "正文「一、本周工作内容」节，编号条目")
    render_list_section(doc, ledger, "二、下周工作计划", "下周工作计划", True,
                        plan, "正文「二、下周工作计划」节，编号条目")
    render_list_section(doc, ledger, "三、问题与需协调事项", "问题与需协调事项", False,
                        issues, "正文「三、问题与需协调事项」节，编号条目")

    sig = ["填报人：%s" % (person if person is not None else PLACEHOLDER),
           "报送日期：%s" % (date if date is not None else PLACEHOLDER)]
    add_signature(doc, sig)
    ledger.add("填报人", True, person,
               "文末落款（右对齐，右空4字）" if person is not None
               else "文末落款（右对齐）；必填缺失，以 %s 占位" % PLACEHOLDER)
    ledger.add("报送日期", True, date,
               "文末落款（右对齐，右空4字）" if date is not None
               else "文末落款（右对齐）；必填缺失，以 %s 占位" % PLACEHOLDER)
    return title


def build_qingshihan(doc, data, ledger):
    subject = val(data, "请示事由")
    recipient = val(data, "主送机关")
    reason = val(data, "请示缘由")
    matter = val(data, "请示事项")
    unit = val(data, "请示单位")
    contact = val(data, "联系人")
    phone = val(data, "联系电话")
    date = val(data, "成文日期")

    title = "关于%s的请示" % (subject if subject is not None else PLACEHOLDER)
    add_title(doc, title)
    ledger.add("请示事由", True, subject, "大标题「%s」" % title)

    add_recipient(doc, "%s：" % (recipient if recipient is not None else PLACEHOLDER))
    ledger.add("主送机关", True, recipient,
               "正文首行顶格「×××：」" if recipient is not None
               else "正文首行顶格；必填缺失，以 %s 占位" % PLACEHOLDER)

    for name, v, how in (("请示缘由", reason, "正文段落（首行缩进2字）"),
                         ("请示事项", matter, "正文段落（首行缩进2字）")):
        if v is not None:
            add_body(doc, v)
            ledger.add(name, True, v, how)
        else:
            add_body(doc, PLACEHOLDER)
            ledger.add(name, True, None, "%s；必填缺失，以 %s 占位" % (how, PLACEHOLDER))

    add_body(doc, "妥否，请批示。")  # 请示类公文格式套语

    sig = [unit if unit is not None else PLACEHOLDER,
           date if date is not None else PLACEHOLDER]
    add_signature(doc, sig)
    ledger.add("请示单位", True, unit,
               "文末落款（右对齐，右空4字）" if unit is not None
               else "文末落款（右对齐）；必填缺失，以 %s 占位" % PLACEHOLDER)
    ledger.add("成文日期", True, date,
               "文末落款日期行（右空4字）" if date is not None
               else "文末落款日期行；必填缺失，以 %s 占位" % PLACEHOLDER)

    render_contact_note(doc, ledger, contact, phone)
    return title


def build_huiyitongzhi(doc, data, ledger):
    name = val(data, "会议名称")
    unit = val(data, "召开单位")
    to = val(data, "主送对象")
    mtime = val(data, "会议时间")
    place = val(data, "会议地点")
    attendees = val(data, "参会人员")
    topic = val(data, "会议议题")
    demands = val(data, "会议要求")
    contact = val(data, "联系人")
    phone = val(data, "联系电话")
    date = val(data, "发文日期")

    n_disp = name if name is not None else PLACEHOLDER
    title = "关于召开%s的通知" % n_disp
    add_title(doc, title)
    ledger.add("会议名称", True, name, "大标题「%s」" % title)

    add_recipient(doc, "%s：" % (to if to is not None else PLACEHOLDER))
    ledger.add("主送对象", True, to,
               "正文首行顶格「×××：」" if to is not None
               else "正文首行顶格；必填缺失，以 %s 占位" % PLACEHOLDER)

    add_body(doc, "现将%s有关事项通知如下：" % n_disp)

    def labeled(no, label, v, fname, how):
        p = para(doc, indent_chars=2)
        seg(p, "%s、%s" % (no, label), east=HEITI)
        if v is not None:
            seg(p, "：%s" % v)
            ledger.add(fname, True, v, how)
        else:
            seg(p, "：%s" % PLACEHOLDER)
            ledger.add(fname, True, None, "%s；必填缺失，以 %s 占位" % (how, PLACEHOLDER))

    labeled("一", "会议时间", mtime, "会议时间", "正文「一、会议时间」条（黑体标签+仿宋取值）")
    labeled("二", "会议地点", place, "会议地点", "正文「二、会议地点」条（黑体标签+仿宋取值）")
    labeled("三", "参会人员", attendees, "参会人员", "正文「三、参会人员」条（黑体标签+仿宋取值）")
    labeled("四", "会议议题", topic, "会议议题", "正文「四、会议议题」条（黑体标签+仿宋取值）")
    render_list_section(doc, ledger, "五、会议要求", "会议要求", False,
                        demands, "正文「五、会议要求」节，编号条目")

    sig = [unit if unit is not None else PLACEHOLDER,
           date if date is not None else PLACEHOLDER]
    add_signature(doc, sig)
    ledger.add("召开单位", True, unit,
               "文末落款（右对齐，右空4字）" if unit is not None
               else "文末落款（右对齐）；必填缺失，以 %s 占位" % PLACEHOLDER)
    ledger.add("发文日期", True, date,
               "文末落款日期行（右空4字）" if date is not None
               else "文末落款日期行；必填缺失，以 %s 占位" % PLACEHOLDER)

    render_contact_note(doc, ledger, contact, phone)
    return title


def build_gongzuozongjie(doc, data, ledger):
    unit = val(data, "总结主体")
    period = val(data, "总结时段")
    review = val(data, "工作回顾")
    achievements = val(data, "主要成绩")
    problems = val(data, "存在问题")
    plans = val(data, "下一步工作打算")
    date = val(data, "成文日期")

    title = "%s%s工作总结" % (unit if unit is not None else PLACEHOLDER,
                              period if period is not None else PLACEHOLDER)
    add_title(doc, title)
    ledger.add("总结主体", True, unit, "大标题「%s」" % title)
    ledger.add("总结时段", True, period, "大标题「%s」" % title)

    render_list_section(doc, ledger, "一、工作回顾", "工作回顾", True,
                        review, "正文「一、工作回顾」节，编号条目")
    render_list_section(doc, ledger, "二、主要成绩", "主要成绩", True,
                        achievements, "正文「二、主要成绩」节，编号条目")
    render_list_section(doc, ledger, "三、存在问题", "存在问题", False,
                        problems, "正文「三、存在问题」节，编号条目")
    render_list_section(doc, ledger, "四、下一步工作打算", "下一步工作打算", True,
                        plans, "正文「四、下一步工作打算」节，编号条目")

    sig = [unit if unit is not None else PLACEHOLDER,
           date if date is not None else PLACEHOLDER]
    add_signature(doc, sig)
    ledger.add("成文日期", True, date,
               "文末落款日期行（右空4字）" if date is not None
               else "文末落款日期行；必填缺失，以 %s 占位" % PLACEHOLDER)
    return title


BUILDERS = {
    "周报": build_zhoubao,
    "请示函": build_qingshihan,
    "会议通知": build_huiyitongzhi,
    "工作总结": build_gongzuozongjie,
}


def main():
    ap = argparse.ArgumentParser(
        description="GB/T 9704 风格公文生成器（周报/请示函/会议通知/工作总结）")
    ap.add_argument("--template", required=True,
                    choices=["周报", "请示函", "会议通知", "工作总结"],
                    help="模板类型")
    ap.add_argument("--data", required=True, help="数据 JSON 路径")
    ap.add_argument("--outdir", required=True, help="输出目录（生成 文书.docx 与 fields.json）")
    args = ap.parse_args()

    with open(args.data, "r", encoding="utf-8") as f:
        data = json.load(f)
    if not isinstance(data, dict):
        print("[gen_doc] 错误：数据 JSON 顶层必须是对象，实际为 %s" % type(data).__name__,
              file=sys.stderr)
        sys.exit(2)

    os.makedirs(args.outdir, exist_ok=True)

    doc = new_document()
    ledger = Ledger()
    title = BUILDERS[args.template](doc, data, ledger)
    summary = ledger.summary()

    fields_doc = {
        "template": args.template,
        "title": title,
        "fields": ledger.fields,
        "summary": summary,
    }
    docx_path = os.path.join(args.outdir, "文书.docx")
    json_path = os.path.join(args.outdir, "fields.json")
    doc.save(docx_path)
    with open(json_path, "w", encoding="utf-8") as f:
        json.dump(fields_doc, f, ensure_ascii=False, indent=2)
        f.write("\n")
    print("[gen_doc] 模板=%s → %s + %s（filled=%d missing=%d total=%d）"
          % (args.template, docx_path, json_path,
             summary["filled"], summary["missing"], summary["total"]))


if __name__ == "__main__":
    main()
