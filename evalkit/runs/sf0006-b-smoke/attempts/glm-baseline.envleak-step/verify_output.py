#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""自检：核对 out/纪要.docx 与 out/summary.json 是否满足 brief.md 全部硬性要求。"""
import json
import os
import re

from docx import Document

BASE = os.path.dirname(os.path.abspath(__file__))
OUT = os.path.join(BASE, "out")

with open(os.path.join(BASE, "fixtures", "transcript.txt"), encoding="utf-8") as f:
    transcript_raw = f.read()

# 说话人 -> 其说过的句子集合（用于核对条目确为该说话人所述）
said_by = {}
for line in transcript_raw.splitlines():
    m = re.match(r"^\[\d{2}:\d{2}\]\s*(\S+):\s*(.+)$", line.strip())
    if m:
        said_by.setdefault(m.group(1), []).append(m.group(2))

speakers = set(said_by)
transcript_text = "\n".join(s for ss in said_by.values() for s in ss)

doc = Document(os.path.join(OUT, "纪要.docx"))
paras = [p.text for p in doc.paragraphs]
print("=== docx 段落 ===")
for i, t in enumerate(paras):
    print(f"[{i}] {t!r}")

print("\n=== docx 表格 ===")
for tbl in doc.tables:
    for row in tbl.rows:
        print([c.text for c in row.cells])

ok = True
def check(name, cond, detail=""):
    global ok
    mark = "PASS" if cond else "FAIL"
    if not cond:
        ok = False
    print(f"[{mark}] {name}" + (f" — {detail}" if detail else ""))

# 1. 第一段为标题，且包含会议主题「星帆订单系统」
check("第一段为纪要标题且含会议主题", paras[0] == "星帆订单系统项目例会会议纪要", paras[0])

# 2. 三个节标题逐字存在且顺序正确
h_idx = [i for i, t in enumerate(paras) if t in ("一、决议事项", "二、待办事项", "三、风险与关注")]
check("三个节标题逐字且按序", [paras[i] for i in h_idx] == ["一、决议事项", "二、待办事项", "三、风险与关注"],
      str([paras[i] for i in h_idx]))
check("节标题出现且互不重复", len(h_idx) == 3)

# 3. 决议事项/风险与关注条目格式与溯源
pat = re.compile(r"^(\d+)\. 【(决议事项|风险与关注) · (\S+)】(.+)$")
items = {"决议事项": [], "风险与关注": []}
for t in paras:
    m = pat.match(t)
    if m:
        items[m.group(2)].append((int(m.group(1)), m.group(3), m.group(4)))

for sec in ("决议事项", "风险与关注"):
    lst = items[sec]
    check(f"{sec} 条目存在", len(lst) > 0)
    check(f"{sec} 编号从 1 连续", [n for n, _, _ in lst] == list(range(1, len(lst) + 1)))
    for n, person, content in lst:
        check(f"{sec} 第{n}条 相关人「{person}」为转写稿说话人", person in speakers)
        check(f"{sec} 第{n}条 内容逐字可溯源", content in transcript_text, content)
        check(f"{sec} 第{n}条 内容确为该说话人「{person}」所述",
              any(s.rstrip("。") == content for s in said_by.get(person, [])))

# 4. 待办事项表格
check("恰好一个表格", len(doc.tables) == 1)
tbl = doc.tables[0]
header = [c.text for c in tbl.rows[0].cells]
check("表头逐字为 事项/负责人/期限", header == ["事项", "负责人", "期限"], str(header))
todo_rows = [[c.text for c in r.cells] for r in tbl.rows[1:]]
check("待办表格有数据行", len(todo_rows) > 0)
for item, owner, deadline in todo_rows:
    check(f"待办「{item[:12]}…」事项逐字可溯源", item in transcript_text, item)
    check(f"待办「{item[:12]}…」事项确为负责人「{owner}」所述",
          any(s.rstrip("。") == item for s in said_by.get(owner, [])))
    check(f"待办「{item[:12]}…」负责人为说话人", owner in speakers, owner)
    check(f"待办「{item[:12]}…」期限非空", bool(deadline), deadline)

# 5. summary.json
with open(os.path.join(OUT, "summary.json"), encoding="utf-8") as f:
    summary = json.load(f)
print("\nsummary.json =", summary)
check("summary 四个键名逐字", set(summary.keys()) == {"决议事项", "待办事项", "风险与关注", "合计"})
check("决议事项计数一致", summary["决议事项"] == len(items["决议事项"]))
check("风险与关注计数一致", summary["风险与关注"] == len(items["风险与关注"]))
check("待办事项计数一致", summary["待办事项"] == len(todo_rows))
check("合计为三者之和", summary["合计"] == summary["决议事项"] + summary["待办事项"] + summary["风险与关注"])

print("\n" + ("ALL CHECKS PASSED" if ok else "SOME CHECKS FAILED"))
raise SystemExit(0 if ok else 1)
