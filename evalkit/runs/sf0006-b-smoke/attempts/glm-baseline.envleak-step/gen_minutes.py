#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""依据 fixtures/transcript.txt 生成三节式会议纪要：out/纪要.docx 与 out/summary.json。

所有条目内容逐字取自转写稿（去掉句末句号），人名均为转写稿中出现过的说话人。
"""
import json
import os

from docx import Document
from docx.shared import Pt
from docx.oxml.ns import qn

BASE = os.path.dirname(os.path.abspath(__file__))
OUT = os.path.join(BASE, "out")
os.makedirs(OUT, exist_ok=True)

TITLE = "星帆订单系统项目例会会议纪要"

# (相关人, 内容逐字取自转写稿对应句子，去句末句号)
DECISIONS = [
    ("赵静怡", "星帆订单系统整体进度正常，会议同意按原计划推进，里程碑保持不变"),
    ("黄雨桐", "会议同意星帆订单系统项目预算300万元，走专项列支，由市场部跟进执行"),
]

# (事项逐字取自转写稿, 负责人, 期限；转写稿未提及时限填「未明确」)
TODOS = [
    ("黄雨桐在本周五前完成复盘纪要并同步给财务部", "黄雨桐", "本周五前"),
    ("徐静怡负责把数据迁移脚本整理成7月4日归档", "徐静怡", "7月4日"),
]

RISKS = [
    ("吴志远", "财务部侧存在履约风险，去年高峰期出现过积压，今年需要提前准备预案"),
    ("赵静怡", "灰度放量方案的数据口径尚未对齐，复盘前需要补齐基线数据"),
]


def build_docx():
    doc = Document()

    normal = doc.styles["Normal"]
    normal.font.name = "SimSun"
    normal.font.size = Pt(12)
    normal.element.rPr.rFonts.set(qn("w:eastAsia"), "SimSun")

    # 第一段：纪要标题
    doc.add_paragraph(TITLE)

    # 一、决议事项
    doc.add_heading("一、决议事项", level=1)
    for i, (person, content) in enumerate(DECISIONS, 1):
        doc.add_paragraph(f"{i}. 【决议事项 · {person}】{content}")

    # 二、待办事项（表格）
    doc.add_heading("二、待办事项", level=1)
    table = doc.add_table(rows=1, cols=3)
    table.style = "Table Grid"
    for cell, header in zip(table.rows[0].cells, ["事项", "负责人", "期限"]):
        cell.text = header
    for item, owner, deadline in TODOS:
        row = table.add_row().cells
        row[0].text = item
        row[1].text = owner
        row[2].text = deadline

    # 三、风险与关注
    doc.add_heading("三、风险与关注", level=1)
    for i, (person, content) in enumerate(RISKS, 1):
        doc.add_paragraph(f"{i}. 【风险与关注 · {person}】{content}")

    path = os.path.join(OUT, "纪要.docx")
    doc.save(path)
    return path


def build_summary():
    summary = {
        "决议事项": len(DECISIONS),
        "待办事项": len(TODOS),
        "风险与关注": len(RISKS),
        "合计": len(DECISIONS) + len(TODOS) + len(RISKS),
    }
    path = os.path.join(OUT, "summary.json")
    with open(path, "w", encoding="utf-8") as f:
        json.dump(summary, f, ensure_ascii=False, indent=2)
        f.write("\n")
    return path, summary


if __name__ == "__main__":
    docx_path = build_docx()
    json_path, summary = build_summary()
    print("docx :", docx_path)
    print("json :", json_path)
    print("summary:", json.dumps(summary, ensure_ascii=False))
