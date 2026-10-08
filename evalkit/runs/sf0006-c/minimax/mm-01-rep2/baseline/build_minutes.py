# -*- coding: utf-8 -*-
"""由 fixtures/transcript.txt 生成 out/纪要.docx 与 out/summary.json。
所有条目内容均逐字取自转写稿，不做归纳改写。"""
import json
import os

from docx import Document

BASE = os.path.dirname(os.path.abspath(__file__))
TRANSCRIPT = os.path.join(BASE, "fixtures", "transcript.txt")
OUT_DIR = os.path.join(BASE, "out")

TITLE = "云岭数据看板项目例会会议纪要"

DECISIONS = [
    ("陈静怡", "会议同意按原计划推进，里程碑保持不变"),
    ("徐丽娜", "会议同意云岭数据看板项目预算16万元，走专项列支，由客户成功部跟进执行"),
]

TODOS = [
    ("完成压测报告并同步给综合办", "徐丽娜", "7月4日前"),
    ("把验收清单整理成下周三归档", "张文博", "下周三"),
    ("复盘前需要补齐基线数据", "陈静怡", "复盘前"),
]

RISKS = [
    ("黄丽娜", "综合办侧存在履约风险，去年高峰期出现过积压，今年需要提前准备预案"),
    ("陈静怡", "灰度放量方案的数据口径尚未对齐"),
]


def check_traceable():
    """自证：每条内容、负责人、事项都必须在转写稿中逐字出现。"""
    raw = open(TRANSCRIPT, encoding="utf-8").read()
    missing = []
    for person, text in DECISIONS + RISKS:
        if text not in raw:
            missing.append(text)
        if person not in raw:
            missing.append("人名:" + person)
    for item, owner, _ in TODOS:
        if item not in raw:
            missing.append(item)
        if owner not in raw:
            missing.append("人名:" + owner)
    if missing:
        raise SystemExit("以下条目无法在转写稿中逐字溯源: %s" % missing)


def entry(idx, category, who, text):
    return "%d. 【%s · %s】%s" % (idx, category, who, text)


def main():
    check_traceable()
    os.makedirs(OUT_DIR, exist_ok=True)

    doc = Document()

    p = doc.add_paragraph()
    run = p.add_run(TITLE)
    run.bold = True
    run.font.size = __import__("docx").shared.Pt(16)

    doc.add_paragraph("一、决议事项")
    for i, (who, text) in enumerate(DECISIONS, 1):
        doc.add_paragraph(entry(i, "决议事项", who, text))

    doc.add_paragraph("二、待办事项")
    table = doc.add_table(rows=1, cols=3)
    table.style = "Table Grid"
    for cell, head in zip(table.rows[0].cells, ["事项", "负责人", "期限"]):
        cell.text = ""
        r = cell.paragraphs[0].add_run(head)
        r.bold = True
    for item, owner, deadline in TODOS:
        cells = table.add_row().cells
        cells[0].text = item
        cells[1].text = owner
        cells[2].text = deadline

    doc.add_paragraph("三、风险与关注")
    for i, (who, text) in enumerate(RISKS, 1):
        doc.add_paragraph(entry(i, "风险与关注", who, text))

    doc.save(os.path.join(OUT_DIR, "纪要.docx"))

    summary = {
        "决议事项": len(DECISIONS),
        "待办事项": len(TODOS),
        "风险与关注": len(RISKS),
        "合计": len(DECISIONS) + len(TODOS) + len(RISKS),
    }
    with open(os.path.join(OUT_DIR, "summary.json"), "w", encoding="utf-8") as f:
        json.dump(summary, f, ensure_ascii=False)

    print("done", summary)


if __name__ == "__main__":
    main()