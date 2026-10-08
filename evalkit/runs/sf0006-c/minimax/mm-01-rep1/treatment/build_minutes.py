#!/usr/bin/env python3
"""按 brief.md 的格式契约，从转写稿生成 out/纪要.docx 与 out/summary.json。

所有条目内容均逐字取自转写稿，不改写、不概括。
"""
import json
import os
import re
import sys

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Pt

TRANSCRIPT = "fixtures/transcript.txt"
OUTDIR = "out"
LINE_RE = re.compile(r"^\s*\[(\d{1,2}:\d{2})\]\s*([^：:]+?)\s*[：:]\s*(.+?)\s*$")

# (时间, 说话人, 归类, 待办负责人, 待办期限)
# 归类依据：句中出现「同意/议定」→ 决议；明确的行动项 → 待办；出现「风险」→ 风险。
ENTRIES = [
    ("09:31", "赵静怡", "决议事项", None, None),
    ("09:36", "黄雨桐", "决议事项", None, None),
    ("09:37", "主持人", "决议事项", None, None),
    ("09:39", "黄雨桐", "待办事项", "黄雨桐", "本周五前"),
    ("09:44", "徐静怡", "待办事项", "徐静怡", "7月4日"),
    ("09:47", "赵静怡", "待办事项", "赵静怡", "复盘前"),
    ("09:46", "吴志远", "风险与关注", None, None),
]

NO_DEADLINE = "未明确"


def load_lines(path):
    with open(path, encoding="utf-8-sig") as fh:
        out = []
        for raw in fh:
            m = LINE_RE.match(raw)
            if m:
                out.append((m.group(1), m.group(2).strip(), m.group(3)))
    return out


def set_font(run, name="宋体"):
    run.font.name = name
    run._element.rPr.rFonts.set(qn("w:eastAsia"), name)


def add_para(doc, text, size=None, bold=False, align=None):
    p = doc.add_paragraph()
    if align is not None:
        p.alignment = align
    run = p.add_run(text)
    run.bold = bold
    if size:
        run.font.size = Pt(size)
    set_font(run)
    return p


def main():
    lines = load_lines(TRANSCRIPT)
    index = {(t, s): c for t, s, c in lines}

    # 逐字溯源自检：每条内容必须能在其 (时间, 说话人) 出处行中逐字找到
    for t, sp, cat, owner, due in ENTRIES:
        if (t, sp) not in index:
            sys.exit(f"转写稿中找不到出处行：[{t}] {sp}")
        content = index[(t, sp)]
        # 期限必须是原句中的逐字时间语
        if due and due != NO_DEADLINE and due not in content:
            sys.exit(f"[{t}] {sp} 的期限「{due}」不在原句中：{content}")
        # 负责人：原句中逐字出现，或由该行说话人第一人称承担
        if owner and owner not in content and owner != sp:
            sys.exit(f"[{t}] {sp} 的负责人「{owner}」既不在原句中也不是说话人")
    speakers = {s for _, s, _ in lines}
    for _, _, _, owner, _ in ENTRIES:
        if owner and owner not in speakers:
            sys.exit(f"负责人「{owner}」不是转写稿中的说话人")

    decisions = [e for e in ENTRIES if e[2] == "决议事项"]
    todos = [e for e in ENTRIES if e[2] == "待办事项"]
    risks = [e for e in ENTRIES if e[2] == "风险与关注"]
    total = len(decisions) + len(todos) + len(risks)

    os.makedirs(OUTDIR, exist_ok=True)
    doc = Document()
    add_para(doc, "星帆订单系统项目例会纪要", size=18, bold=True,
             align=WD_ALIGN_PARAGRAPH.CENTER)
    add_para(
        doc,
        "来源：{} ｜ 收录条目 {} 条（决议 {} · 待办 {} · 风险 {}）".format(
            os.path.basename(TRANSCRIPT), total,
            len(decisions), len(todos), len(risks)),
    )

    for title, items in [("一、决议事项", decisions)]:
        add_para(doc, title, size=14, bold=True)
        for i, (t, sp, cat, _, _) in enumerate(items, start=1):
            add_para(doc, "{}. 【{} · {}】{}".format(i, cat, sp, index[(t, sp)]))

    add_para(doc, "二、待办事项", size=14, bold=True)
    table = doc.add_table(rows=1, cols=3)
    table.style = "Table Grid"
    for cell, head in zip(table.rows[0].cells, ("事项", "负责人", "期限")):
        cell.text = ""
        run = cell.paragraphs[0].add_run(head)
        run.bold = True
        set_font(run)
    for t, sp, _, owner, due in todos:
        row = table.add_row().cells
        for cell, val in zip(row, (index[(t, sp)], owner, due or NO_DEADLINE)):
            cell.text = ""
            run = cell.paragraphs[0].add_run(val)
            set_font(run)

    add_para(doc, "三、风险与关注", size=14, bold=True)
    for i, (t, sp, cat, _, _) in enumerate(risks, start=1):
        add_para(doc, "{}. 【{} · {}】{}".format(i, cat, sp, index[(t, sp)]))

    docx_path = os.path.join(OUTDIR, "纪要.docx")
    doc.save(docx_path)

    summary = {
        "决议事项": len(decisions),
        "待办事项": len(todos),
        "风险与关注": len(risks),
        "合计": total,
    }
    with open(os.path.join(OUTDIR, "summary.json"), "w", encoding="utf-8") as fh:
        json.dump(summary, fh, ensure_ascii=False, indent=2)
        fh.write("\n")
    print("OK", docx_path, "| 决议", len(decisions), "· 待办", len(todos),
          "· 风险", len(risks), "· 合计", total)


if __name__ == "__main__":
    main()
