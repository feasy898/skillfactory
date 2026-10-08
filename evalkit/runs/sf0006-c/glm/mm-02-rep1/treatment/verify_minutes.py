#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""按 brief.md 逐条机械验收 out/纪要.docx 与 out/summary.json。"""
import json
import os
import re
import sys

from docx import Document

BASE = os.path.dirname(os.path.abspath(__file__))
TRANSCRIPT = os.path.join(BASE, "fixtures", "transcript.txt")
OUTDIR = os.path.join(BASE, "out")

failures = []


def check(name, ok, detail=""):
    print("%s  %s%s" % ("PASS" if ok else "FAIL", name, ("  | " + detail) if detail and not ok else ""))
    if not ok:
        failures.append(name)


transcript_text = open(TRANSCRIPT, encoding="utf-8-sig").read()
turn_count = {}
for raw in transcript_text.splitlines():
    m = re.match(r"^([^：:]+?)\s*[：:]\s*(.+?)\s*$", raw.strip())
    if m:
        turn_count[m.group(1)] = turn_count.get(m.group(1), 0) + 1
speaker_names = set(turn_count)

docx_path = os.path.join(OUTDIR, "纪要.docx")
json_path = os.path.join(OUTDIR, "summary.json")
check("产物存在: 纪要.docx", os.path.isfile(docx_path))
check("产物存在: summary.json", os.path.isfile(json_path))

doc = Document(docx_path)
paras = [p.text for p in doc.paragraphs]
tables = doc.tables
all_text = "\n".join(paras) + "\n" + "\n".join(
    c.text for t in tables for r in t.rows for c in r.cells
)

# 1. 第一段为标题，含会议主题
check("第一段为标题且含主题「望舒客服工作台」", bool(paras) and "望舒客服工作台" in paras[0], repr(paras[0] if paras else ""))

# 2. 三节标题逐字齐全且有序
expected_headers = ["一、决议事项", "二、待办事项", "三、风险与关注"]
positions = []
for h in expected_headers:
    pos = [i for i, t in enumerate(paras) if t.strip() == h]
    check("节标题逐字存在: %s" % h, len(pos) == 1, "出现 %d 次" % len(pos))
    positions.append(pos[0] if pos else -1)
check("三节标题顺序正确", positions == sorted(positions) and all(p >= 0 for p in positions))

ENTRY_RE = re.compile(r"^(\d+)\. 【(决议事项|风险与关注) · (.+?)】(.+)$")


def parse_entries(start, end):
    out = []
    for t in paras[start + 1 : end]:
        if not t.strip():
            continue
        m = ENTRY_RE.match(t.strip())
        out.append((t.strip(), m) if m else (t.strip(), None))
    return out


dec_entries = parse_entries(positions[0], positions[1])
risk_entries = parse_entries(positions[2], len(paras))

# 3. 决议/风险条目格式与逐字溯源
def check_section(entries, category, section_name):
    for idx, (text, m) in enumerate(entries, 1):
        tag = "%s 第%d条" % (section_name, idx)
        if not m:
            check("%s 格式" % tag, False, repr(text))
            continue
        num, cat, person, content = m.groups()
        check("%s 编号连续" % tag, num == str(idx), "编号=%s" % num)
        check("%s 类别正确" % tag, cat == category, "类别=%s" % cat)
        check("%s 相关人是转写稿说话人" % tag, person in speaker_names, repr(person))
        body = content[: -len("（归属待核）")] if content.endswith("（归属待核）") else content
        check("%s 内容逐字溯源" % tag, body in transcript_text, repr(content))
        # 同名归属：多次发言的姓名（稿内无区分标记）无法唯一归属 → 必须带后缀；否则不得带
        ambiguous = turn_count.get(person, 0) > 1
        has_flag = content.endswith("（归属待核）")
        check(
            "%s 归属待核后缀正确（%s）" % (tag, "同名多轮→应有" if ambiguous else "唯一归属→不应有"),
            has_flag == ambiguous,
            "has_flag=%s ambiguous=%s" % (has_flag, ambiguous),
        )


check_section(dec_entries, "决议事项", "决议事项")
check_section(risk_entries, "风险与关注", "风险与关注")

# 4. 待办表格
check("待办为表格", len(tables) == 1, "表格数=%d" % len(tables))
table = tables[0]
check("表格 3 列", len(table.columns) == 3)
header = [c.text.strip() for c in table.rows[0].cells]
check("表头逐字为 事项/负责人/期限", header == ["事项", "负责人", "期限"], repr(header))
todo_rows = []
for i, row in enumerate(table.rows[1:], 1):
    cells = [c.text.strip() for c in row.cells]
    item, owner, ddl = cells
    check("待办%d 事项逐字溯源" % i, item in transcript_text, repr(item))
    check("待办%d 负责人是说话人" % i, owner in speaker_names, repr(owner))
    if ddl != "未明确":
        check("待办%d 期限在转写稿中有出处" % i, ddl in transcript_text, repr(ddl))
    todo_rows.append(cells)
check("至少一条期限为「未明确」", any(r[2] == "未明确" for r in todo_rows))

# 5. summary.json
summary = json.load(open(json_path, encoding="utf-8"))
check("summary 四键逐字", list(summary.keys()) == ["决议事项", "待办事项", "风险与关注", "合计"], repr(list(summary.keys())))
check("summary 决议数=docx 实际", summary.get("决议事项") == len(dec_entries), "%s vs %d" % (summary.get("决议事项"), len(dec_entries)))
check("summary 待办数=docx 实际", summary.get("待办事项") == len(todo_rows), "%s vs %d" % (summary.get("待办事项"), len(todo_rows)))
check("summary 风险数=docx 实际", summary.get("风险与关注") == len(risk_entries), "%s vs %d" % (summary.get("风险与关注"), len(risk_entries)))
check("summary 合计=三者之和", summary.get("合计") == summary.get("决议事项") + summary.get("待办事项") + summary.get("风险与关注"))

# 6. 纪律：无 [HH:MM] 时间；无虚构区分标记
check("全文无 [HH:MM] 形式时间", re.search(r"\[\d{1,2}:\d{2}\]", all_text) is None)
invented = re.findall(r"(刘文博|黄静怡|黄文博|王晓雯|张晓雯)[（(（]", all_text)
check("无虚构区分标记（名字后带括号）", not invented, repr(invented))

print()
if failures:
    print("验收未通过，%d 项 FAIL: %s" % (len(failures), failures))
    sys.exit(1)
print("全部检查通过。")
