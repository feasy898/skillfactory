#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""gen_minutes.py — 从 fixtures/transcript.txt 生成 out/纪要.docx + out/summary.json。

纪律（brief.md）：
- 转写稿为唯一事实来源；条目内容逐字取自转写稿（不含句末句号）；
- 不补造 [HH:MM] 时间戳；不虚构同名区分标记；
- 同名说话人（刘文博，稿内无区分标记）无法唯一归属的条目，内容末尾追加（归属待核）；
- 无期限待办的期限填「未明确」。
"""

import io
import json
import os

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Pt

HERE = os.path.dirname(os.path.abspath(__file__))
TRANSCRIPT = os.path.join(HERE, "fixtures", "transcript.txt")
OUT_DIR = os.path.join(HERE, "out")

SAME_NAME = "刘文博"  # 同名说话人：主持/决议陈述/风险补充多角色共用此名，稿内无区分标记

DECISIONS = [
    ("刘文博", "望舒客服工作台整体进度正常，会议同意按原计划推进，里程碑保持不变"),
    ("黄静怡", "会议同意望舒客服工作台项目预算300万元，走专项列支，由综合办跟进执行"),
    ("黄文博", "望舒客服工作台的验收标准按新版执行，旧版同时作废"),
]

TODOS = [
    ("黄静怡在6月30日前完成复盘纪要并同步给市场部", "黄静怡", "6月30日前"),
    ("王晓雯负责把灰度放量方案整理成7月中旬归档", "王晓雯", "7月中旬"),
    ("张晓雯牵头梳理望舒客服工作台的责任分工", "张晓雯", "未明确"),
]

RISKS = [
    ("张晓雯", "市场部侧存在履约风险，去年高峰期出现过积压，今年需要提前准备预案"),
    ("刘文博", "验收清单的数据口径尚未对齐，复盘前需要补齐基线数据"),
]


def main():
    with io.open(TRANSCRIPT, encoding="utf-8") as f:
        transcript = f.read()

    # 自证：所有条目内容/事项/期限逐字存在于转写稿，人名均为说话人
    speakers = set()
    for line in transcript.splitlines():
        if ":" in line[:13]:
            speakers.add(line.split(":", 1)[0].strip())
    for owner, content in DECISIONS + RISKS:
        assert owner in speakers, owner
        assert content in transcript, content
    for item, owner, dl in TODOS:
        assert item in transcript, item
        assert owner in speakers, owner
        assert dl == "未明确" or (dl in item and dl in transcript), dl

    doc = Document()
    normal = doc.styles["Normal"]
    normal.font.name = "Times New Roman"
    normal.font.size = Pt(12)
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")

    title = doc.add_paragraph("望舒客服工作台项目例会会议纪要")
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = title.runs[0]
    run.bold = True
    run.font.size = Pt(16)

    def add_header(text):
        p = doc.add_paragraph(text)
        p.runs[0].bold = True
        p.runs[0].font.size = Pt(14)
        return p

    def add_item(idx, cat, owner, content):
        suffix = "（归属待核）" if owner == SAME_NAME else ""
        doc.add_paragraph("%d. 【%s · %s】%s%s" % (idx, cat, owner, content, suffix))

    add_header("一、决议事项")
    for i, (owner, content) in enumerate(DECISIONS, 1):
        add_item(i, "决议事项", owner, content)

    add_header("二、待办事项")
    table = doc.add_table(rows=1 + len(TODOS), cols=3)
    table.style = "Table Grid"
    for j, h in enumerate(("事项", "负责人", "期限")):
        cell = table.rows[0].cells[j]
        cell.text = h
        cell.paragraphs[0].runs[0].bold = True
    for i, (item, owner, dl) in enumerate(TODOS, 1):
        for j, v in enumerate((item, owner, dl)):
            table.rows[i].cells[j].text = v

    add_header("三、风险与关注")
    for i, (owner, content) in enumerate(RISKS, 1):
        add_item(i, "风险与关注", owner, content)

    os.makedirs(OUT_DIR, exist_ok=True)
    doc.save(os.path.join(OUT_DIR, "纪要.docx"))

    counts = {"决议事项": len(DECISIONS), "待办事项": len(TODOS), "风险与关注": len(RISKS)}
    counts["合计"] = sum(counts.values())
    with io.open(os.path.join(OUT_DIR, "summary.json"), "w", encoding="utf-8", newline="\n") as f:
        f.write(json.dumps(counts, ensure_ascii=False, indent=2) + "\n")

    print("OK: out/纪要.docx + out/summary.json  %s" % counts)


if __name__ == "__main__":
    main()
