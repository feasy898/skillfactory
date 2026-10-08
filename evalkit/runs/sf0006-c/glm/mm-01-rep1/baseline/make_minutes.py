# -*- coding: utf-8 -*-
"""由 fixtures/transcript.txt 生成 out/纪要.docx 与 out/summary.json，并自证纪律约束。

纪律：条目内容/事项全部逐字取自转写稿（允许去掉句末句号）；
相关人/负责人必须是转写稿中出现过的说话人；不得编造。
"""
import json
import re
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt

BASE = Path(__file__).resolve().parent
OUT = BASE / "out"
TRANSCRIPT = BASE / "fixtures" / "transcript.txt"

transcript = TRANSCRIPT.read_text(encoding="utf-8")

# ---------- 解析说话人 ----------
speakers = set()
for m in re.finditer(r"^\[\d{2}:\d{2}\]\s*([^:：]+)\s*[:：]", transcript, re.MULTILINE):
    speakers.add(m.group(1).strip())

# ---------- 条目数据（逐字取自转写稿，句末句号去掉） ----------
DECISIONS = [
    ("赵静怡", "星帆订单系统整体进度正常，会议同意按原计划推进，里程碑保持不变"),
    ("黄雨桐", "会议同意星帆订单系统项目预算300万元，走专项列支，由市场部跟进执行"),
]
TODOS = [
    # (事项, 负责人, 期限)
    ("黄雨桐在本周五前完成复盘纪要并同步给财务部", "黄雨桐", "本周五前"),
    ("徐静怡负责把数据迁移脚本整理成7月4日归档", "徐静怡", "7月4日"),
]
RISKS = [
    ("吴志远", "财务部侧存在履约风险，去年高峰期出现过积压，今年需要提前准备预案"),
    ("赵静怡", "灰度放量方案的数据口径尚未对齐，复盘前需要补齐基线数据"),
]


def assert_verbatim(text: str, what: str) -> None:
    """内容必须逐字存在于转写稿（含或不含句末句号均可）。"""
    ok = (text in transcript) or (text + "。" in transcript) or (text + "。" in transcript.replace("\n", ""))
    assert ok, f"[自证失败] {what} 无法在转写稿中逐字找到: {text!r}"


def assert_speaker(name: str, what: str) -> None:
    assert name in speakers, f"[自证失败] {what} 不是转写稿说话人: {name!r} (说话人={sorted(speakers)})"


# ---------- 生成 docx ----------
OUT.mkdir(exist_ok=True)
doc = Document()

title = doc.add_paragraph()
title.alignment = WD_ALIGN_PARAGRAPH.CENTER
run = title.add_run("星帆订单系统项目例会会议纪要")
run.bold = True
run.font.size = Pt(16)

def add_section_heading(text: str) -> None:
    p = doc.add_paragraph()
    r = p.add_run(text)
    r.bold = True
    r.font.size = Pt(14)

# 一、决议事项
add_section_heading("一、决议事项")
for i, (person, content) in enumerate(DECISIONS, 1):
    doc.add_paragraph(f"{i}. 【决议事项 · {person}】{content}")

# 二、待办事项（表格）
add_section_heading("二、待办事项")
table = doc.add_table(rows=1, cols=3)
table.style = "Table Grid"
for cell, text in zip(table.rows[0].cells, ("事项", "负责人", "期限")):
    cell.text = text
    cell.paragraphs[0].runs[0].bold = True
for item, owner, due in TODOS:
    cells = table.add_row().cells
    cells[0].text = item
    cells[1].text = owner
    cells[2].text = due

# 三、风险与关注
add_section_heading("三、风险与关注")
for i, (person, content) in enumerate(RISKS, 1):
    doc.add_paragraph(f"{i}. 【风险与关注 · {person}】{content}")

docx_path = OUT / "纪要.docx"
doc.save(docx_path)

summary = {
    "决议事项": len(DECISIONS),
    "待办事项": len(TODOS),
    "风险与关注": len(RISKS),
    "合计": len(DECISIONS) + len(TODOS) + len(RISKS),
}
json_path = OUT / "summary.json"
json_path.write_text(json.dumps(summary, ensure_ascii=False) + "\n", encoding="utf-8")

# ---------- 自证 ----------
print("== 自证开始 ==")
for person, content in DECISIONS:
    assert_speaker(person, "决议相关人")
    assert_verbatim(content, "决议内容")
for item, owner, due in TODOS:
    assert_speaker(owner, "待办负责人")
    assert_verbatim(item, "待办事项")
    assert due == "未明确" or due in transcript, f"[自证失败] 期限未见于转写稿: {due!r}"
for person, content in RISKS:
    assert_speaker(person, "风险相关人")
    assert_verbatim(content, "风险内容")

# 回读 docx 复核结构与计数
doc2 = Document(docx_path)
paras = [p.text for p in doc2.paragraphs]
assert paras[0] == "星帆订单系统项目例会会议纪要", f"首段应为标题, 实际: {paras[0]!r}"
for h in ("一、决议事项", "二、待办事项", "三、风险与关注"):
    assert h in paras, f"缺少节标题 {h!r}"
item_re = re.compile(r"^(\d+)\. 【(决议事项|风险与关注) · ([^】]+)】(.+)$")
counts = {"决议事项": 0, "风险与关注": 0}
for p in paras:
    m = item_re.match(p)
    if m:
        n, cat, person, content = m.groups()
        counts[cat] += 1
        assert_speaker(person, f"{cat}条目相关人")
        assert_verbatim(content, f"{cat}条目内容")
        assert int(n) == counts[cat], f"{cat} 条目编号不连续: {p!r}"
assert counts["决议事项"] == len(DECISIONS), "docx 决议条数与数据不符"
assert counts["风险与关注"] == len(RISKS), "docx 风险条数与数据不符"
t = doc2.tables[0]
assert [c.text for c in t.rows[0].cells] == ["事项", "负责人", "期限"], "表头不符"
assert len(t.rows) - 1 == len(TODOS), "待办行数不符"
for row, (item, owner, due) in zip(t.rows[1:], TODOS):
    vals = [c.text for c in row.cells]
    assert vals[0] == item and vals[1] == owner and vals[2] == due, f"待办行不符: {vals!r}"
    assert_verbatim(vals[0], "表格事项")
    assert_speaker(vals[1], "表格负责人")

loaded = json.loads(json_path.read_text(encoding="utf-8"))
assert loaded == {
    "决议事项": counts["决议事项"],
    "待办事项": len(t.rows) - 1,
    "风险与关注": counts["风险与关注"],
    "合计": counts["决议事项"] + (len(t.rows) - 1) + counts["风险与关注"],
}, f"summary.json 与 docx 实际条目数不一致: {loaded}"
assert list(loaded.keys()) == ["决议事项", "待办事项", "风险与关注", "合计"], "summary.json 键名/顺序不符"

print("== 自证全部通过 ==")
print("说话人:", sorted(speakers))
print("summary:", loaded)
print("产出:", docx_path, json_path)
