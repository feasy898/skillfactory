# -*- coding: utf-8 -*-
"""按 brief.md 生成 out/纪要.docx 与 out/summary.json。
所有条目内容均逐字取自 fixtures/transcript.txt。"""
import json
import os

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH

BASE = os.path.dirname(os.path.abspath(__file__))
OUT = os.path.join(BASE, "out")
os.makedirs(OUT, exist_ok=True)

TITLE = "云岭数据看板项目例会纪要"

# 「决议事项」「风险与关注」条目：N. 【类别 · 相关人】内容
DECISIONS = [
    "【决议事项 · 陈静怡】云岭数据看板整体进度正常，会议同意按原计划推进，里程碑保持不变。",
    "【决议事项 · 徐丽娜】会议同意云岭数据看板项目预算16万元，走专项列支，由客户成功部跟进执行。",
]
RISKS = [
    "【风险与关注 · 黄丽娜】综合办侧存在履约风险，去年高峰期出现过积压，今年需要提前准备预案。",
    "【风险与关注 · 陈静怡】灰度放量方案的数据口径尚未对齐，复盘前需要补齐基线数据。",
]
# 「待办事项」表格：事项须为转写稿逐字子串
TODOS = [
    ("完成压测报告并同步给综合办", "徐丽娜", "7月4日前"),
    ("把验收清单整理成下周三归档", "张文博", "下周三"),
]

doc = Document()

title = doc.add_paragraph(TITLE)
title.alignment = WD_ALIGN_PARAGRAPH.CENTER
title.runs[0].bold = True

doc.add_heading("一、决议事项", level=1)
for i, item in enumerate(DECISIONS, 1):
    doc.add_paragraph(f"{i}. {item}")

doc.add_heading("二、待办事项", level=1)
table = doc.add_table(rows=1, cols=3)
table.style = "Table Grid"
hdr = table.rows[0].cells
for cell, text in zip(hdr, ("事项", "负责人", "期限")):
    cell.text = text
for row in TODOS:
    cells = table.add_row().cells
    for cell, text in zip(cells, row):
        cell.text = text

doc.add_heading("三、风险与关注", level=1)
for i, item in enumerate(RISKS, 1):
    doc.add_paragraph(f"{i}. {item}")

docx_path = os.path.join(OUT, "纪要.docx")
doc.save(docx_path)

summary = {
    "决议事项": len(DECISIONS),
    "待办事项": len(TODOS),
    "风险与关注": len(RISKS),
    "合计": len(DECISIONS) + len(TODOS) + len(RISKS),
}
json_path = os.path.join(OUT, "summary.json")
with open(json_path, "w", encoding="utf-8") as f:
    json.dump(summary, f, ensure_ascii=False)

print("saved:", docx_path, json_path)
print(json.dumps(summary, ensure_ascii=False))
