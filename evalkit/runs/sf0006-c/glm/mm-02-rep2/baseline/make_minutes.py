#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""按 brief.md 从 fixtures/transcript.txt 生成 out/纪要.docx 与 out/summary.json。

纪律：转写稿为唯一事实来源——所有条目内容逐字取自转写稿（句末句号可省），
不补造 [HH:MM] 时间戳，不虚构同名说话人区分标记。
同名说话人：周丽娜（稿内多个角色共用该姓名、无任何区分标记，其条目无法唯一归属，
按 brief 追加（归属待核）后缀）。
"""

import io
import json
import os
import re

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt

HERE = os.path.dirname(os.path.abspath(__file__))
OUT_DIR = os.path.join(HERE, "out")
TRANSCRIPT = io.open(os.path.join(HERE, "fixtures", "transcript.txt"), encoding="utf-8").read()

SAME_NAME = "周丽娜"
ATTRIB_SUFFIX = "（归属待核）"

# （相关人, 内容）——内容均为转写稿逐字原句（省句末句号）
DECISIONS = [
    ("周丽娜", "云岭数据看板整体进度正常，会议同意按原计划推进，里程碑保持不变"),
    ("徐志远", "会议同意云岭数据看板项目预算45万元，走专项列支，由财务部跟进执行"),
    ("吴雨桐", "云岭数据看板的验收标准按新版执行，旧版同时作废"),
]
RISKS = [
    ("吴志远", "供应链组侧存在履约风险，去年高峰期出现过积压，今年需要提前准备预案"),
    ("周丽娜", "灰度放量方案的数据口径尚未对齐，复盘前需要补齐基线数据"),
]
# （事项, 负责人, 期限）——事项逐字取自转写稿；无时限者期限=未明确
TODOS = [
    ("徐志远在下周一前完成应急预案演练脚本并同步给供应链组", "徐志远", "下周一前"),
    ("刘晓雯负责把验收清单整理成本周五归档", "刘晓雯", "本周五"),
    ("吴志远牵头梳理云岭数据看板的责任分工", "吴志远", "未明确"),
]


def selfcheck():
    """写盘前自证：逐字溯源、相关人均为说话人、无补造时间戳。"""
    speakers = set()
    for line in TRANSCRIPT.splitlines():
        m = re.match(r"^(?:\[\d{1,2}:\d{2}\]\s*)?([^:]{1,12}):\s", line)
        if m:
            speakers.add(m.group(1).strip())
    texts = []
    for owner, content in DECISIONS + RISKS:
        assert content in TRANSCRIPT, "决议/风险内容不可逐字溯源: %s" % content
        assert owner in speakers, "相关人非转写稿说话人: %s" % owner
        texts.append(content)
    for item, owner, dl in TODOS:
        assert item in TRANSCRIPT, "待办事项不可逐字溯源: %s" % item
        assert owner in speakers, "负责人非转写稿说话人: %s" % owner
        assert dl == "未明确" or dl in item, "期限非该条来源: %s" % dl
        texts.append(item)
    texts.append(SAME_NAME)
    for t in texts:
        assert not re.search(r"\[\d{1,2}:\d{2}\]", t), "出现补造时间戳: %s" % t
    return speakers


def entry_line(n, cat, owner, content):
    suffix = ATTRIB_SUFFIX if owner == SAME_NAME else ""
    return "%d. 【%s · %s】%s%s" % (n, cat, owner, content, suffix)


def build():
    doc = Document()

    title = doc.add_paragraph("云岭数据看板项目例会会议纪要")
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    for run in title.runs:
        run.bold = True
        run.font.size = Pt(16)

    doc.add_paragraph("一、决议事项")
    for n, (owner, content) in enumerate(DECISIONS, 1):
        doc.add_paragraph(entry_line(n, "决议事项", owner, content))

    doc.add_paragraph("二、待办事项")
    table = doc.add_table(rows=1 + len(TODOS), cols=3)
    try:
        table.style = "Table Grid"
    except KeyError:
        pass
    for j, header in enumerate(("事项", "负责人", "期限")):
        cell = table.rows[0].cells[j]
        cell.text = header
        for p in cell.paragraphs:
            for run in p.runs:
                run.bold = True
    for i, (item, owner, dl) in enumerate(TODOS, 1):
        table.rows[i].cells[0].text = item
        table.rows[i].cells[1].text = owner
        table.rows[i].cells[2].text = dl

    doc.add_paragraph("三、风险与关注")
    for n, (owner, content) in enumerate(RISKS, 1):
        doc.add_paragraph(entry_line(n, "风险与关注", owner, content))

    counts = {"决议事项": len(DECISIONS), "待办事项": len(TODOS),
              "风险与关注": len(RISKS)}
    counts["合计"] = sum(counts.values())

    os.makedirs(OUT_DIR, exist_ok=True)
    doc.save(os.path.join(OUT_DIR, "纪要.docx"))
    with io.open(os.path.join(OUT_DIR, "summary.json"), "w", encoding="utf-8", newline="\n") as f:
        f.write(json.dumps(counts, ensure_ascii=False, indent=2) + "\n")
    return counts


if __name__ == "__main__":
    selfcheck()
    print("speakers:", sorted(selfcheck()))
    print("counts:", json.dumps(build(), ensure_ascii=False))
